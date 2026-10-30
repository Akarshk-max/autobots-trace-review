# tests/ingestion/unit/test_spreadsheet_ingestion_and_polling.py
"""
Exhaustive test suite for spreadsheet ingestion.

Coverage goals:
- Locator encoding/decoding + parsing (including partial/legacy locators)
- Header region detection, multi-row headers, and "dense title row" penalty
- Blank-row segmentation behavior and repeated-header skipping
- Multi-table handling including the tricky EOF-after-break case (table vs bottom notes)
- Span expansion logic (header span too narrow, data extends right)
- Side-note emission bounds (left/right col windows and min-chars threshold)
- Render chunking no-loss behavior: row/col splitting and oversize single-cell rejection
- Merge map building behavior (skips oversized ranges, warning emitted)
- apply_merge_fill_to_row semantics (no padding)
- XLSX zip preflight: sharedStrings.xml guard (with monkeypatched small caps)
- Polling readiness_probe logic for Weaviate and Neo4j:
  - UUID dedup
  - thresholds
  - timeout behavior
  - transient errors

Notes:
- No external Weaviate/Neo4j required; fakes used.
- Openpyxl-specific tests are skipped if openpyxl isn't installed.
"""

from __future__ import annotations

import asyncio
import re
import types
import uuid
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pytest

from src.autobots.ingestion.polling import readiness_probe
from src.autobots.ingestion.spreadsheet.layout import (
    detect_header_region,
    flatten_header_rows,
    infer_table_span_from_header_rows,
    is_blank_row,
    row_matches_header,
)
from src.autobots.ingestion.spreadsheet.locator import (
    build_source_locator,
    decode_sheet_name,
    encode_sheet_name,
    parse_source_locator,
)
from src.autobots.ingestion.spreadsheet.merges import (
    apply_merge_fill_to_row,
    build_merge_fill_map,
)
from src.autobots.ingestion.spreadsheet.pipeline import (
    SheetIngestConfig,
    ingest_sheet_from_rows_iter,
)
from src.autobots.ingestion.spreadsheet.render import emit_row_chunks_no_loss

# -------------------------------------------------------------------------------------------------
# Helpers for tests
# -------------------------------------------------------------------------------------------------

def _excel_col_to_index_0(col_letters: str) -> int:
    """
    Excel letters (A, B, Z, AA, AB, ...) -> 0-based index.
    """
    s = (col_letters or "").strip().upper()
    if not s or not re.fullmatch(r"[A-Z]+", s):
        return -1
    n = 0
    for ch in s:
        n = n * 26 + (ord(ch) - ord("A") + 1)
    return n - 1


def _cols_range_to_span_0(cols: Optional[str]) -> Tuple[int, int]:
    """
    'A-E' -> (0,4) inclusive span. 'C' -> (2,2).
    Returns (-1,-1) on parse failure.
    """
    if not cols:
        return (-1, -1)
    s = cols.strip().upper()
    if "-" in s:
        a, b = s.split("-", 1)
        ia = _excel_col_to_index_0(a.strip())
        ib = _excel_col_to_index_0(b.strip())
        return (ia, ib)
    i = _excel_col_to_index_0(s)
    return (i, i)


def _mk_cfg(
    *,
    header_scan_rows: int = 10,
    blank_row_break: int = 3,
    blank_row_empty_ratio: float = 0.95,
    max_rows_per_sheet: int = 10_000,
    max_rows_per_chunk: int = 3,
    max_cols_included: int = 20,
    max_cols_hard: int = 50,
    max_chars_per_chunk: int = 20_000,
    forward_fill_cols: int = 2,
    enable_side_notes: bool = False,
    side_note_max_cols: int = 5,
    side_note_min_chars: int = 30,
) -> SheetIngestConfig:
    return SheetIngestConfig(
        header_scan_rows=header_scan_rows,
        blank_row_break=blank_row_break,
        blank_row_empty_ratio=blank_row_empty_ratio,
        max_rows_per_sheet=max_rows_per_sheet,
        max_rows_per_chunk=max_rows_per_chunk,
        max_cols_included=max_cols_included,
        max_cols_hard=max_cols_hard,
        max_chars_per_chunk=max_chars_per_chunk,
        forward_fill_cols=forward_fill_cols,
        enable_side_notes=enable_side_notes,
        side_note_max_cols=side_note_max_cols,
        side_note_min_chars=side_note_min_chars,
    )


def _collect_emitted():
    emitted: List[Dict[str, Any]] = []

    def add_doc_cb(
        text: str,
        *,
        doc_type: str,
        source_locator: str,
        order_val: int,
        seed: str,
    ):
        emitted.append(
            {
                "doc_type": doc_type,
                "locator": source_locator,
                "order": order_val,
                "seed": seed,
                "text": text,
            }
        )

    return emitted, add_doc_cb


# -------------------------------------------------------------------------------------------------
# Locator tests
# -------------------------------------------------------------------------------------------------

@pytest.mark.parametrize(
    "sheet_name",
    [
        "Simple",
        "USA-GIE",
        "Name with spaces",
        "Semi;Colon",
        "Unicode—测试",
        "Percent%Sign",
        "Weird=Equals",
        "Slash/Back\\",
    ],
)
def test_sheet_name_encode_decode_roundtrip(sheet_name: str):
    enc = encode_sheet_name(sheet_name)
    dec = decode_sheet_name(enc)
    assert dec == sheet_name


def test_source_locator_roundtrip_all_fields():
    loc = build_source_locator(
        sheet_name="Sheet;1",
        kind="rows",
        table=2,
        rows=(10, 20),
        cols="A-D",
        header_rows=(3, 4),
    )
    info = parse_source_locator(loc)
    assert info.sheet == "Sheet;1"
    assert info.kind == "rows"
    assert info.table == 2
    assert info.rows == (10, 20)
    assert info.cols == "A-D"
    assert info.header_rows == (3, 4)


def test_parse_source_locator_backward_compat_missing_fields():
    # Legacy-ish: no kind/table/cols/header_rows
    legacy = "Sheet=USA-GIE; rows=1-80"
    info = parse_source_locator(legacy)
    assert info.sheet == "USA-GIE"
    assert info.kind is None
    assert info.table == 1
    assert info.rows == (1, 80)
    assert info.cols is None
    assert info.header_rows is None


def test_parse_source_locator_handles_missing_sheet_key():
    info = parse_source_locator("kind=rows; table=7; rows=1-2")
    assert info.sheet == "__UNKNOWN__"
    assert info.table == 7
    assert info.rows == (1, 2)


# -------------------------------------------------------------------------------------------------
# Layout/header tests
# -------------------------------------------------------------------------------------------------

def test_flatten_header_rows_two_levels():
    rows = [
        ["Revenue", "Revenue", "Cost", "Cost"],
        ["Q1", "Q2", "Q1", "Q2"],
    ]
    flat = flatten_header_rows(rows, max_cols=4)
    assert flat == ["Revenue / Q1", "Revenue / Q2", "Cost / Q1", "Cost / Q2"]


def test_infer_table_span_from_header_rows():
    rows = [
        [None, None, "A", "B", None],
        [None, None, None, "C", "D"],
    ]
    c0, c1 = infer_table_span_from_header_rows(rows, max_cols=5)
    # leftmost header at col 2, rightmost at col 4 => end_excl should be 5
    assert c0 == 2
    assert c1 == 5


def test_detect_header_region_prefers_real_header_over_dense_title_row():
    # Row 1: dense repeated "Holdings Report" (simulated merged fill)
    # Row 2: real header
    preview_rows = [
        ["Holdings Report"] * 20,
        ["Name", "Ticker", "Qty", "Price", "Value"] + [None] * 15,
        ["Apple", "AAPL", 10, 100, 1000] + [None] * 15,
    ]
    preview_nums = [1, 2, 3]
    hdr = detect_header_region(preview_rows, preview_nums, max_cols_hard=60, max_header_rows=3)
    assert hdr is not None
    assert hdr.header_rows_1based[0] == 2
    assert hdr.headers_flat[0].lower().startswith("name")


def test_detect_header_region_expands_upward_for_group_header():
    # row 1 = group header (sparse but stringy), row 2 = real header
    preview_rows = [
            # Make this genuinely "group header"-like per detect_header_region() heuristic:
            # sparse (density <= 0.40) but stringy.
            [None, "Revenue", None, None, "Cost", None],
        ["Region", "Q1", "Q2", "Q3", "Q1", "Q2"],
        ["NA", 1, 2, 3, 4, 5],
    ]
    preview_nums = [10, 11, 12]
    hdr = detect_header_region(preview_rows, preview_nums, max_cols_hard=60, max_header_rows=3)
    assert hdr is not None
    assert hdr.header_rows_1based == (10, 11)  # consumed two header rows


def test_row_matches_header_detects_repeated_header_leaf_match():
    headers_flat = ["Revenue / Q1", "Revenue / Q2", "Cost / Q1"]
    # repeated header row in data: leaves match
    row = ["Q1", "Q2", "Q1"]
    assert row_matches_header(row, headers_flat, col_start_0=0, col_end_excl_0=3) is True


def test_is_blank_row_key_cols_prevents_false_breaks():
    # A wide sparse table: first col has a label, rest empty.
    # Should NOT be treated as blank if key_cols includes first column.
    row = ["RowLabel"] + [None] * 19
    assert is_blank_row(row, col_start_0=0, col_end_excl_0=20, empty_ratio_threshold=0.95, key_cols=2) is False
    # But if key_cols=0, it becomes blank-ish
    assert is_blank_row(row, col_start_0=0, col_end_excl_0=20, empty_ratio_threshold=0.95, key_cols=0) is True


# -------------------------------------------------------------------------------------------------
# Merge fill tests (pure)
# -------------------------------------------------------------------------------------------------

@dataclass
class _FakeRange:
    min_row: int
    max_row: int
    min_col: int
    max_col: int


class _FakeMergedCells:
    def __init__(self, ranges):
        self.ranges = ranges


class _FakeWS:
    def __init__(self, anchor_values: Dict[Tuple[int, int], object], ranges):
        self._anchors = anchor_values
        self.merged_cells = _FakeMergedCells(ranges)

    def cell(self, row: int, column: int):
        class _Cell:
            def __init__(self, v):
                self.value = v
        return _Cell(self._anchors.get((row, column)))


def test_build_merge_fill_map_skips_oversized_range_and_emits_warning():
    warnings: List[str] = []
    ws = _FakeWS(
        # IMPORTANT: the oversized range must have a non-empty anchor cell (row=3,col=1),
        # otherwise build_merge_fill_map() will skip it *before* counting it as "skipped due to cap".
        anchor_values={(1, 1): "TITLE", (2, 2): "Group", (3, 1): "BIG"},
        ranges=[
            _FakeRange(1, 1, 1, 10),   # 10 cells
            _FakeRange(2, 2, 2, 3),    # 2 cells
            _FakeRange(3, 500, 1, 50), # huge, should be skipped by max_cells
        ],
    )
    mm = build_merge_fill_map(
        ws,
        max_row_1based=10,
        max_col_1based=10,
        max_cells=20,  # will skip huge range
        warnings=warnings,
        sheet_name="S1",
    )
    # Ensure small ranges included
    assert mm[(1, 1)] == "TITLE"
    assert mm[(1, 10)] == "TITLE"
    assert mm[(2, 2)] == "Group"
    assert mm[(2, 3)] == "Group"
    # Ensure a warning about skipped ranges
    assert any("skipped" in w.lower() for w in warnings)


def test_apply_merge_fill_to_row_does_not_pad_row():
    merge_map = {(1, 5): "X"}  # would require padding if we padded
    row = [None, None, None]   # len 3
    out = apply_merge_fill_to_row(1, row, merge_map, max_col_1based=50)
    # Must not grow the row; no padding should occur
    assert len(out) == 3


# -------------------------------------------------------------------------------------------------
# Render tests (no-loss chunking)
# -------------------------------------------------------------------------------------------------

def test_emit_row_chunks_no_loss_splits_by_rows_when_needed():
    emitted, add_doc_cb = _collect_emitted()
    headers = ["A", "B", "C"]
    buf = [(i, ["x" * 50, "y" * 50, "z" * 50]) for i in range(1, 10)]
    # Force small chunk size so it can't take many rows
    emit_row_chunks_no_loss(
        headers=headers,
        buffer=buf,
        workbook_name="W",
        sheet_name="S",
        table_id=1,
        header_rows_1based=(1, 1),
        kind="rows",
        doc_type="spreadsheet_rows",
        add_doc_cb=add_doc_cb,
        global_order_ref=[0],
        max_rows_per_chunk=10,
        max_cols_included=10,
        max_cols_hard=50,
        max_chars_per_chunk=700,
        col_span=(0, 3),
    )
    assert len(emitted) >= 2
    # buffer must be consumed
    assert buf == []


def test_emit_row_chunks_no_loss_splits_by_columns_when_needed():
    emitted, add_doc_cb = _collect_emitted()
    headers = [f"H{i}" for i in range(50)]
    # One row with many populated columns
    row = ["v" * 10] * 50
    buf = [(1, row)]
    emit_row_chunks_no_loss(
        headers=headers,
        buffer=buf,
        workbook_name="W",
        sheet_name="S",
        table_id=1,
        header_rows_1based=(1, 1),
        kind="rows",
        doc_type="spreadsheet_rows",
        add_doc_cb=add_doc_cb,
        global_order_ref=[0],
        max_rows_per_chunk=5,
        max_cols_included=10,  # forces windowing
        max_cols_hard=300,
        max_chars_per_chunk=800,
        col_span=(0, 50),
    )
    assert len(emitted) >= 2
    # locators include cols ranges
    assert any("cols=" in e["locator"] for e in emitted)


def test_emit_row_chunks_no_loss_raises_on_single_cell_oversize():
    emitted, add_doc_cb = _collect_emitted()
    headers = ["A"]
    buf = [(1, ["X" * 50_000])]
    with pytest.raises(ValueError):
        emit_row_chunks_no_loss(
            headers=headers,
            buffer=buf,
            workbook_name="W",
            sheet_name="S",
            table_id=1,
            header_rows_1based=(1, 1),
            kind="rows",
            doc_type="spreadsheet_rows",
            add_doc_cb=add_doc_cb,
            global_order_ref=[0],
            max_rows_per_chunk=1,
            max_cols_included=1,
            max_cols_hard=10,
            max_chars_per_chunk=2000,
            col_span=(0, 1),
        )


# -------------------------------------------------------------------------------------------------
# Pipeline behavior tests
# -------------------------------------------------------------------------------------------------

def test_pipeline_emits_sheet_summary_once_even_if_multiple_note_blocks():
    rows = [
        (1, ["Note 1", None]),
        (2, ["Note 2", None]),
        (3, ["More notes", None]),
        (4, [None, None]),
        (5, ["Even more notes", None]),
    ]
    emitted, add_doc_cb = _collect_emitted()
    cfg = _mk_cfg(header_scan_rows=3, enable_side_notes=False)

    ingest_sheet_from_rows_iter(
        rows_iter=iter(rows),
        workbook_name="W",
        sheet_name="S",
        add_doc_cb=add_doc_cb,
        global_order_ref=[0],
        cfg=cfg,
        warnings=[],
    )

    summaries = [e for e in emitted if e["doc_type"] == "spreadsheet_sheet_summary"]
    assert len(summaries) == 1


def test_pipeline_single_table_with_top_notes_emits_notes_top_and_rows():
    rows = [
        (1, ["Q4 Report", None, None]),
        (2, ["All values USD", None, None]),
        (3, ["Name", "Ticker", "Qty"]),
        (4, ["Apple", "AAPL", 10]),
        (5, ["MSFT", "MSFT", 5]),
    ]
    emitted, add_doc_cb = _collect_emitted()
    cfg = _mk_cfg(header_scan_rows=5, max_rows_per_chunk=2, enable_side_notes=False)

    ingest_sheet_from_rows_iter(
        rows_iter=iter(rows),
        workbook_name="W",
        sheet_name="Sheet;1",
        add_doc_cb=add_doc_cb,
        global_order_ref=[0],
        cfg=cfg,
        warnings=[],
    )

    assert any(e["doc_type"] == "spreadsheet_sheet_summary" for e in emitted)
    assert any("kind=notes_top" in e["locator"] for e in emitted)
    assert any("kind=rows" in e["locator"] and "table=1" in e["locator"] for e in emitted)


def test_pipeline_multi_table_eof_tail_is_table_emits_table2_rows():
    # This is the failure you observed: EOF occurs soon after table break, but tail is actually a 2nd table.
    rows = [
        (1, ["Q4 Report", None, None]),
        (2, ["All values USD", None, None]),
        (3, ["Revenue", "Revenue", "Cost"]),
        (4, ["Q1", "Q2", "Q1"]),
        (5, [100, 200, 50]),
        (6, [110, 210, 55]),
        (7, [120, 220, 60]),
        (8, [None, None, None]),
        (9, [None, None, None]),
        (10, [None, None, None]),
        (11, ["Region", "Sales"]),
        (12, ["NA", 1000]),
        (13, ["EU", 900]),
    ]
    emitted, add_doc_cb = _collect_emitted()
    cfg = _mk_cfg(header_scan_rows=10, max_rows_per_chunk=3, enable_side_notes=False)

    ingest_sheet_from_rows_iter(
        rows_iter=iter(rows),
        workbook_name="W",
        sheet_name="Sheet;1",
        add_doc_cb=add_doc_cb,
        global_order_ref=[0],
        cfg=cfg,
        warnings=[],
    )

    assert any("kind=rows" in e["locator"] and "table=1" in e["locator"] for e in emitted)
    assert any("kind=rows" in e["locator"] and "table=2" in e["locator"] for e in emitted)


def test_pipeline_multi_table_eof_tail_is_notes_emits_notes_bottom():
    rows = [
        (1, ["Title", None]),
        (2, ["A", "B"]),
        (3, [1, 2]),
        (4, [None, None]),
        (5, [None, None]),
        (6, [None, None]),
        (7, ["footnote: generated", None]),
        (8, ["units: USD", None]),
    ]
    emitted, add_doc_cb = _collect_emitted()
    cfg = _mk_cfg(header_scan_rows=6, max_rows_per_chunk=3, enable_side_notes=False)

    ingest_sheet_from_rows_iter(
        rows_iter=iter(rows),
        workbook_name="W",
        sheet_name="S",
        add_doc_cb=add_doc_cb,
        global_order_ref=[0],
        cfg=cfg,
        warnings=[],
    )

    assert any("kind=notes_bottom" in e["locator"] for e in emitted)


def test_pipeline_repeated_header_row_inside_data_is_skipped():
    rows = [
        (1, ["Name", "Ticker", "Qty"]),
        (2, ["Apple", "AAPL", 10]),
        (3, ["Name", "Ticker", "Qty"]),  # repeated header
        (4, ["MSFT", "MSFT", 5]),
    ]
    emitted, add_doc_cb = _collect_emitted()
    cfg = _mk_cfg(header_scan_rows=4, max_rows_per_chunk=10, enable_side_notes=False)

    ingest_sheet_from_rows_iter(
        rows_iter=iter(rows),
        workbook_name="W",
        sheet_name="S",
        add_doc_cb=add_doc_cb,
        global_order_ref=[0],
        cfg=cfg,
        warnings=[],
    )

    # Ensure row chunks exist
    row_docs = [e for e in emitted if e["doc_type"] == "spreadsheet_rows"]
    assert row_docs
    text_all = "\n".join(e["text"] for e in row_docs)
    # Repeated header row should not appear as a data row line with Name=Name etc.
    assert "Name=Name" not in text_all


def test_pipeline_span_expansion_includes_data_beyond_sparse_header():
    # Header has labels only in first 2 columns, but data includes columns 3..5.
    # IMPORTANT: provide enough data rows so that even if header detection uses 2 header rows,
    # we still have >=3 rows in scan_rows to meet the current expansion threshold (>=3).
    rows = [
        (1, ["A", "B", None, None, None]),
        (2, [1, 2, "X", "Y", "Z"]),
        (3, [3, 4, "P", "Q", "R"]),
        (4, [5, 6, "L", "M", "N"]),
        (5, [7, 8, "U", "V", "W"]),
        (6, [9, 10, "S", "T", "U"]),
    ]
    emitted, add_doc_cb = _collect_emitted()
    cfg = _mk_cfg(
        header_scan_rows=6,
        max_rows_per_chunk=10,
        max_cols_included=20,  # allow wide window so locator reflects full span
        enable_side_notes=False,
    )

    ingest_sheet_from_rows_iter(
        rows_iter=iter(rows),
        workbook_name="W",
        sheet_name="S",
        add_doc_cb=add_doc_cb,
        global_order_ref=[0],
        cfg=cfg,
        warnings=[],
    )

    # We expect the table span to expand so at least one rows chunk includes columns through E (index 4).
    row_docs = [e for e in emitted if e["doc_type"] == "spreadsheet_rows"]
    row_locs = [e["locator"] for e in row_docs]
    assert row_locs
    max_end = -1
    for loc in row_locs:
        info = parse_source_locator(loc)
        a, b = _cols_range_to_span_0(info.cols)
        if b > max_end:
            max_end = b
    assert max_end >= 4, f"Expected span expansion to reach at least col E; got max_end={max_end}, locs={row_locs}"

    # And the rendered text should show C/D/E keys for those columns (because headers were synthesized).
    text = "\n".join(e["text"] for e in row_docs)
    assert ("C=" in text) and ("D=" in text) and ("E=" in text)


def test_pipeline_side_notes_emitted_and_bounded_to_max_cols():
    # Table span should be cols 1-2, but notes appear in far right.
    # IMPORTANT:
    # - Do NOT put "RightNoteHeader" in the header row; otherwise the header span will include that column.
    # - Ensure the note column is NOT consistently populated in >=3 scanned rows, or span expansion will absorb it.
    rows = [
        (1, ["H1", "H2", None, None, None]),
        (2, ["A", 1, None, None, "This is a right note with enough chars"]),
        (3, ["B", 2, None, None, None]),
        (4, ["C", 3, None, None, "Another right note with enough chars"]),
        (5, ["D", 4, None, None, None]),
    ]
    emitted, add_doc_cb = _collect_emitted()
    cfg = _mk_cfg(
        header_scan_rows=5,
        max_rows_per_chunk=10,
        max_cols_included=10,
        enable_side_notes=True,
        side_note_max_cols=2,
        side_note_min_chars=10,
    )

    ingest_sheet_from_rows_iter(
        rows_iter=iter(rows),
        workbook_name="W",
        sheet_name="S",
        add_doc_cb=add_doc_cb,
        global_order_ref=[0],
        cfg=cfg,
        warnings=[],
    )

    # Must have main rows
    assert any(e["doc_type"] == "spreadsheet_rows" for e in emitted), (
        "Expected spreadsheet_rows docs but none were emitted. "
        f"Emitted locators: {[e['locator'] for e in emitted]}"
    )

    # Must have side right notes
    right_notes = [e for e in emitted if e["doc_type"] == "spreadsheet_notes" and "kind=notes_side_right" in e["locator"]]
    assert right_notes

    # Ensure right-note locators are bounded to <= side_note_max_cols columns
    for e in right_notes:
        info = parse_source_locator(e["locator"])
        a, b = _cols_range_to_span_0(info.cols)
        assert a >= 0 and b >= 0
        width = (b - a + 1)
        assert width <= cfg.side_note_max_cols, f"Expected bounded side-note width <= {cfg.side_note_max_cols}, got {width} in {e['locator']}"


def test_pipeline_side_notes_not_emitted_if_below_min_chars():
    rows = [
        (1, ["H1", "H2", None, "N"]),
        (2, ["A", "B", None, "x"]),  # too short
        (3, ["C", "D", None, "y"]),  # too short
    ]
    emitted, add_doc_cb = _collect_emitted()
    cfg = _mk_cfg(
        header_scan_rows=3,
        enable_side_notes=True,
        side_note_max_cols=2,
        side_note_min_chars=10,  # blocks short notes
    )

    ingest_sheet_from_rows_iter(
        rows_iter=iter(rows),
        workbook_name="W",
        sheet_name="S",
        add_doc_cb=add_doc_cb,
        global_order_ref=[0],
        cfg=cfg,
        warnings=[],
    )

    assert not any("kind=notes_side_right" in e["locator"] for e in emitted)
    assert not any("kind=notes_side_left" in e["locator"] for e in emitted)


def test_pipeline_orders_monotonic_and_increasing():
    rows = [
        (1, ["Name", "Qty"]),
        (2, ["A", 1]),
        (3, ["B", 2]),
        (4, ["C", 3]),
    ]
    emitted, add_doc_cb = _collect_emitted()
    cfg = _mk_cfg(header_scan_rows=4, max_rows_per_chunk=1, enable_side_notes=False)

    ingest_sheet_from_rows_iter(
        rows_iter=iter(rows),
        workbook_name="W",
        sheet_name="S",
        add_doc_cb=add_doc_cb,
        global_order_ref=[0],
        cfg=cfg,
        warnings=[],
    )

    orders = [e["order"] for e in emitted]
    assert orders == sorted(orders)
    assert len(set(orders)) == len(orders)


# -------------------------------------------------------------------------------------------------
# XLSX zip preflight tests (sharedStrings guard)
# -------------------------------------------------------------------------------------------------

def test_xlsx_zip_preflight_rejects_large_shared_strings(monkeypatch, tmp_path: Path):
    # Build a tiny zip, but monkeypatch caps so it rejects at small size.
    from src.autobots.ingestion.spreadsheet import ingest as ingest_mod

    # Patch caps down for test
    monkeypatch.setattr(ingest_mod, "SHARED_STRINGS_MAX_ABS_BYTES", 1024)  # 1KB
    monkeypatch.setattr(ingest_mod, "SHARED_STRINGS_MAX_FRACTION_OF_CAP", 1.0)
    monkeypatch.setattr(ingest_mod, "SHARED_STRINGS_MIN_BYTES", 0)

    xlsx_like = tmp_path / "fake.xlsx"
    with zipfile.ZipFile(xlsx_like, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/sharedStrings.xml", b"x" * 5000)  # 5KB > 1KB cap
        zf.writestr("xl/workbook.xml", b"<xml/>")

    with pytest.raises(ValueError) as e:
        ingest_mod.xlsx_zip_preflight_or_reject(xlsx_like, compressed_bytes=xlsx_like.stat().st_size)
    assert "sharedStrings" in str(e.value)


# -------------------------------------------------------------------------------------------------
# readiness_probe tests (Weaviate + Neo4j fakes)
# -------------------------------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_readiness_probe_zero_chunks_is_success():
    ok = await readiness_probe(
        collection_name="C",
        user_id="t",
        chunk_uuids=[],
        targets=["weaviate", "graphrag"],
        timeout_seconds=1,
        success_threshold=1.0,
    )
    assert ok is True


@pytest.mark.asyncio
async def test_readiness_probe_deduplicates_chunk_ids(monkeypatch):
    # Make sleep and to_thread deterministic/fast
    async def no_sleep(_):
        return None
    async def fake_to_thread(fn, *args, **kwargs):
        return fn(*args, **kwargs)
    monkeypatch.setattr(asyncio, "sleep", no_sleep)
    monkeypatch.setattr(asyncio, "to_thread", fake_to_thread)

    # Patch wvc filter builder
    class _ById:
        def contains_any(self, batch):
            return ("contains_any", list(batch))
    class _Filter:
        @staticmethod
        def by_id():
            return _ById()
    fake_wvc = types.SimpleNamespace(query=types.SimpleNamespace(Filter=_Filter))
    import src.autobots.ingestion.polling as polling_mod
    monkeypatch.setattr(polling_mod, "wvc", fake_wvc)

    # Fake Weaviate: always returns everything it is asked for
    class FakeResp:
        def __init__(self, n):
            self.objects = [object()] * n

    class FakeQuery:
        def fetch_objects(self, *, filters, limit, return_properties):
            batch = filters[1]
            return FakeResp(len(batch))

    class FakeCollection:
        def __init__(self):
            self.query = FakeQuery()
        def with_tenant(self, tenant_id: str):
            return self

    class FakeWeaviateSDK:
        def __init__(self):
            self.collections = types.SimpleNamespace(get=lambda name: FakeCollection())

    class FakeWeaviateClient:
        def __init__(self):
            self.client = FakeWeaviateSDK()
        async def close(self):
            return None

    # Fake Neo4j: returns everything
    class FakeNeo4jClient:
        def query(self, query: str, tenant_id: str, params: Dict[str, Any]):
            return [{"count": len(set(params["chunkIds"]))}]

    async def fake_get_client():
        return FakeNeo4jClient()

    import src.autobots.ingestion.neo4j_loader as neo4j_loader_mod
    monkeypatch.setattr(neo4j_loader_mod.Neo4jManager, "get_client", staticmethod(fake_get_client))

    a = str(uuid.uuid4())
    chunk_ids = [a, a, a]  # duplicates
    ok = await readiness_probe(
        collection_name="C",
        user_id="t",
        chunk_uuids=chunk_ids,
        targets=["weaviate", "graphrag"],
        existing_weaviate_client=FakeWeaviateClient(),
        timeout_seconds=1,
        success_threshold=1.0,
    )
    assert ok is True


@pytest.mark.asyncio
async def test_readiness_probe_weaviate_threshold_failure(monkeypatch):
    async def no_sleep(_):
        return None
    async def fake_to_thread(fn, *args, **kwargs):
        return fn(*args, **kwargs)
    monkeypatch.setattr(asyncio, "sleep", no_sleep)
    monkeypatch.setattr(asyncio, "to_thread", fake_to_thread)

    class _ById:
        def contains_any(self, batch):
            return ("contains_any", list(batch))
    class _Filter:
        @staticmethod
        def by_id():
            return _ById()
    fake_wvc = types.SimpleNamespace(query=types.SimpleNamespace(Filter=_Filter))
    import src.autobots.ingestion.polling as polling_mod
    monkeypatch.setattr(polling_mod, "wvc", fake_wvc)

    class FakeResp:
        def __init__(self, n):
            self.objects = [object()] * n

    class FakeQuery:
        def __init__(self, present_ids: set):
            self.present_ids = present_ids
        def fetch_objects(self, *, filters, limit, return_properties):
            batch = filters[1]
            found = sum(1 for x in batch if x in self.present_ids)
            return FakeResp(found)

    class FakeCollection:
        def __init__(self, present_ids: set):
            self.query = FakeQuery(present_ids)
        def with_tenant(self, tenant_id: str):
            return self

    class FakeWeaviateSDK:
        def __init__(self, present_ids: set):
            self.collections = types.SimpleNamespace(get=lambda name: FakeCollection(present_ids))

    class FakeWeaviateClient:
        def __init__(self, present_ids: set):
            self.client = FakeWeaviateSDK(present_ids)
        async def close(self):
            return None

    chunk_ids = [str(uuid.uuid4()) for _ in range(20)]
    present_ids = set(chunk_ids[:5])  # 25% present; threshold 75% should fail
    ok = await readiness_probe(
        collection_name="C",
        user_id="t",
        chunk_uuids=chunk_ids,
        targets=["weaviate"],
        existing_weaviate_client=FakeWeaviateClient(present_ids),
        timeout_seconds=1,
        success_threshold=0.75,
    )
    assert ok is False


@pytest.mark.asyncio
async def test_readiness_probe_weaviate_transient_tenant_not_found_then_success(monkeypatch):
    async def no_sleep(_):
        return None
    async def fake_to_thread(fn, *args, **kwargs):
        return fn(*args, **kwargs)
    monkeypatch.setattr(asyncio, "sleep", no_sleep)
    monkeypatch.setattr(asyncio, "to_thread", fake_to_thread)

    class _ById:
        def contains_any(self, batch):
            return ("contains_any", list(batch))
    class _Filter:
        @staticmethod
        def by_id():
            return _ById()
    fake_wvc = types.SimpleNamespace(query=types.SimpleNamespace(Filter=_Filter))
    import src.autobots.ingestion.polling as polling_mod
    monkeypatch.setattr(polling_mod, "wvc", fake_wvc)

    calls = {"n": 0}

    class FakeResp:
        def __init__(self, n):
            self.objects = [object()] * n

    class FakeQuery:
        def fetch_objects(self, *, filters, limit, return_properties):
            calls["n"] += 1
            if calls["n"] == 1:
                raise Exception("tenant not found")  # triggers transient path
            batch = filters[1]
            return FakeResp(len(batch))

    class FakeCollection:
        def __init__(self):
            self.query = FakeQuery()
        def with_tenant(self, tenant_id: str):
            return self

    class FakeWeaviateSDK:
        def __init__(self):
            self.collections = types.SimpleNamespace(get=lambda name: FakeCollection())

    class FakeWeaviateClient:
        def __init__(self):
            self.client = FakeWeaviateSDK()
        async def close(self):
            return None

    chunk_ids = [str(uuid.uuid4()) for _ in range(5)]
    ok = await readiness_probe(
        collection_name="C",
        user_id="t",
        chunk_uuids=chunk_ids,
        targets=["weaviate"],
        existing_weaviate_client=FakeWeaviateClient(),
        timeout_seconds=1,
        success_threshold=1.0,
    )
    assert ok is True
    assert calls["n"] >= 2


@pytest.mark.asyncio
async def test_readiness_probe_neo4j_transient_errors_then_success(monkeypatch):
    async def no_sleep(_):
        return None
    async def fake_to_thread(fn, *args, **kwargs):
        return fn(*args, **kwargs)
    monkeypatch.setattr(asyncio, "sleep", no_sleep)
    monkeypatch.setattr(asyncio, "to_thread", fake_to_thread)

    chunk_ids = [str(uuid.uuid4()) for _ in range(10)]
    present = set(chunk_ids)

    class FakeNeo4jClient:
        def __init__(self):
            self.calls = 0
        def query(self, query: str, tenant_id: str, params: Dict[str, Any]):
            self.calls += 1
            if self.calls < 3:
                raise Exception("transient")
            want = set(params["chunkIds"])
            return [{"count": len(want.intersection(present))}]

    async def fake_get_client():
        return FakeNeo4jClient()

    import src.autobots.ingestion.neo4j_loader as neo4j_loader_mod
    monkeypatch.setattr(neo4j_loader_mod.Neo4jManager, "get_client", staticmethod(fake_get_client))

    ok = await readiness_probe(
        collection_name="C",
        user_id="t",
        chunk_uuids=chunk_ids,
        targets=["graphrag"],
        existing_weaviate_client=None,
        timeout_seconds=1,
        success_threshold=0.75,
    )
    assert ok is True


# -------------------------------------------------------------------------------------------------
# Optional openpyxl end-to-end ingest test (skipped if missing)
# -------------------------------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_ingest_xlsx_end_to_end_without_http(monkeypatch, tmp_path: Path):
    try:
        import openpyxl  # noqa: F401
    except Exception:
        pytest.skip("openpyxl not installed")

    import openpyxl

    from src.autobots.data_model.context import Context
    from src.autobots.ingestion.spreadsheet import ingest as ingest_mod

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet;A"
    ws["A1"] = "Holdings Report"
    ws.merge_cells("A1:E1")
    ws["A2"] = "Name"
    ws["B2"] = "Ticker"
    ws["C2"] = "Qty"
    ws["D2"] = "Price"
    ws["E2"] = "Value"
    ws.append(["Apple", "AAPL", 10, 100, 1000])
    ws.append(["MSFT", "MSFT", 5, 300, 1500])

    xlsx_path = tmp_path / "test.xlsx"
    wb.save(xlsx_path)

    async def fake_download_to_tempfile(ctx, url, suffix, max_bytes, timeout_seconds=240.0):
        return xlsx_path, xlsx_path.stat().st_size

    monkeypatch.setattr(ingest_mod, "download_to_tempfile", fake_download_to_tempfile)

    async def fake_to_thread(fn, *args, **kwargs):
        return fn(*args, **kwargs)

    monkeypatch.setattr(asyncio, "to_thread", fake_to_thread)

    ctx = Context(user_id="u1", trace_id="t1")
    res = await ingest_mod.ingest_spreadsheet_from_url(
        ctx=ctx,
        file_url="https://example.invalid/test.xlsx",
        original_filename="test.xlsx",
        url_for_metadata="https://s3.example/permanent/test.xlsx",
        tenant_id="u1",
    )
    processed = res.processed_data
    assert "documents" in processed
    assert "chunk_uuids" in processed
    assert len(processed["documents"]) == len(processed["chunk_uuids"])
    assert any(d.metadata.get("doc_type") == "spreadsheet_workbook_summary" for d in processed["documents"])
    assert any(d.metadata.get("doc_type") == "spreadsheet_sheet_summary" for d in processed["documents"])
    assert any(d.metadata.get("doc_type") == "spreadsheet_rows" for d in processed["documents"])
