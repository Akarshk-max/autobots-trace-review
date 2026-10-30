import asyncio
import copy
import io
import json
import os
import re
import tempfile
import time
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Literal, Optional, Set, Tuple, Type  # Optional

import httpx
import pandas as pd
from azure.ai.inference.aio import ChatCompletionsClient
from docx import Document
from fastapi import HTTPException
from loguru import logger
from motor.motor_asyncio import AsyncIOMotorDatabase
from openai import (
    APIError,
    AsyncAzureOpenAI,
    AsyncOpenAI,
    AsyncStream,
    BadRequestError,
    RateLimitError,
    _legacy_response,
)
from openai.types import FileObject, VectorStore
from openai.types.responses import (
    FileSearchToolParam,
    Response,
    ResponseCodeInterpreterToolCall,
    ResponseCodeInterpreterToolCallParam,
    ResponseCompactionItem,
    ResponseCompactionItemParamParam,
    ResponseComputerToolCall,
    ResponseComputerToolCallParam,
    ResponseFileSearchToolCall,
    ResponseFileSearchToolCallParam,
    ResponseFunctionToolCall,
    ResponseFunctionToolCallParam,
    ResponseFunctionWebSearch,
    ResponseFunctionWebSearchParam,
    ResponseInputItemParam,
    ResponseInputTextParam,
    ResponseOutputItem,
    ResponseOutputMessage,
    ResponseOutputMessageParam,
    ResponseReasoningItem,
    ResponseReasoningItemParam,
    ResponseStreamEvent,
    ToolParam,
    response_output_item,
)
from openai.types.responses.file_search_tool_param import RankingOptions
from openai.types.responses.response_input_item_param import (
    FunctionCallOutput,
    ImageGenerationCall,
    McpApprovalRequest,
    McpCall,
    McpListTools,
)
from openai.types.responses.response_input_param import LocalShellCall, Message
from openai.types.responses.tool_param import (
    CodeInterpreter,
    CodeInterpreterContainerCodeInterpreterToolAuto,
    Mcp,
)
from pptx import Presentation
from pydantic import BaseModel, ConfigDict, Field
from tenacity import AsyncRetrying, retry_if_exception, stop_after_attempt, wait_exponential

from src.autobots.action.action.common_action_models import DataBlock
from src.autobots.conn.aws.aws_s3 import get_public_s3
from src.autobots.conn.openai.azure_openai_client_factory import AzureOpenAIModelClientFactory
from src.autobots.conn.openai.openai_common_models import (
    OPENAI_FILE_FORMATS,
    OPENAI_FILE_FORMATS_CODE_INTERPRETER,
)
from src.autobots.conn.openai.openai_files.openai_files import OpenaiFiles
from src.autobots.conn.openai.openai_files.openai_files_model import FileCreate
from src.autobots.conn.openai.openai_responses.debug_container_files import get_container_files_sync
from src.autobots.conn.openai.openai_responses.replay_sanitizer import sanitize_outgoing_payload
from src.autobots.conn.openai.openai_responses.response_model import OpenaiResponsesReq
from src.autobots.conn.openai.openai_responses.tool_args_streamer import ToolArgsStreamer
from src.autobots.conn.openai.openai_vector_stores.openai_vector_stores.openai_vector_stores import (
    OpenaiVectorStores,
)
from src.autobots.conn.openai.openai_vector_stores.openai_vector_stores.openai_vector_stores_model import (
    CreateVectorStore,
    RetrieveVectorStore,
)
from src.autobots.core.cancellation.run_instructions import RunInstruction, drain_instructions
from src.autobots.core.logging.app_code import AppCode
from src.autobots.core.settings import SettingsProvider
from src.autobots.core.utils import gen_uuid, truncate_str
from src.autobots.data_model.app_state import AppState
from src.autobots.data_model.context import Context
from src.autobots.files.file_downloader import FileDownloader
from src.autobots.files.user_files import UserFiles
from src.autobots.llm.chat_snapshot.provider_snapshots import snapshot_openai_responses
from src.autobots.llm.tools.action_chat_history_query_openai import (
    ActionChatHistoryQueryOpenai,
    ChatHistoryQueryReq,
)
from src.autobots.llm.tools.common_default_tools import OPENAI_RESPONSES_V2_DEFAULT_TOOLS
from src.autobots.llm.tools.llm_func_check_retry_urls import llm_func_check_retry_urls
from src.autobots.llm.tools.tool_factory import ToolFactory
from src.autobots.mcp.user_mcp import UserMCP
from src.autobots.server_sent_events.server_sent_events import Notify, NotifyPayload
from src.autobots.server_sent_events.topics import text_topic
from src.autobots.token_usage.token_usage_service import LlmSurface, TokenUsageService
from src.autobots.user.user_orm_model import UserORM


class ContainerResetError(Exception):
    """Raised when code interpreter container appears to have been reset."""
    pass


class AuthKillswitchError(Exception):
    """Raised when auth check fails in response loop - signals intentional run termination."""
    pass


# OpenAI 429 codes that are NOT transient — retrying won't restore access.
# See https://developers.openai.com/api/docs/guides/error-codes
BILLING_QUOTA_ERROR_CODES: Set[str] = {
    "credit_balance_exhausted",
    "organization_spend_limit_exceeded",
    "project_spend_limit_exceeded",
    "organization_usage_limit_exceeded",
}

# Phrases in error messages that indicate non-retryable billing/quota errors.
# Needed for streaming errors where the SDK falls back to APIError (no structured body / code).
BILLING_QUOTA_ERROR_PHRASES: Tuple[str, ...] = (
    "no credits remaining",
    "credit_balance_exhausted",
    "spend limit",
    "usage limit",
)


def _is_retryable_openai_responses_exception(exc: BaseException) -> bool:
    """Predicate for tenacity retry_if_exception.

    Returns False (don't retry) for:
      - Any BaseException that is not an Exception, i.e. cancellation and interpreter
        shutdown. Tenacity swallows these into the retry state, so without this a
        cancelled run would be re-issued as a fresh API call.
      - AuthKillswitchError (intentional run termination)
      - RateLimitError whose `code` is a billing/quota error (not transient)
      - APIError whose message matches a billing/quota phrase (streaming fallback, no structured code)
    Returns True (retry) for everything else, including plain rate-limit 429s.
    """
    if not isinstance(exc, Exception):
        return False
    if isinstance(exc, AuthKillswitchError):
        return False
    if isinstance(exc, RateLimitError):
        if getattr(exc, "code", None) in BILLING_QUOTA_ERROR_CODES:
            return False
    if isinstance(exc, APIError):
        msg = str(exc).lower()
        if any(phrase in msg for phrase in BILLING_QUOTA_ERROR_PHRASES):
            return False
    return True


class CompactionTooLargeError(Exception):
    """Raised when compaction encrypted_content exceeds OpenAI's limit."""
    pass


COMPACTION_ENCRYPTED_CONTENT_LIMIT = 10_485_760  # 10MB - OpenAI's max encrypted_content size


class TurnArchiveEntry(BaseModel):
    """Archive entry for a single conversation turn, preserving input/output even after/before pruning."""
    turn_number: int
    input: List[Dict[str, Any]] = Field(default_factory=list)
    output: List[Dict[str, Any]] = Field(default_factory=list)
    in_turn_messages: List[Dict[str, Any]] = Field(default_factory=list)


class VectorStoreLite(VectorStore):
    name: str | None = None


class LlmResponsesOpenaiV2Config(OpenaiResponsesReq):
    # 1. Ignore extra fields (id, summary, generator) so they don't cause validation errors
    model_config = ConfigDict(extra="ignore")
    # 2. Override 'input' to accept the List of Dicts (conversation history) that caused the crash
    # input: Union[str, List[Dict[str, Any]], Any]

    # summary: Optional[str] = None
    # id: Optional[str] = None
    # generator: Optional[Any] = None

    provider: Literal["azure", "openai"]
    tools: (
        List[ToolParam | str | Literal["file_search", "code_interpreter", "web_search"]] | None
    ) = list(OPENAI_RESPONSES_V2_DEFAULT_TOOLS)
    store: bool = True
    file_urls: List[str] = []
    file_url_to_file_object_map: Dict[str, FileObject] = {}
    file_id_to_vector_store_map: Dict[str, VectorStoreLite] = {}
    sandbox_url_to_s3_url: Dict[str, str] = {}
    tool_output_text_to_file_char_limit: int | None = 75000
    consolidated_file_ids: List[str] = Field(default_factory=list)  # original file_ids already inside a consolidation zip
    consolidation_zip_ids: List[str] = Field(default_factory=list)  # zip file_ids we created (don't re-zip these)
    compaction_triggered: bool = False
    initial_input_snapshot: List[Dict[str, Any]] = Field(default_factory=list)
    container_memory_limit: Literal["1g", "4g", "16g", "64g"] = "1g"
    create_retry_multiplier: int = 10  # Exponential backoff: 10s → 20s → 40s
    execute_retry_wait: int = 180  # Fixed wait between execute retries (seconds)
    sandbox_retry_current: int = 0  # per-conversation retry counter for sandbox-path detection
    timeout: float | None = 900  # Override parent's None default on a per request basis; user can change or set None to disable
    _keep_alive_task: asyncio.Task | None = None
    input_archive: List[TurnArchiveEntry] = Field(default_factory=list)  # Archive of all turns for logging/debugging
    visible_turns: int = Field(default=20, ge=1, le=200)  # Bounded conversation context: only last N turns sent to the model


class OpenaiResponsesV2:
    SANDBOX_PATH_PREFIX = "sandbox:/mnt/data/"
    MNT_DATA_PATH_PREFIX = "/mnt/data/"
    SANDBOX_URLS_NOT_FOUND_MESSAGE= "No valid sandbox URLs found in the assistant output."
    MAX_OPENAI_FILE_SIZE_BYTES = 510 * 1024 * 1024  # 510 MB - OpenAI limit is 512 MB
    FILE_IDS_CONSOLIDATION_THRESHOLD = 45
    MEMORY_ESCALATION = ["1g", "4g", "16g", "64g"]
    RETRY_MAX_COUNT = 4

    _keep_alive_interval: int = 300  # 5 minutes in seconds

    def __init__(
        self,
        provider: Literal["azure", "openai"],
        client_factory: Type[AzureOpenAIModelClientFactory] | None = None,  # for azure
        openai_client: AsyncOpenAI | None = None,  # for openai
    ):
        assert client_factory or openai_client, (
            "Either client_factory or openai_client must be provided"
        )
        self.provider = provider
        self.client = openai_client
        self.client_factory = client_factory

    async def _start_keep_alive_task(
        self,
        client: AsyncAzureOpenAI | ChatCompletionsClient | AsyncOpenAI,
        responses_req: LlmResponsesOpenaiV2Config,
        container_id: str,
        ctx: Context | None = None,
    ) -> None:
        """Start background keep-alive task for this container and store on responses_req."""
        task = asyncio.create_task(
            self._keep_alive_background_loop(client, container_id, ctx)
        )
        responses_req._keep_alive_task = task
        log = logger.bind(ctx=ctx) if ctx else logger
        log.info(f"Started background keep-alive task for container {container_id}")

    async def _cancel_keep_alive_task(
        self,
        responses_req: LlmResponsesOpenaiV2Config,
        ctx: Context | None = None,
    ) -> None:
        """Cancel the background keep-alive task for this chat."""
        task = getattr(responses_req, '_keep_alive_task', None)
        if task and not task.done():
            task.cancel()
            try:
                await task  # Wait for task to process the cancellation
            except asyncio.CancelledError:
                pass  # Expected - task was cancelled
            log = logger.bind(ctx=ctx) if ctx else logger
            log.info("Cancelled background keep-alive task for containers")
            responses_req._keep_alive_task = None

    async def _keep_alive_background_loop(
        self,
        client: AsyncAzureOpenAI | ChatCompletionsClient | AsyncOpenAI,
        container_id: str,
        ctx: Context | None = None,
    ) -> None:
        """Runs keep_containers_alive every 5 minutes in the background until cancelled."""
        log = logger.bind(ctx=ctx) if ctx else logger
        try:
            while True:
                await asyncio.sleep(OpenaiResponsesV2._keep_alive_interval)
                try:
                    log.info(f"Background keep-alive task running for container {container_id}")
                    await self.keep_containers_alive(client=client, container_id=container_id, ctx=ctx)
                    log.info(f"Background keep-alive task completed for container {container_id}")
                except Exception as e:
                    log.warning(f"Background keep-alive failed for container {container_id}: {e}")
        except asyncio.CancelledError:
            log.info(f"Background keep-alive task was cancelled for container {container_id}")
            raise

    def _is_container_reset(
        self,
        output_item: ResponseCodeInterpreterToolCall,
        ctx: Context | None = None,
    ) -> bool:
        """Detect if container appears to have been reset (AI running trivial test code)."""
        code = (output_item.code or "").strip()
        
        RESET_PATTERNS = [
            'print("hello")',
            "print('hello')",
            'print("hi")',
            "print('hi')",
            'print("1+1")',
            "print('1+1')",
            "1+1",
            "1 + 1",
        ]
        
        if code in RESET_PATTERNS:
            logger.bind(
                ctx=ctx,
                container_id=output_item.container_id,
                code=str(code)[:50],
            ).warning("Detected container reset - AI running trivial test code")
            return True
        
        return False

    async def _check_auth_killswitch(
        self,
        ctx: Context,
        db: AsyncIOMotorDatabase | None,
        user: UserORM | None,
        curr_loop: int,
    ) -> None:
        """Check auth at loop start - acts as killswitch if auth is revoked/removed/disconnected.
        
        If auth is no longer valid (e.g., removed from database), raises AuthKillswitchError
        which bypasses retry logic and terminates the run immediately.
        """
        # Lazy import to avoid circular dependency
        from src.autobots.integrations.app_conn.app_conn import AppConn
        
        if db is None or user is None:
            return  # Cannot check without db/user context
        
        for attempt in range(2):  # 1 retry
            conn = await AppConn.get_openai(ctx=ctx, db=db, user=user, provider=self.provider)
            if not isinstance(conn, Exception):
                return  # Auth valid, continue
            
            if attempt == 0:
                logger.bind(ctx=ctx, loop=curr_loop).warning(
                    "Auth killswitch check failed, retrying..."
                )
                await asyncio.sleep(1)
            else:
                logger.bind(ctx=ctx, loop=curr_loop).error(
                    "Auth killswitch triggered - auth connection no longer valid"
                )
                raise AuthKillswitchError("Run terminated: auth connection no longer valid")

    async def create(
        self,
        responses_req: LlmResponsesOpenaiV2Config,
        app_state: AppState,
        user: UserORM | None = None,
        ctx: Context | None = None,
        db: AsyncIOMotorDatabase | None = None,
        notify: Notify | None = None,

    ) -> Response | AsyncStream[ResponseStreamEvent] | None:
        # Strip any previous "Time of run:" marker from prior execute() retries
        if isinstance(responses_req.input, list) and len(responses_req.input) > 0:
            first_msg = responses_req.input[0]
            stripped = False
            
            if isinstance(first_msg, dict):
                content = first_msg.get("content", "")
                if isinstance(content, str) and content.startswith("Time of run:") and "\n\n" in content:
                    first_msg["content"] = content.split("\n\n", 1)[1]
                    stripped = True
            elif hasattr(first_msg, "content") and first_msg.content:
                if isinstance(first_msg.content, str) and first_msg.content.startswith("Time of run:") and "\n\n" in first_msg.content:
                    first_msg.content = first_msg.content.split("\n\n", 1)[1]
                    stripped = True
            
            if stripped:
                logger.bind(ctx=ctx).info("Stripped previous 'Time of run:' marker from input (execute retry)")
        
        # Snapshot input ONCE for all retries
        input_snapshot = copy.deepcopy(responses_req.input)
        
        async for attempt in AsyncRetrying(
            stop=stop_after_attempt(5),
            wait=wait_exponential(multiplier=responses_req.create_retry_multiplier, min=10, max=300),
            retry=retry_if_exception(_is_retryable_openai_responses_exception),
            reraise=True
        ):
            with attempt:
                try:
                    logger.bind(ctx=ctx).trace("Starting OpenAI Responses")
                    client: AsyncAzureOpenAI | ChatCompletionsClient | AsyncOpenAI | None = None
                    if self.provider == "azure":
                        logger.bind(ctx=ctx).info(
                            f"Using Azure OpenAI client with model: {responses_req.model}"
                        )
                        client: (
                            AsyncAzureOpenAI | ChatCompletionsClient
                        ) = await self.client_factory.create_client(model=responses_req.model)
                    elif self.provider == "openai":
                        logger.bind(ctx=ctx).info("Using OpenAI responses client")
                        client: AsyncOpenAI = self.client
                    assert client, "Openai LLM responses Client not found"

                    user_files: UserFiles | None = None
                    if user:
                        s3 = get_public_s3()
                        user_files = UserFiles(user=user, s3=s3)
                    res = None
                    if responses_req.stream:
                        res = await self.response_loop_stream(
                            ctx=ctx,
                            app_state=app_state,
                            client=client,
                            responses_req=responses_req,
                            user=user,
                            db=db,
                            user_files=user_files,
                            notify=notify,

                        )
                    else:
                        res = await self.response_loop(
                            ctx=ctx,
                            app_state=app_state,
                            client=client,
                            responses_req=responses_req,
                            user=user,
                            db=db,
                            user_files=user_files,
                            notify=notify,

                        )
                    logger.bind(
                        ctx=ctx,
                        res_type=type(res).__name__,
                        res_is_none=res is None,
                        res_output_text=res.output_text[:100] if res and hasattr(res, 'output_text') and res.output_text else None,
                    ).info("response_loop returned")
                    logger.trace("Completed OpenAI Responses")
                    if isinstance(res, AsyncStream):
                        self.yield_responses_chunks(res)
                    else:
                        return res
                except Exception as e:
                    # Check for container crash/reset error and escalate memory if possible
                    container_error = (
                        (isinstance(e, BadRequestError) and "container is not running" in str(e).lower())
                        or isinstance(e, ContainerResetError)
                        or "container" in str(e).lower()
                    )
                    if container_error:
                        current_limit = responses_req.container_memory_limit or "1g"
                        try:
                            current_idx = self.MEMORY_ESCALATION.index(current_limit)
                            if current_idx < len(self.MEMORY_ESCALATION) - 1:
                                next_limit = self.MEMORY_ESCALATION[current_idx + 1]
                                logger.bind(
                                    ctx=ctx,
                                    error_type=type(e).__name__,
                                    current_memory=current_limit,
                                    next_memory=next_limit,
                                    escalation_step=current_idx + 1,
                                    max_escalations=len(self.MEMORY_ESCALATION) - 1,
                                ).warning(
                                    "Container crashed/reset, escalating memory and retrying"
                                )
                                responses_req.input = copy.deepcopy(input_snapshot)
                                responses_req.container_memory_limit = next_limit
                            else:
                                logger.bind(
                                    ctx=ctx,
                                    error_type=type(e).__name__,
                                    memory_limit=current_limit,
                                ).error(
                                    "Container crashed/reset at max memory (64g), cannot escalate further"
                                )
                                responses_req.input = copy.deepcopy(input_snapshot)
                        except ValueError:
                            pass  # current_limit not in MEMORY_ESCALATION list, skip escalation

                    # On stream connection drop, fall back to non-stream for retry
                    if responses_req.stream and "peer closed connection" in str(e).lower():
                        responses_req.stream = False
                        logger.bind(ctx=ctx, error_type=type(e).__name__).warning(
                            "Stream connection dropped (RemoteProtocolError), disabling stream for retry"
                        )

                    if not isinstance(e, RecursionError):
                        responses_req.input = copy.deepcopy(input_snapshot)
                        # Clear in_turn_messages for current turn on retry
                        if responses_req.input_archive:
                            responses_req.input_archive[-1].in_turn_messages = []
                        
                        # Prepend timestamp to bust prompt cache
                        attempt_text = None
                        if isinstance(responses_req.input, list) and len(responses_req.input) > 0:
                            first_msg = responses_req.input[0]
                            attempt_text = f"Time of run: {datetime.now().isoformat()}\n\n"
                            
                            if isinstance(first_msg, dict):
                                content = first_msg.get("content", "")
                                if isinstance(content, str):
                                    first_msg["content"] = attempt_text + content
                            elif hasattr(first_msg, "content"):
                                if isinstance(first_msg.content, str):
                                    first_msg.content = attempt_text + first_msg.content
                        
                        logger.bind(
                            ctx=ctx,
                            error_type=type(e).__name__,
                            attempt_text=attempt_text.strip() if attempt_text else None,
                        ).info("Resetting input for retry except in the case of RecursionError")

                    logger.bind(
                        ctx=ctx, responses_req_model=responses_req.model
                    ).exception(f"Error in Openai Responses: {str(e)}")
                    raise
                finally:
                    # Ensure cleanup of keep-alive task
                    await self._cancel_keep_alive_task(responses_req, ctx)

    @staticmethod
    def _deliver_instructions(
        ctx: Context,
        responses_req: Any,
        instructions: list[RunInstruction],
        notify: Notify | None,
    ) -> None:
        """Append run instructions as user input items, mirrored into the turn archive."""
        for instruction in instructions:
            item = {"role": "user", "content": instruction.as_user_message(), "type": "message"}
            responses_req.input.append(item)
            if responses_req.input_archive:
                responses_req.input_archive[-1].in_turn_messages.append(dict(item))
            logger.bind(ctx=ctx).info(f"Run instruction applied: {instruction.text}")
            notify(ctx, f"Instruction applied: {instruction.text}") if notify else None


    async def _instructions_at_exit(
        self, ctx: Context, app_state: AppState, responses_req: Any, notify: Notify | None
    ) -> bool:
        pending = await drain_instructions(app_state)
        if not pending:
            return False
        self._deliver_instructions(ctx, responses_req, pending, notify)
        return True

    def _get_max_loop(self, responses_req: LlmResponsesOpenaiV2Config) -> int:
        if not responses_req.tools:
            return 20
        for tool in responses_req.tools:
            if isinstance(tool, str) and tool == "run_terminal_command":
                return 50
            if isinstance(tool, dict) and tool.get("name") == "run_terminal_command":
                return 50
            if hasattr(tool, "name") and tool.name == "run_terminal_command":
                return 50
        return 50

    async def response_loop(
        self,
        ctx: Context,
        app_state: AppState,
        client: AsyncAzureOpenAI | ChatCompletionsClient | AsyncOpenAI,
        responses_req: LlmResponsesOpenaiV2Config,
        user: UserORM | None = None,
        db: AsyncIOMotorDatabase | None = None,
        user_files: UserFiles | None = None,
        notify: Notify | None = None,

    ) -> Response | AsyncStream[ResponseStreamEvent]:
        # is_continue: bool = True
        # max_loop = 20
        max_loop = self._get_max_loop(responses_req)
        logger.bind(ctx=ctx).info(f"Response max loop is {max_loop}")
        curr_loop = 0
        while True:
            app_state.raise_if_cancelled()
            snapshot_openai_responses(ctx, app_state, responses_req, curr_loop)
            self._deliver_instructions(ctx, responses_req, await drain_instructions(app_state), notify)
            # Killswitch: check if auth is still valid (allows external run cancellation)
            await self._check_auth_killswitch(ctx=ctx, db=db, user=user, curr_loop=curr_loop)
            
            notify(ctx, "OpenAI thinking...") if notify else None
            if curr_loop >= max_loop:
                logger.bind(ctx=ctx).error(f"Responses loop exceeded max loop: {max_loop}")
                raise RecursionError({"detail": f"Responses loop exceeded max loop: {max_loop}"})
            curr_loop += 1
            logger.bind(ctx=ctx).info(f"Responses Loop: {curr_loop}")

            # responses_config = await OpenaiResponses.get_model_based_config(ctx=ctx, responses_req=responses_req)

            response: Response | AsyncStream[ResponseStreamEvent] | None = None
            # response_error: Exception | None = None
            payload = responses_req.model_dump(exclude_none=True)
            openai_responses_req = OpenaiResponsesReq(**payload)
            final_payload = openai_responses_req.model_dump(exclude_none=True)

            # Non-reasoning models: gpt-3.5* and gpt-4* do not support reasoning
            model = final_payload.get("model", "")
            if model.startswith("gpt-3.5") or model.startswith("gpt-4"):
                logger.bind(ctx=ctx, model=model).info(
                    "Removing reasoning for non-reasoning model in openai responses file"
                )
                final_payload.pop("reasoning", None)

            logger.bind(ctx=ctx).info("Starting replay sanitizer process")
            # Replay hygiene: strip/repair fields in the outgoing *wire payload* only.
            # We keep full-fidelity annotations in stored outputs/history for UI/audit.
            final_payload, replay_sanitization_report = sanitize_outgoing_payload(final_payload)
            if replay_sanitization_report.cleared_lists:
                logger.bind(ctx=ctx, stats=replay_sanitization_report.to_dict()).warning(
                    "Replay sanitizer cleared annotations in outgoing payload"
                )
            elif replay_sanitization_report.repaired or replay_sanitization_report.dropped:
                logger.bind(ctx=ctx, stats=replay_sanitization_report.to_dict()).debug(
                    "Replay sanitizer adjusted outgoing payload"
                )

            logger.bind(ctx=ctx).info("Completed replay sanitizer process")
            
            try:
                response: (
                    Response | AsyncStream[ResponseStreamEvent]
                ) = await client.responses.create(
                    **final_payload
                )
                logger.bind(
                    ctx=ctx,
                    responses_req_model=responses_req.model,
                    response=truncate_str(response.output_text, limit=100),
                ).info("Openai Responses response")
            except Exception as e:
                logger.bind(
                    ctx=ctx,
                ).exception(
                    f"Error in Openai Responses: {str(e)}, continuing...",
                )
                raise

            # response_2_error: Exception | None = None
            if response is None:
                raise Exception("Response is None")

            TokenUsageService.record(
                ctx=ctx,
                surface=LlmSurface.openai_responses,
                app_state=app_state,
                usage=getattr(response, "usage", None),
                model=str(getattr(response, "model", None) or responses_req.model),
                loop=curr_loop,
            )
            #     # Note: Removing the last reasoning item (which precedes output item) to not get 400 error sayiing reasoning item is unpaired
            #     logger.bind(ctx=ctx).info(
            #         "Removing the last reasoning item (which precedes output item) to not get 400 error saying reasoning item is unpaired"
            #     )
            #     responses_req.input.pop(-3)
            #     openai_responses_req = OpenaiResponsesReq(
            #         **responses_req.model_dump(exclude_none=True)
            #     )
            #     try:
            #         response: (
            #             Response | AsyncStream[ResponseStreamEvent]
            #         ) = await client.responses.create(
            #             **openai_responses_req.model_dump(exclude_none=True)
            #         )
            #         logger.bind(
            #             ctx=ctx, responses_req=responses_req, response=response.output_text
            #         ).info("Openai Responses response")
            #     except Exception as e:
            #         logger.bind(
            #             ctx=ctx,
            #         ).error(f"Error in Openai Responses: {str(e)}")
            #         response_2_error = e

            # if response_error and response_2_error:
            #     logger.bind(
            #         ctx=ctx,
            #         responses_req=responses_req,
            #         response_error=response_error,
            #         response_2_error=response_2_error,
            #     ).error(f"Raising error: {str(response_error)}")
            #     raise response_error

            match response:
                # ERROR: The error non-empty 'return' inside asynchronous generator means that
                # you are trying to return response from within an async def function
                # that is also a generator (i.e. it uses yield).
                # case AsyncStream():
                #     yield response
                case Response():
                    response.output

            logger.bind(
                ctx=ctx,
                # responses_req=responses_req,
                # response=response.output_text
            ).info("Openai Responses response")

            is_next_loop_needed: bool = False
            response_outputs: List[ResponseOutputItem] = response.output
            for response_output in response_outputs:
                logger.bind(ctx=ctx,  #response_output=response_output.model_dump()
                            ).info(
                    "Openai Response Output"
                )

                if await self.is_next_loop_needed(response_output):
                    is_next_loop_needed = True

                # Run any required side-effects for this output item first (notably: function tool calls),
                # so that when we store the output item into `responses_req.input`, it reflects any post-processing
                # (e.g., tool arguments linkified after execution).
                response_input_item_param, override_arguments = await self.process_response_output_item(
                    ctx=ctx,
                    app_state=app_state,
                    responses_req=responses_req,
                    response_output_item=response_output,
                    client=client,
                    user=user,
                    db=db,
                    notify=notify,

                )

                responses_output_to_input = await self.convert_output_item_to_input_item_param(
                    ctx=ctx,
                    responses_req=responses_req,
                    output_item=response_output,
                    store=responses_req.store,
                    client=client,
                    user=user,
                    db=db,
                    user_files=user_files,
                    override_arguments=override_arguments,
                )
                if isinstance(responses_output_to_input, tuple):
                    output_param, retry_msg = responses_output_to_input
                    responses_req.input.append(output_param)
                    responses_req.input.append(retry_msg)
                    is_next_loop_needed = True
                elif responses_output_to_input:
                    responses_req.input.append(responses_output_to_input)
                    # Archive in_turn_messages (exclude output messages)
                    if not isinstance(response_output, ResponseOutputMessage) and responses_req.input_archive:
                        responses_req.input_archive[-1].in_turn_messages.append(
                            responses_output_to_input.model_dump() if hasattr(responses_output_to_input, 'model_dump') else responses_output_to_input
                        )

                if response_input_item_param:
                    responses_req.input.append(response_input_item_param)
                    # Archive tool results to in_turn_messages
                    if responses_req.input_archive:
                        responses_req.input_archive[-1].in_turn_messages.append(
                            response_input_item_param.model_dump() if hasattr(response_input_item_param, 'model_dump') else response_input_item_param
                        )

            if not is_next_loop_needed:
                if await self._instructions_at_exit(ctx, app_state, responses_req, notify):
                    continue
                return response

    async def yield_responses_chunks(self, responses_res: AsyncStream[ResponseStreamEvent]):
        try:
            async for part in responses_res:
                yield part
        except Exception as e:
            logger.error(str(e))

    async def response_loop_stream(
        self,
        ctx: Context,
        app_state: AppState,
        client: AsyncAzureOpenAI | ChatCompletionsClient | AsyncOpenAI,
        responses_req: LlmResponsesOpenaiV2Config,
        user: UserORM | None = None,
        db: AsyncIOMotorDatabase | None = None,
        user_files: UserFiles | None = None,
        notify: Notify | None = None,

    ) -> Response:
        # max_loop = 20
        max_loop = self._get_max_loop(responses_req)
        curr_loop = 0
        final_response = None
        # def send(ev: str, **kw) -> bytes:
        #     logger.bind(
        #         ctx=ctx, responses_req=responses_req.model_dump(exclude_none=True)
        #     ).info(f"Sending event: {ev} with data: {kw}")
        #     return (json.dumps({"event": ev, **kw}, ensure_ascii=False) + "\n").encode("utf-8")

        content_main = ""
        content_complimentary = ""
        tool_args_streamers: Dict[str, ToolArgsStreamer] = {}
        code_interpreter_code_preview: Dict[str, str] = {}
        code_interpreter_code_pending: Dict[str, int] = {}
        code_interpreter_code_total: Dict[str, int] = {}

        def _cleanup_tool_args_streamer(streamer: ToolArgsStreamer) -> None:
            keys = [k for k, v in tool_args_streamers.items() if v is streamer]
            for k in keys:
                tool_args_streamers.pop(k, None)


        notify_main = (
            Notify(topic=text_topic(notify.topic), payload=NotifyPayload.snapshot)
            if notify
            else None
        )
        notify_complimentary = Notify(topic=notify.topic) if notify else None
        while True:
            app_state.raise_if_cancelled()
            snapshot_openai_responses(ctx, app_state, responses_req, curr_loop)
            self._deliver_instructions(ctx, responses_req, await drain_instructions(app_state), notify)
            # Killswitch: check if auth is still valid (allows external run cancellation)
            await self._check_auth_killswitch(ctx=ctx, db=db, user=user, curr_loop=curr_loop)
            
            if curr_loop >= max_loop:
                logger.bind(ctx=ctx).error(f"Responses loop exceeded max loop: {max_loop}")
                raise RecursionError({"detail": f"Responses loop exceeded max loop: {max_loop}"})
            curr_loop += 1
            logger.bind(ctx=ctx).info(f"Responses Loop: {curr_loop}")

            # responses_config = await OpenaiResponses.get_model_based_config(ctx=ctx, responses_req=responses_req)

            response: Response | None = None
            # response_error: Exception | None = None
            if not responses_req.stream:
                logger.bind(ctx=ctx).info("Setting stream=True for responses_req")
                responses_req.stream = True
            payload = responses_req.model_dump(exclude_none=True)
            openai_responses_req = OpenaiResponsesReq(**payload)
            final_payload = openai_responses_req.model_dump(exclude_none=True)

            # Non-reasoning models: gpt-3.5* and gpt-4* do not support reasoning
            model = final_payload.get("model", "")
            if model.startswith("gpt-3.5") or model.startswith("gpt-4"):
                logger.bind(ctx=ctx, model=model).info(
                    "Removing reasoning for non-reasoning model in openai responses file"
                )
                final_payload.pop("reasoning", None)

            logger.bind(ctx=ctx).info("Starting replay sanitizer process")
            # Replay hygiene: strip/repair fields in the outgoing *wire payload* only.
            # We keep full-fidelity annotations in stored outputs/history for UI/audit.
            final_payload, replay_sanitization_report = sanitize_outgoing_payload(final_payload)
            if replay_sanitization_report.cleared_lists:
                logger.bind(ctx=ctx, stats=replay_sanitization_report.to_dict()).warning(
                    "Replay sanitizer cleared annotations in outgoing payload (stream)"
                )
            elif replay_sanitization_report.repaired or replay_sanitization_report.dropped:
                logger.bind(ctx=ctx, stats=replay_sanitization_report.to_dict()).debug(
                    "Replay sanitizer adjusted outgoing payload (stream)"
                )
            
            logger.bind(ctx=ctx).info("Completed replay sanitizer process")
            
            try:
                # print("responses_req: ", responses_req)
                logger.bind(
                    ctx=ctx,  #responses_req=responses_req.model_dump(exclude_none=True)
                ).info(
                    "Creating OpenAI responses stream object"
                )
                notify_complimentary(ctx, "Calling OpenAI stream API...")
                stream_obj: AsyncStream[ResponseStreamEvent] = await client.responses.create(
                    **final_payload
                )
                # print("stream_obj: ", stream_obj)
                logger.bind(
                    ctx=ctx,  #responses_req=responses_req.model_dump(exclude_none=True)
                ).info("Created OpenAI responses stream object")
                loop_needed: bool = False
                notify_complimentary(ctx, "OpenAI thinking...")
                try:
                    # iterate the stream as you already do
                    async for event in stream_obj:  # type: ignore
                        # print("event: ", event)
                        # logger.bind(
                        #     ctx=ctx,  #responses_req=responses_req.model_dump(exclude_none=True)
                        # ).info(f"Received event: {event}")
                        match event.type:
                            # text & refusal deltas
                            case "response.created":
                                # yield send("token_2", message="Starting to create your response...")
                                content_complimentary = "Starting to create your response..."
                                notify_complimentary(ctx, content_complimentary)
                            case "response.in_progress":
                                # yield send("token_2", message="Response in progress...")
                                content_complimentary = "Response in progress..."
                                notify_complimentary(ctx, content_complimentary)
                            case "response.output_text.delta":
                                # yield send("token", text=event.delta)
                                content_main += event.delta
                                notify_main(ctx, content_main)
                                content_complimentary = "Generating message response..."
                                notify_complimentary(ctx, content_complimentary)
                            case "response.refusal.delta":
                                content_complimentary += event.delta
                                notify_complimentary(ctx, content_complimentary)
                            case "response.content_part.added":
                                if event.part.type == "refusal":
                                    content_complimentary = (
                                        f"Unfortunate Refusal from the model : {event.part.refusal}"
                                    )
                                    notify_complimentary(ctx, content_complimentary)
                                if event.part.type == "reasoning_text":
                                    content_complimentary = f"Model's Reasoning : {event.part.text}"
                                    notify_complimentary(ctx, content_complimentary)
                            case "response.content_part.done":
                                if event.part.type == "refusal":
                                    content_complimentary = (
                                        f"Refusal finalized : {event.part.refusal}"
                                    )
                                    notify_complimentary(ctx, content_complimentary)
                                if event.part.type == "reasoning_text":
                                    content_complimentary = f"Model reasoning finalized : {event.part.text}"
                                    notify_complimentary(ctx, content_complimentary)
                            case "response.reasoning_text.delta":
                                content_complimentary += event.delta
                                notify_complimentary(ctx, content_complimentary)
                            # case "response.output_text.done":
                            #     # yield send("token", text=event.text)
                            #     content_main += event.text
                            #     notify_main(ctx,content_main)
                            # case "respons"

                            # case "response.refusal.delta":
                            #     # yield send("token_2", text=event.delta)

                            #     notify_complimentary(ctx,event.delta)
                            # code interpreter lifecycle (bounded preview; no unbounded accumulation)
                            case "response.code_interpreter_call.in_progress":
                                item_id = getattr(event, "item_id", "")
                                if item_id:
                                    code_interpreter_code_preview[item_id] = ""
                                    code_interpreter_code_pending[item_id] = 0
                                    code_interpreter_code_total[item_id] = 0
                                content_complimentary = (
                                    "Code interpreter: generating code\n"
                                    "Preview (first 400 chars):\n"
                                    "(streaming...)"
                                )
                                notify_complimentary(ctx, content_complimentary)
                            case "response.code_interpreter_call.interpreting":
                                content_complimentary = "Code interpreter interpreting..."
                                notify_complimentary(ctx, content_complimentary)
                            case "response.code_interpreter_call.completed":
                                content_complimentary = "Code interpreter completed."
                                notify_complimentary(ctx, content_complimentary)
                                item_id = getattr(event, "item_id", "")
                                if item_id:
                                    code_interpreter_code_preview.pop(item_id, None)
                                    code_interpreter_code_pending.pop(item_id, None)
                                    code_interpreter_code_total.pop(item_id, None)
                            case "response.code_interpreter_call_code.delta":
                                item_id = getattr(event, "item_id", "")
                                delta = getattr(event, "delta", "") or ""
                                # Maintain only a small HEAD preview for UI.
                                preview = code_interpreter_code_preview.get(item_id, "")
                                total = code_interpreter_code_total.get(item_id, 0) + len(delta)
                                pending = code_interpreter_code_pending.get(item_id, 0) + len(delta)

                                updated_preview = False
                                if len(preview) < 400 and delta:
                                    needed = 400 - len(preview)
                                    if needed > 0:
                                        preview += delta[:needed]
                                        updated_preview = True

                                if item_id:
                                    code_interpreter_code_preview[item_id] = preview
                                    code_interpreter_code_total[item_id] = total
                                    # Only track pending while preview is still growing; avoids runaway.
                                    code_interpreter_code_pending[item_id] = pending if len(preview) < 400 else 0

                                # Throttle updates to keep notify traffic reasonable.
                                if updated_preview and (pending >= 400 or "\n" in delta or len(preview) == 400):
                                    content_complimentary = (
                                        f"Code interpreter: generating code ({total} chars)\n"
                                        "Preview (first 400 chars):\n"
                                        + preview
                                        + ("\n... (truncated)" if total > 400 else "")
                                    )
                                    notify_complimentary(ctx, content_complimentary)
                                    if item_id:
                                        code_interpreter_code_pending[item_id] = 0
                            case "response.code_interpreter_call_code.done":
                                item_id = getattr(event, "item_id", "")
                                code = getattr(event, "code", "") or ""
                                total = len(code)
                                preview = code[:400]
                                if item_id:
                                    code_interpreter_code_preview[item_id] = preview
                                    code_interpreter_code_total[item_id] = total
                                    code_interpreter_code_pending[item_id] = 0
                                content_complimentary = (
                                    f"Code interpreter: code generation done ({total} chars)\n"
                                    "Preview (first 400 chars):\n"
                                    + preview
                                    + ("\n... (truncated)" if total > 400 else "")
                                )
                                notify_complimentary(ctx, content_complimentary)
                                # if item_id:
                                #     code_interpreter_code_preview.pop(item_id, None)
                                #     code_interpreter_code_pending.pop(item_id, None)
                                #     code_interpreter_code_total.pop(item_id, None)
                            # file search lifecycle
                            case "response.file_search_call.in_progress":
                                # yield send("token_2", text="file search in progress")
                                content_complimentary = "File search in progress"
                                notify_complimentary(ctx, content_complimentary)
                            case "response.file_search_call.searching":
                                # yield send("token_2", text="file search searching")
                                content_complimentary = "File search searching"
                                notify_complimentary(ctx, content_complimentary)
                            case "response.file_search_call.completed":
                                # yield send("token_2", text="file search completed")
                                content_complimentary = "File search completed"
                                notify_complimentary(ctx, content_complimentary)

                            case "response.web_search_call.in_progress":
                                # yield send("token_2", text="web search in progress")
                                content_complimentary = "Web search in progress"
                                notify_complimentary(ctx, content_complimentary)
                            case "response.web_search_call.searching":
                                # yield send("token_2", text="web search searching")
                                content_complimentary = "Web search searching"
                                notify_complimentary(ctx, content_complimentary)
                            case "response.web_search_call.completed":
                                # yield send("token_2", text="web search completed")
                                content_complimentary = "Web search completed"
                                notify_complimentary(ctx, content_complimentary)

                            # reasoning tokens (o-series)
                            case "response.reasoning_summary_part.added":
                                # yield send("reasoning-summary", text="Reasoning summary being added...")
                                content_complimentary = "Reasoning summary being added..."
                                notify_complimentary(ctx, content_complimentary)
                            case "response.reasoning_summary_part.done":
                                # yield send("reasoning-summary", text=event.part.text)
                                content_complimentary += event.part.text
                                notify_complimentary(ctx, content_complimentary)
                            case "response.reasoning_summary_text.delta":
                                # yield send("reasoning-summary", text=event.delta)
                                content_complimentary += event.delta
                                notify_complimentary(ctx, content_complimentary)
                            # case "response.reasoning_text.delta":
                            #     yield send("reasoning", text=f"{event.delta}")

                            # image generation lifecycle (names per docs; adjust if your tool emits slightly different ones)
                            case "response.image_generation_call.in_progress":
                                # yield send("token_2", text="image generation in progress...")
                                content_complimentary = "Image generation in progress..."
                                notify_complimentary(ctx, content_complimentary)
                            case "response.image_generation_call.completed":
                                # yield send("token_2", text="image generation completed")
                                content_complimentary = "Image generation completed"
                                notify_complimentary(ctx, content_complimentary)
                            case "response.image_generation_call.generating":
                                # yield send("token_2", text="image generation generating")
                                content_complimentary = "Image generation generating"
                                notify_complimentary(ctx, content_complimentary)
                            case "response.image_generation_call.partial_image":
                                # yield send("partial-image", text=event.partial_image_b64)
                                notify_complimentary(ctx, event.partial_image_b64)
                            # output item add: function tool call starting
                            case "response.mcp_call.completed":
                                content_complimentary = "MCP Call completed..."
                                notify_complimentary(ctx, content_complimentary)
                            case "response.mcp_call.failed":
                                content_complimentary = "MCP Call failed..."
                                notify_complimentary(ctx, content_complimentary)
                            case "response.mcp_call.in_progress":
                                content_complimentary = "MCP Call in progress..."
                                notify_complimentary(ctx, content_complimentary)
                            case "response.output_item.added":
                                # your body expects `fc.name` already bound
                                if event.item.type == "function_call":
                                    # Start a structured tool-arguments stream. Avoid emitting raw JSON.
                                    item_id = getattr(event.item, "id", None)
                                    call_id = getattr(event.item, "call_id", None)
                                    tool_name = getattr(event.item, "name", None) or "tool"
                                    streamer = ToolArgsStreamer(
                                        tool_name, key_char_limit=100, value_char_limit=100
                                    )
                                    # Store under both IDs to match subsequent delta events (which use item_id).
                                    if item_id:
                                        tool_args_streamers[item_id] = streamer
                                    if call_id:
                                        tool_args_streamers[call_id] = streamer
                                    if not item_id and not call_id:
                                        tool_args_streamers[str(gen_uuid())] = streamer
                                    content_complimentary = streamer.render()
                                    # logger.bind(ctx=ctx).info(
                                    #     f"(Notify) Started tool arguments streamer for function call -- content complimentary : {content_complimentary}")
                                    notify_complimentary(ctx, content_complimentary)
                                    streamer.mark_clean()
                                if event.item.type == "message":
                                    content_complimentary = "Generating message response..."
                                    notify_complimentary(ctx, content_complimentary)
                                if event.item.type == "mcp_call":
                                    content_complimentary = f"MCP Call initiated for server : {event.item.server_label}..."
                                    notify_complimentary(ctx, content_complimentary)
                                if event.item.type == "mcp_list_tools":
                                    content_complimentary = "Listing MCP Tools on the server..."
                                    notify_complimentary(ctx, content_complimentary)
                            case "response.output_item.done":
                                if event.item.type == "function_call":
                                    content_complimentary = f"Function call completed: {event.item.name}"
                                    notify_complimentary(ctx, content_complimentary)

                                    # Cleanup tool args streamer bookkeeping (item_id/call_id aliases).
                                    item_id = getattr(event.item, "id", None)
                                    call_id = getattr(event.item, "call_id", None)
                                    streamer = None
                                    if item_id:
                                        streamer = tool_args_streamers.get(item_id) or streamer
                                    if call_id:
                                        streamer = tool_args_streamers.get(call_id) or streamer
                                    if streamer is not None:
                                        _cleanup_tool_args_streamer(streamer)
                                elif event.item.type == "message":
                                    content_complimentary = "Message generation completed."
                                    notify_complimentary(ctx, content_complimentary)
                                elif event.item.type == "file_search_call":
                                    content_complimentary = "File search item completed."
                                    notify_complimentary(ctx, content_complimentary)
                                elif event.item.type == "web_search_call":
                                    content_complimentary = "Web search item completed."
                                    notify_complimentary(ctx, content_complimentary)
                                elif event.item.type == "code_interpreter_call":
                                    content_complimentary = "Code interpreter item completed."
                                    notify_complimentary(ctx, content_complimentary)
                                elif event.item.type == "mcp_call":
                                    content_complimentary = "MCP call item completed."
                                    notify_complimentary(ctx, content_complimentary)
                                elif event.item.type == "reasoning":
                                    content_complimentary = "Reasoning completed."
                                    notify_complimentary(ctx, content_complimentary)
                                elif event.item.type == "compaction":
                                    content_complimentary = "Context compaction completed."
                                    notify_complimentary(ctx, content_complimentary)
                                else:
                                    logger.bind(ctx=ctx).warning(f"Unhandled output_item.done type: {event.item.type}")

                            # function call arguments
                            case "response.function_call_arguments.delta":
                                streamer = tool_args_streamers.get(getattr(event, "item_id", ""))
                                if streamer is None:
                                    # Fallback: if we didn't see output_item.added, still stream structurally.
                                    # logger.bind(ctx=ctx).info(
                                    #     f"Received function_call_arguments.delta for item_id {getattr(event, 'item_id', '')} without existing streamer, creating fallback streamer"
                                    # )
                                    streamer = ToolArgsStreamer("tool", key_char_limit=100, value_char_limit=100)
                                    tool_args_streamers[getattr(event, "item_id", "")] = streamer
                                streamer.ingest(getattr(event, "delta", ""))
                                if streamer.has_updates():
                                    content_complimentary = streamer.render()
                                    # logger.bind(ctx=ctx).info(
                                    #     f"(Notify) Updated tool arguments streamer for function call -- content complimentary : {content_complimentary}")
                                    notify_complimentary(ctx, content_complimentary)
                                    streamer.mark_clean()
                            case "response.function_call_arguments.done":
                                streamer = tool_args_streamers.get(getattr(event, "item_id", ""))
                                if streamer is None:
                                    streamer = ToolArgsStreamer(
                                        getattr(event, "name", None) or "tool",
                                        key_char_limit=100,
                                        value_char_limit=100,
                                    )
                                    tool_args_streamers[getattr(event, "item_id", "")] = streamer
                                # Final pass from full JSON to fill any gaps.
                                full = getattr(event, "arguments", None)
                                if full:
                                    streamer.finalize_from_full(full)
                                content_complimentary = streamer.render()
                                # logger.bind(ctx=ctx).info(
                                #     f"(Notify) Finalized tool arguments streamer for function call -- content complimentary : {content_complimentary}")
                                notify_complimentary(ctx, content_complimentary)
                                streamer.mark_clean()
                                _cleanup_tool_args_streamer(streamer)

                            # output text finalization
                            case "response.output_text.done":
                                pass  # Text already accumulated via deltas
                            case "response.output_text.annotation.added":
                                content_complimentary = "Annotation added."
                                notify_complimentary(ctx, content_complimentary)

                            # keepalive heartbeat
                            case "keepalive":
                                content_complimentary = "Keeping connection alive..."
                                notify_complimentary(ctx, content_complimentary)

                            # completed: final response object
                            case "response.completed":
                                # print("response completed")
                                response = event.response
                                is_next_loop_needed = False
                                final_response = response
                                TokenUsageService.record(
                                    ctx=ctx,
                                    surface=LlmSurface.openai_responses,
                                    app_state=app_state,
                                    usage=getattr(response, "usage", None),
                                    model=str(getattr(response, "model", None) or responses_req.model),
                                    loop=curr_loop,
                                )
                                logger.bind(
                                    ctx=ctx,
                                    # event=event,
                                    # response=response.model_dump(exclude_none=True),
                                ).info("Response completed")
                                # content_complimentary = "OpenAI Response Completed!"
                                # notify_complimentary(ctx, content_complimentary)
                                response_outputs: List[ResponseOutputItem] = response.output
                                for response_output in response_outputs:
                                    logger.bind(
                                        ctx=ctx,
                                    ).info("Openai Response Output")

                                    if await self.is_next_loop_needed(response_output):
                                        is_next_loop_needed = True

                                    if is_next_loop_needed:
                                        content_complimentary = "Response completed for this loop. Proceeding to next loop."
                                    else:
                                        content_complimentary = "OpenAI Response Completed! Finishing up..."
                                    notify_complimentary(ctx, content_complimentary)

                                    # Run any required side-effects for this output item first (notably: function tool calls),
                                    # so that when we store the output item into `responses_req.input`, it reflects any post-processing
                                    # (e.g., tool arguments linkified after execution).
                                    response_input_item_param, override_arguments = await self.process_response_output_item(
                                        ctx=ctx,
                                        app_state=app_state,
                                        responses_req=responses_req,
                                        client=client,
                                        response_output_item=response_output,
                                        user=user,
                                        db=db,
                                        notify=notify_complimentary,
            
                                    )

                                    responses_output_to_input = await self.convert_output_item_to_input_item_param(
                                        ctx=ctx,
                                        responses_req=responses_req,
                                        output_item=response_output,
                                        store=responses_req.store,
                                        client=client,
                                        user_files=user_files,
                                        user=user,
                                        db=db,
                                        override_arguments=override_arguments,
                                    )
                                    if isinstance(responses_output_to_input, tuple):
                                        output_param, retry_msg = responses_output_to_input
                                        responses_req.input.append(output_param)
                                        responses_req.input.append(retry_msg)
                                        is_next_loop_needed = True
                                    elif responses_output_to_input:
                                        responses_req.input.append(responses_output_to_input)
                                        # Archive in_turn_messages (exclude output messages)
                                        if not isinstance(response_output, ResponseOutputMessage) and responses_req.input_archive:
                                            responses_req.input_archive[-1].in_turn_messages.append(
                                                responses_output_to_input.model_dump() if hasattr(responses_output_to_input, 'model_dump') else responses_output_to_input
                                            )

                                    if is_next_loop_needed:
                                        content_complimentary = "Response completed for this loop. Proceeding to next loop."
                                    else:
                                        content_complimentary = "OpenAI Response Completed! Finishing up..."
                                    notify_complimentary(ctx, content_complimentary)

                                    if response_input_item_param:
                                        responses_req.input.append(response_input_item_param)
                                        # Archive tool results to in_turn_messages
                                        if responses_req.input_archive:
                                            responses_req.input_archive[-1].in_turn_messages.append(
                                                response_input_item_param.model_dump() if hasattr(response_input_item_param, 'model_dump') else response_input_item_param
                                            )

                                    loop_needed = is_next_loop_needed

                            # failed
                            case "response.failed":
                                logger.bind(ctx=ctx).error("Response failed")
                                response = event.response
                                err = getattr(response, "error", None)
                                inc = getattr(response, "incomplete_details", None)
                                details = {
                                    "response_id": response.id,
                                    "status": response.status,
                                    "model": getattr(response, "model", None),
                                    "error": {
                                        "type": getattr(err, "type", None),
                                        "code": getattr(err, "code", None),
                                        "message": getattr(err, "message", None)
                                                   or (str(err) if err else None),
                                        "param": getattr(err, "param", None),
                                    }
                                    if err
                                    else None,
                                    "incomplete_details": {
                                        "reason": getattr(inc, "reason", None),
                                        "duration": getattr(inc, "duration", None),
                                        "token_limit": getattr(inc, "token_limit", None),
                                    }
                                    if inc
                                    else None,
                                }
                                raise HTTPException(
                                    status_code=502, detail=details
                                )  # Bad Gateway to indicate upstream failure
                            case "response.incomplete":
                                logger.bind(ctx=ctx).error("Response incomplete")
                                response = event.response
                                err = getattr(response, "error", None)
                                inc = getattr(response, "incomplete_details", None)
                                details = {
                                    "response_id": response.id,
                                    "status": response.status,
                                    "max_output_tokens": getattr(response, "max_output_tokens", None),
                                    "model": getattr(response, "model", None),
                                    "error": {
                                        "type": getattr(err, "type", None),
                                        "code": getattr(err, "code", None),
                                        "message": getattr(err, "message", None)
                                                   or (str(err) if err else None),
                                        "param": getattr(err, "param", None),
                                    }
                                    if err
                                    else None,
                                    "incomplete_details": {
                                        "reason": getattr(inc, "reason", None),
                                        "duration": getattr(inc, "duration", None),
                                        "token_limit": getattr(inc, "token_limit", None),
                                    }
                                    if inc
                                    else None,
                                }
                                raise HTTPException(
                                    status_code=502, detail=details
                                )  # Bad Gateway to indicate upstream failure

                            case _:
                                logger.bind(ctx=ctx).warning(f"Unhandled event type: {event.type}")
                                continue
                    if loop_needed:
                        # yield send("info", message="Loop needed, continuing to next loop")

                        continue
                    else:
                        # yield send("done",message="No loop needed, finishing")
                        # yield send("response-loop", result=final_response.model_dump(exclude_none=True))
                        notify_complimentary(ctx, "Finishing")
                        logger.bind(ctx=ctx).info("Responses loop finished, no next loop needed")
                        # break
                        # logger.bind(ctx=ctx).info(f"Final response: {final_response}")
                        if final_response is None:
                            raise Exception("Stream ended without 'response.completed' event")
                        if await self._instructions_at_exit(ctx, app_state, responses_req, notify):
                            continue
                        return final_response

                except Exception:
                    # yield send("error", message=f"stream_exception: {e}")
                    raise

                # logger.bind(ctx=ctx).info("Responses loop finished")

            except Exception as e:
                logger.bind(ctx=ctx, error=str(e)).exception(
                    f"Error in response_loop_stream: {str(e)}"
                )
                raise e
                # yield send("response-loop", error=f"response_loop_exception: {str(e)}")

    @staticmethod
    async def get_model_based_config(
        ctx: Context, responses_req: LlmResponsesOpenaiV2Config
    ) -> Dict[str, Any]:
        responses_req_dict = responses_req.model_dump(exclude_none=True)
        return responses_req_dict

    async def keep_containers_alive(
        self,
        client: AsyncAzureOpenAI | ChatCompletionsClient | AsyncOpenAI,
        container_id: str,
        ctx: Context | None = None,
    ) -> None:
        """
        Keep container alive by retrieving it.
        This resets its last_active_at timestamp to prevent expiration.
        """
        try:
            await client.containers.retrieve(container_id)
            log = logger.bind(ctx=ctx) if ctx else logger
            log.info(f"Kept container {container_id} alive")
        except Exception as e:
            log = logger.bind(ctx=ctx) if ctx else logger
            log.warning(f"Failed to keep container {container_id} alive: {str(e)}")

    async def convert_output_item_to_input_item_param(
        self,
        ctx: Context,
        responses_req: LlmResponsesOpenaiV2Config,
        output_item: ResponseOutputItem,
        store: bool,
        client: AsyncAzureOpenAI | ChatCompletionsClient | AsyncOpenAI,
        user: UserORM,
        db: AsyncIOMotorDatabase,
        user_files: UserFiles | None = None,
        override_arguments: Any | None = None,
    ) -> ResponseInputItemParam | None | Tuple[ResponseInputItemParam, Message]:
        if store:
            match output_item:
                case ResponseOutputMessage():
                    # Cancel background keep-alive task when response is received
                    await self._cancel_keep_alive_task(responses_req, ctx)
                    
                    has_sandbox_path = await OpenaiResponsesV2.update_openai_urls_with_public_s3_urls(
                        ctx=ctx,
                        responses_req=responses_req,
                        output_item=output_item,
                        client=client,
                        user_files=user_files,
                    )
                    await self.update_new_public_s3_urls_with_openai_files(
                        ctx=ctx,
                        user=user,
                        db=db,
                        client=client,
                        responses_req=responses_req,
                        user_files=user_files,
                    )
                    output_param = ResponseOutputMessageParam(**output_item.model_dump(exclude_none=True))
                    if has_sandbox_path:
                        logger.bind(ctx=ctx).info(
                            "has_sandbox_path is True."
                        )
                        try:
                            sandbox_urls_message = await llm_func_check_retry_urls(
                                ctx=ctx,
                                user=user,
                                db=db,
                                assistant_output=output_item.content[0].text
                            )
                        except Exception as e:
                            logger.bind(ctx=ctx).exception(
                                f"llm_func_check_retry_urls failed; falling back to generic retry prompt: {e}"
                            )
                            sandbox_urls_message = ""
                        retry_msg = None
                        if OpenaiResponsesV2.SANDBOX_URLS_NOT_FOUND_MESSAGE not in sandbox_urls_message and responses_req.sandbox_retry_current < OpenaiResponsesV2.RETRY_MAX_COUNT:
                            logger.bind(ctx=ctx).info(
                                f"Sandbox paths detected in output and new public S3 URLs found, prompting user to retry with updated URLs with Retry count : {responses_req.sandbox_retry_current}"
                            )
                            responses_req.sandbox_retry_current += 1
                            msg_text = OpenaiResponsesV2.SANDBOX_PATH_RETRY_MESSAGE + "\n\n" + sandbox_urls_message
                            retry_msg = Message(
                                content=[
                                    ResponseInputTextParam(
                                        text=msg_text,
                                        type="input_text",
                                    )
                                ],
                                role="user",
                            )
                        else:
                            logger.bind(ctx=ctx).warning(
                                f"Resetting sandbox retry count: either no valid URLs extracted or max retry count reached. Current retry count: {responses_req.sandbox_retry_current}"
                            )
                            responses_req.sandbox_retry_current = 0
                        if retry_msg:
                            return (output_param, retry_msg)
                    else:
                        responses_req.sandbox_retry_current = 0
                    return output_param
                case ResponseFileSearchToolCall():
                    return ResponseFileSearchToolCallParam(**output_item.model_dump(exclude_none=True))
                case ResponseFunctionToolCall():
                    data = output_item.model_dump(exclude_none=True)
                    if override_arguments is not None:
                        data["arguments"] = override_arguments
                    return ResponseFunctionToolCallParam(**data)
                case ResponseFunctionWebSearch():
                    data = output_item.model_dump(exclude_none=True)
                    # SDK requires deprecated 'query' field for ActionSearch validation
                    if "action" in data and data["action"].get("type") == "search":
                        data["action"].setdefault("query", "")
                    # SDK requires 'url' and 'pattern' fields for ActionFind validation
                    if "action" in data and data["action"].get("type") == "find_in_page":
                        data["action"].setdefault("url", "")
                        data["action"].setdefault("pattern", "")
                    return ResponseFunctionWebSearchParam(**data)
                case ResponseComputerToolCall():
                    return ResponseComputerToolCallParam(
                        **output_item.model_dump(exclude_none=True)
                    )
                case ResponseReasoningItem():
                    return ResponseReasoningItemParam(**output_item.model_dump(exclude_none=True))
                case ResponseCompactionItem():
                    encrypted_len = len(output_item.encrypted_content or "")
                    if encrypted_len > COMPACTION_ENCRYPTED_CONTENT_LIMIT:
                        logger.bind(
                            ctx=ctx,
                            compaction_id=output_item.id,
                            encrypted_len=encrypted_len,
                            limit=COMPACTION_ENCRYPTED_CONTENT_LIMIT,
                        ).error("Compaction encrypted_content exceeds limit")
                        raise CompactionTooLargeError(
                            f"encrypted_content size {encrypted_len} exceeds limit {COMPACTION_ENCRYPTED_CONTENT_LIMIT}"
                        )
                    logger.bind(
                        ctx=ctx,
                        compaction_id=output_item.id,
                        encrypted_len=encrypted_len,
                        store=store,
                    ).info("Context compaction triggered by server")
                    responses_req.compaction_triggered = True
                    return ResponseCompactionItemParamParam(**output_item.model_dump(exclude_none=True))
                case response_output_item.ImageGenerationCall():
                    return ImageGenerationCall(**output_item.model_dump(exclude_none=True))
                case ResponseCodeInterpreterToolCall():
                    logger.bind(
                        ctx=ctx,
                        code=output_item.code[:100] if output_item.code else "",
                    ).info(
                        f"Processing code interpreter tool call with container {output_item.container_id}, status: {output_item.status}"
                    )
                    
                    # Check for container reset
                    if self._is_container_reset(output_item, ctx):
                        raise ContainerResetError(f"Container {output_item.container_id} reset detected")
                    
                    # Start keep-alive task for this container if not already running
                    if not getattr(responses_req, '_keep_alive_task', None):
                        logger.bind(ctx=ctx).info(
                            f"Starting keep-alive task for code interpreter container {output_item.container_id}"
                        )
                        await self._start_keep_alive_task(client, responses_req, output_item.container_id, ctx)
                    
                    data = output_item.model_dump(exclude_none=True)
                    data.setdefault("outputs", [])
                    return ResponseCodeInterpreterToolCallParam(**data)
                case response_output_item.LocalShellCall():
                    return LocalShellCall(**output_item.model_dump(exclude_none=True))
                case response_output_item.McpCall():
                    return McpCall(**output_item.model_dump(exclude_none=True))
                case response_output_item.McpListTools():
                    return McpListTools(**output_item.model_dump(exclude_none=True))
                case response_output_item.McpApprovalRequest():
                    return McpApprovalRequest(**output_item.model_dump(exclude_none=True))
                case _:
                    logger.error(f"Unhandled output item type: {type(output_item)}")
                    raise HTTPException(
                        status_code=501, detail=f"Unhandled output item type: {type(output_item)}"
                    )
        else:
            match output_item:
                case ResponseOutputMessage():
                    # Cancel background keep-alive task when response is received
                    await self._cancel_keep_alive_task(responses_req, ctx)
                    
                    has_sandbox_path = await OpenaiResponsesV2.update_openai_urls_with_public_s3_urls(
                        ctx=ctx,
                        responses_req=responses_req,
                        output_item=output_item,
                        client=client,
                        user_files=user_files,
                    )
                    await self.update_new_public_s3_urls_with_openai_files(
                        ctx=ctx,
                        user=user,
                        db=db,
                        client=client,
                        responses_req=responses_req,
                        user_files=user_files,
                    )
                    output_param = ResponseOutputMessageParam(**output_item.model_dump(exclude_none=True))
                    if has_sandbox_path:
                        logger.bind(ctx=ctx).info(
                            "has_sandbox_path is True."
                        )
                        try:
                            sandbox_urls_message = await llm_func_check_retry_urls(
                                ctx=ctx,
                                user=user,
                                db=db,
                                assistant_output=output_item.content[0].text
                            )
                        except Exception as e:
                            logger.bind(ctx=ctx).exception(
                                f"llm_func_check_retry_urls failed; falling back to generic retry prompt: {e}"
                            )
                            sandbox_urls_message = ""
                        retry_msg = None
                        if OpenaiResponsesV2.SANDBOX_URLS_NOT_FOUND_MESSAGE not in sandbox_urls_message and responses_req.sandbox_retry_current < OpenaiResponsesV2.RETRY_MAX_COUNT:
                            logger.bind(ctx=ctx).info(
                                f"Sandbox paths detected in output and new public S3 URLs found, prompting user to retry with updated URLs with Retry count : {responses_req.sandbox_retry_current}"
                            )
                            responses_req.sandbox_retry_current += 1
                            msg_text = OpenaiResponsesV2.SANDBOX_PATH_RETRY_MESSAGE + "\n\n" + sandbox_urls_message
                            retry_msg = Message(
                                content=[
                                    ResponseInputTextParam(
                                        text=msg_text,
                                        type="input_text",
                                    )
                                ],
                                role="user",
                            )
                        else:
                            logger.bind(ctx=ctx).warning(
                                f"Resetting sandbox retry count: either no valid URLs extracted or max retry count reached. Current retry count: {responses_req.sandbox_retry_current}"
                            )
                            responses_req.sandbox_retry_current = 0
                        if retry_msg:
                            return (output_param, retry_msg)
                    else:
                        responses_req.sandbox_retry_current = 0
                    return output_param
                case ResponseReasoningItem():
                    return ResponseReasoningItemParam(**output_item.model_dump(exclude_none=True))
                case ResponseCompactionItem():
                    encrypted_len = len(output_item.encrypted_content or "")
                    if encrypted_len > COMPACTION_ENCRYPTED_CONTENT_LIMIT:
                        logger.bind(
                            ctx=ctx,
                            compaction_id=output_item.id,
                            encrypted_len=encrypted_len,
                            limit=COMPACTION_ENCRYPTED_CONTENT_LIMIT,
                        ).error("Compaction encrypted_content exceeds limit")
                        raise CompactionTooLargeError(
                            f"encrypted_content size {encrypted_len} exceeds limit {COMPACTION_ENCRYPTED_CONTENT_LIMIT}"
                        )
                    logger.bind(
                        ctx=ctx,
                        compaction_id=output_item.id,
                        encrypted_len=encrypted_len,
                        store=store,
                    ).info("Context compaction triggered by server")
                    responses_req.compaction_triggered = True
                    return ResponseCompactionItemParamParam(**output_item.model_dump(exclude_none=True))

                case ResponseFunctionToolCall():
                    data = output_item.model_dump(exclude_none=True)
                    if override_arguments is not None:
                        data["arguments"] = override_arguments
                    return ResponseFunctionToolCallParam(**data)

                case (
                ResponseFileSearchToolCall()
                | ResponseFunctionWebSearch()
                | ResponseComputerToolCall()
                | ResponseCodeInterpreterToolCall()
                | ImageGenerationCall()
                | LocalShellCall()
                ):
                    # Log and start keep-alive task if this is a code interpreter call
                    if isinstance(output_item, ResponseCodeInterpreterToolCall):
                        logger.bind(
                            ctx=ctx,
                            code=output_item.code[:100] if output_item.code else "",
                        ).info(
                            f"Processing code interpreter tool call with container {output_item.container_id}, status: {output_item.status}"
                        )
                        
                        # Check for container reset
                        if self._is_container_reset(output_item, ctx):
                            raise ContainerResetError(f"Container {output_item.container_id} reset detected")
                        
                        # Start keep-alive task for this container if not already running
                        if not getattr(responses_req, '_keep_alive_task', None):
                            logger.bind(ctx=ctx).info(
                                f"Starting keep-alive task for code interpreter container {output_item.container_id}"
                            )
                            await self._start_keep_alive_task(client, responses_req, output_item.container_id, ctx)
                    
                    content = output_item.model_dump_json(exclude_none=True)
                    msg: Message = Message(
                        content=[ResponseInputTextParam(text=content, type="input_text")],
                        role="system",
                    )
                    return msg
                case response_output_item.McpCall():
                    return McpCall(**output_item.model_dump(exclude_none=True))
                case response_output_item.McpListTools():
                    return McpListTools(**output_item.model_dump(exclude_none=True))
                case response_output_item.McpApprovalRequest():
                    return McpApprovalRequest(**output_item.model_dump(exclude_none=True))
                case _:
                    logger.error(f"Unhandled output item type: {type(output_item)}")
                    raise HTTPException(
                        status_code=501, detail=f"Unhandled output item type: {type(output_item)}"
                    )

    async def process_response_output_item(
        self,
        ctx: Context,
        app_state: AppState,
        responses_req: LlmResponsesOpenaiV2Config,
        response_output_item: ResponseOutputItem,
        client: AsyncAzureOpenAI | ChatCompletionsClient | AsyncOpenAI,
        user: UserORM | None = None,
        db: AsyncIOMotorDatabase | None = None,
        notify: Notify | None = None,

    ) -> tuple[ResponseInputItemParam | None, Any | None]:
        match response_output_item:
            case ResponseFunctionToolCall():
                override_arguments: Any | None = None
                tool_args_for_execution: Any = response_output_item.arguments

                # Pre-execution: resolve Code Interpreter sandbox paths in tool arguments to public S3 URLs.
                # If we cannot resolve, do NOT run the tool; instead return an actionable message
                # instructing the model to regenerate the file(s).
                refs: set[str] = set()
                try:
                    refs = OpenaiResponsesV2._extract_mnt_data_refs_from_any(response_output_item.arguments)
                except Exception as e:
                    logger.bind(
                        ctx=ctx,
                        tool=response_output_item.name,
                        call_id=getattr(response_output_item, "call_id", None),
                        error=str(e),
                    ).exception("Failed extracting sandbox refs from tool arguments")
                    refs = set()

                if refs:
                    replacements, unresolved = await self._resolve_mnt_data_refs_to_s3(
                        ctx=ctx,
                        responses_req=responses_req,
                        refs=refs,
                        client=client,
                        user=user,
                        notify=notify,
                    )
                    if unresolved:
                        msg_text = OpenaiResponsesV2._format_regenerate_sandbox_files_message(unresolved)
                        tool_call_id = response_output_item.call_id
                        logger.bind(
                            ctx=ctx,
                            tool=response_output_item.name,
                            call_id=tool_call_id,
                            unresolved_paths=len(unresolved),
                            unresolved_paths_list=unresolved,
                        ).warning(
                            "Aborting tool call: sandbox paths in tool arguments could not be resolved; asking to regenerate files"
                        )
                        if notify:
                            notify(ctx=ctx, message="Tool call blocked: sandbox file paths are not accessible. Asking to regenerate.")
                        function_call_output = FunctionCallOutput(
                            type="function_call_output",
                            call_id=tool_call_id,
                            output=msg_text,
                        )
                        return function_call_output, override_arguments

                    rewritten_args = OpenaiResponsesV2._rewrite_tool_args_with_replacements(
                        response_output_item.arguments, replacements
                    )
                    # Ensure we always store string arguments back into responses_req.input for replay safety.
                    if not isinstance(rewritten_args, str):
                        try:
                            rewritten_args = json.dumps(rewritten_args)
                        except Exception:
                            rewritten_args = response_output_item.arguments
                    tool_args_for_execution = rewritten_args
                    override_arguments = rewritten_args

                # Intercept chat_history_query_openai here (before run_tool) because the
                # tool needs access to responses_req.input_archive, which ToolFactory
                # does not have. Same pattern as Claude's chat_history_query_claude.
                if response_output_item.name == "chat_history_query_openai":
                    tool_output = await self._run_chat_history_query_openai(
                        ctx=ctx,
                        app_state=app_state,
                        responses_req=responses_req,
                        response_output_item=response_output_item,
                        tool_args_override=tool_args_for_execution,
                        notify=notify,
                    )
                else:
                    tool_output: DataBlock = await self.run_tool(
                        ctx=ctx,
                        app_state=app_state,
                        response_output_item=response_output_item,
                        tool_args_override=tool_args_for_execution,
                        user=user,
                        db=db,
                        notify=notify,
                        output_text_to_file_char_limit=responses_req.tool_output_text_to_file_char_limit,
                    )

                # Special handling for enable_mcp: inject the MCP into runtime tools config
                if response_output_item.name == "enable_mcp" and tool_output.text:
                    try:
                        result = json.loads(tool_output.text)
                        if result.get("success") and result.get("server_name"):
                            server_name = result["server_name"]
                            user_mcp = UserMCP(user=user, db=db)
                            mcp_doc = await user_mcp.get_mcp_by_server_name(
                                ctx=ctx, mcp_server_name=server_name, decrypt=True
                            )
                            if mcp_doc:
                                # Build Mcp tool param matching OpenAI's expected format
                                description = (
                                    mcp_doc.description
                                    if (mcp_doc.description and mcp_doc.description != "")
                                    else None
                                )
                                request_headers = (
                                    mcp_doc.request_headers
                                    if (mcp_doc.request_headers and len(mcp_doc.request_headers.keys()) > 0)
                                    else None
                                )
                                authorization = (
                                    mcp_doc.oauth_token
                                    if (mcp_doc.oauth_token and mcp_doc.oauth_token != "")
                                    else None
                                )
                                mcp_tool_param = Mcp(
                                    type="mcp",
                                    server_label=mcp_doc.server_name,
                                    server_url=mcp_doc.server_url,
                                    server_description=description,
                                    headers=request_headers,
                                    authorization=authorization,
                                    require_approval="never",
                                )
                                # Check if MCP with same server_label already exists
                                existing_labels = set()
                                for t in responses_req.tools or []:
                                    if hasattr(t, "server_label"):
                                        existing_labels.add(t.server_label)
                                    elif isinstance(t, dict) and t.get("server_label"):
                                        existing_labels.add(t["server_label"])

                                # Inject into runtime tools config only if not duplicate
                                if mcp_doc.server_name not in existing_labels:
                                    if responses_req.tools is None:
                                        responses_req.tools = []
                                    responses_req.tools.append(mcp_tool_param)
                                    logger.bind(
                                        ctx=ctx,
                                        server_name=server_name,
                                    ).info(
                                        f"Dynamically added MCP '{server_name}' to runtime tools config"
                                    )
                                else:
                                    logger.bind(
                                        ctx=ctx,
                                        server_name=server_name,
                                    ).info(
                                        f"MCP '{server_name}' already exists in tools, skipping duplicate injection"
                                    )
                    except Exception as e:
                        logger.bind(
                            ctx=ctx,
                            tool=response_output_item.name,
                            error=str(e),
                        ).exception("Failed to inject MCP into runtime tools config")

                # Post-execution: if tool input contains huge strings, offload them to S3 and
                # override the stored tool-call arguments with URLs so the next loop payload is small.
                # NOTE: This is for *history only*; the tool already ran with tool_args_for_execution.
                try:
                    tool_args_for_storage = (
                        override_arguments if override_arguments is not None else response_output_item.arguments
                    )
                    sanitized_args, input_urls = await ToolFactory.maybe_save_tool_input_to_file_urls(
                        ctx=ctx,
                        user=user,
                        tool_name=response_output_item.name,
                        tool_args=tool_args_for_storage,
                        notify=notify,
                        input_text_to_file_char_limit=10000,
                    )
                    if input_urls:
                        override_arguments = sanitized_args
                        # Ensure stored tool-call arguments are always a JSON string (OpenAI expects `arguments` as string).
                        if override_arguments is not None and not isinstance(override_arguments, str):
                            try:
                                override_arguments = json.dumps(override_arguments)
                            except Exception:
                                override_arguments = str(override_arguments)
                        logger.bind(
                            ctx=ctx,
                            tool=response_output_item.name,
                            call_id=getattr(response_output_item, "call_id", None),
                            offloaded_tool_inputs=len(input_urls),
                            offloaded_tool_input_urls_sample=input_urls[:3],
                        ).info(
                            "Offloaded large tool inputs to URLs and queued them for OpenAI file upload; will override stored tool-call arguments for next loop"
                        )
                        for url in input_urls:
                            responses_req.sandbox_url_to_s3_url[url] = url
                except Exception as e:
                    logger.bind(
                        ctx=ctx,
                        tool=response_output_item.name,
                        tool_input=str(response_output_item.arguments)[:500],
                        error=str(e),
                    ).exception("Error while offloading large tool inputs to file URLs")

                for url in tool_output.urls:
                    responses_req.sandbox_url_to_s3_url[url] = url

                # Extract URL→FileObject mapping from file_uploader_openai tool output
                # This ensures consolidation/zipping logic can find source URLs for these file_ids
                if response_output_item.name == "file_uploader_openai" and tool_output.text:
                    try:
                        parsed = json.loads(tool_output.text)
                        # Handle ToolFactory serialization: entire DataBlock is serialized as JSON
                        # So parsed = {"text": "{url: FileObject}", "urls": [], "file_ids": [...], ...}
                        # We need to extract and parse the nested "text" field
                        if isinstance(parsed, dict) and "text" in parsed and isinstance(parsed.get("text"), str):
                            url_to_file_dict = json.loads(parsed["text"])
                        else:
                            url_to_file_dict = parsed
                        added_count = 0
                        for url, file_dict in url_to_file_dict.items():
                            if url not in responses_req.file_url_to_file_object_map:
                                responses_req.file_url_to_file_object_map[url] = FileObject(**file_dict)
                                added_count += 1
                        logger.bind(ctx=ctx, added_urls=added_count, total_urls=len(url_to_file_dict)).info(
                            "Added file_uploader_openai URLs to file_url_to_file_object_map"
                        )
                    except Exception as e:
                        logger.bind(ctx=ctx, error=str(e)).warning(
                            "Failed to parse file_uploader_openai output for URL mapping, trying fallback"
                        )
                        # Fallback: fetch full content from stored file URL
                        if tool_output.urls:
                            try:
                                async with httpx.AsyncClient(timeout=60.0, follow_redirects=True) as http_client:
                                    response = await http_client.get(str(tool_output.urls[0]))
                                    response.raise_for_status()
                                    full_content = response.text
                                parsed = json.loads(full_content)
                                if isinstance(parsed, dict) and "text" in parsed and isinstance(parsed.get("text"), str):
                                    url_to_file_dict = json.loads(parsed["text"])
                                else:
                                    url_to_file_dict = parsed
                                added_count = 0
                                for url, file_dict in url_to_file_dict.items():
                                    if url not in responses_req.file_url_to_file_object_map:
                                        responses_req.file_url_to_file_object_map[url] = FileObject(**file_dict)
                                        added_count += 1
                                logger.bind(ctx=ctx, added_urls=added_count).info(
                                    "Added file_uploader_openai URLs from fallback"
                                )
                            except Exception as fallback_e:
                                logger.bind(ctx=ctx, fallback_url=str(tool_output.urls[0]), error=str(fallback_e)).warning(
                                    "Failed to parse file_uploader_openai output from fallback URL"
                                )

                # IMPORTANT: This step involves network calls (downloads/uploads) and must not prevent
                # tool-call history from being persisted. Failures should be logged and ignored.
                try:
                    await self.update_new_public_s3_urls_with_openai_files(
                        ctx=ctx,
                        user=user,
                        db=db,
                        client=client,
                        responses_req=responses_req,  # user_files=user_files
                        additional_file_ids=tool_output.file_ids,
                    )
                except Exception as e:
                    logger.bind(
                        ctx=ctx,
                        tool=response_output_item.name,
                        call_id=getattr(response_output_item, "call_id", None),
                        error=str(e),
                    ).exception("Error in update_new_public_s3_urls_with_openai_files after tool call")

                tool_call_id = response_output_item.call_id
                function_call_output = FunctionCallOutput(
                    type="function_call_output", call_id=tool_call_id, output=tool_output.text
                )
                logger.bind(
                    ctx=ctx,
                    tool=response_output_item.name,
                    call_id=tool_call_id,
                    override_arguments_applied=override_arguments is not None,
                ).info("Prepared function_call_output for tool call")
                return function_call_output, override_arguments

        return None, None

    @staticmethod
    def _normalize_input_archive_to_turns(
        responses_req: LlmResponsesOpenaiV2Config,
    ) -> list[dict[str, Any]]:
        """Convert input_archive entries into turn records for chat_history_query_openai.

        Each turn = {"turn": N, "messages": input + in_turn_messages + output}.
        Reasoning and compaction items are filtered out by the tool's _normalize_message,
        so no pre-filtering is needed here.
        """
        turns: list[dict[str, Any]] = []
        for entry in responses_req.input_archive or []:
            messages: list[Any] = []
            for bucket in (entry.input, entry.in_turn_messages, entry.output):
                for msg in bucket or []:
                    if hasattr(msg, "model_dump"):
                        try:
                            messages.append(msg.model_dump(exclude_none=True))
                        except Exception:
                            messages.append(msg)
                    elif isinstance(msg, dict):
                        messages.append(dict(msg))
                    else:
                        messages.append(msg)
            turns.append({"turn": int(entry.turn_number), "messages": messages})
        return turns

    async def _run_chat_history_query_openai(
        self,
        ctx: Context,
        app_state: AppState,
        responses_req: LlmResponsesOpenaiV2Config,
        response_output_item: ResponseFunctionToolCall,
        tool_args_override: Any | None = None,
        notify: Notify | None = None,
    ) -> DataBlock:
        """Execute chat_history_query_openai with the live input_archive injected.

        Mirrors Claude's chat_history_query_claude interception: the tool needs
        access to the conversation archive, which ToolFactory.run_tool cannot provide.
        """
        try:
            raw_args = (
                tool_args_override
                if tool_args_override is not None
                else response_output_item.arguments
            )
            try:
                query_args = (
                    json.loads(raw_args) if isinstance(raw_args, str) else dict(raw_args or {})
                )
            except Exception:
                query_args = {}

            query_req = ChatHistoryQueryReq.model_validate(query_args)
            normalized_turns = self._normalize_input_archive_to_turns(responses_req)
            pinned = [
                dict(m)
                if isinstance(m, dict)
                else (m.model_dump(exclude_none=True) if hasattr(m, "model_dump") else m)
                for m in (responses_req.initial_input_snapshot or [])
            ]
            query_req.set_turn_history(
                turn_history=normalized_turns,
                pinned_messages=pinned,
            )

            query_res = ActionChatHistoryQueryOpenai._execute_query(query_req)
            payload = query_res.model_dump(exclude_none=True)

            logger.bind(
                ctx=ctx,
                tool="chat_history_query_openai",
                available_turn_history_count=len(normalized_turns),
                total_count=payload.get("total_count"),
                returned_count=payload.get("returned_count"),
            ).info("Injected input_archive into chat_history_query_openai")

            return DataBlock(text=json.dumps(payload, ensure_ascii=False))
        except Exception as exc:
            logger.bind(ctx=ctx, tool="chat_history_query_openai").exception(
                "Failed executing chat_history_query_openai with injected history: {}",
                exc,
            )
            return DataBlock(
                text=f"error : Tool chat_history_query_openai failed with error: {exc}"
            )

    async def run_tool(
        self,
        ctx: Context,
        app_state: AppState,
        response_output_item: ResponseFunctionToolCall,
        tool_args_override: Any | None = None,
        user: UserORM | None = None,
        db: AsyncIOMotorDatabase | None = None,
        notify: Notify | None = None,
        output_text_to_file_char_limit: int | None = None,

    ) -> DataBlock:
        tool_name = response_output_item.name
        tool_args = (
            tool_args_override if tool_args_override is not None else response_output_item.arguments
        )
        tool_output = DataBlock()
        try:
            # run tool
            tool_factory = ToolFactory(user=user, db=db)
            tool_output: DataBlock = await tool_factory.run_tool(
                ctx,
                app_state,
                tool_name,
                tool_args,
                notify=notify,
                output_text_to_file_char_limit=output_text_to_file_char_limit,
            )
            logger.bind(
                ctx=ctx,
                app_code=AppCode.ACTION,
                tool=tool_name,
                # tool_input=tool_args,
                # tool_output=tool_output,
            ).info(f"Ran tool: {tool_name}")
        except Exception as e:
            logger.bind(
                ctx=ctx,
                app_code=AppCode.ACTION,
                tool=tool_name,
                tool_input=str(tool_args)[:500],
                tool_error=str(e),
            ).exception(f"Error in tool_call: {str(e)}")
            tool_output = DataBlock(text=str(e))

        return tool_output

    async def is_next_loop_needed(self, response_output_item: ResponseOutputItem) -> bool:
        is_next_loop_needed: bool = False
        match response_output_item:
            case ResponseFunctionToolCall():
                is_next_loop_needed = True
            # case ResponseFileSearchToolCall():
            #     is_next_loop_needed = True
            # case ResponseCodeInterpreterToolCall():
            #     is_next_loop_needed = True
            # case ResponseComputerToolCall():
            #     is_next_loop_needed = True
            # case ResponseFunctionWebSearch():
            #     is_next_loop_needed = True
        return is_next_loop_needed

    # commented out as it was giving unexpected results in text output
    # SANDBOX_PATH_RETRY_MESSAGE = (
    #     "The previous response still contains sandbox file paths (sandbox:/mnt/data/...) "
    #     "that could not be replaced with accessible URLs. Please try generating the file again "
    #     "and ensure the output uses accessible URLs instead of sandbox paths."
    # )
    SANDBOX_PATH_RETRY_MESSAGE = (
        "Try generating the following files again"
    )

    @staticmethod
    def _format_regenerate_sandbox_files_message(unresolved_paths: list[str]) -> str:
        """Tool-output message when a tool call references sandbox paths we cannot resolve.

        The goal is to produce a highly actionable message ("regenerate the file(s)") so the LLM
        can recover without repeatedly calling tools with inaccessible sandbox paths.
        """
        lines = [
            "This tool call references file paths from the Code Interpreter sandbox (sandbox:/mnt/data/... or /mnt/data/...).",
            "External tools cannot access sandbox paths.",
            "",
            "Please regenerate the file(s) and retry the tool call using the resulting public URL(s) (https/S3) instead of sandbox paths.",
            "",
            "Unresolved sandbox paths:",
        ]
        for p in unresolved_paths[:25]:
            lines.append(f"- {p}")
        if len(unresolved_paths) > 25:
            lines.append(f"- ... ({len(unresolved_paths) - 25} more)")
        return "\n".join(lines)

    @staticmethod
    def _extract_mnt_data_refs_from_str(s: str) -> set[str]:
        """Extract sandbox/mnt-data references from an arbitrary string.

        Accepts both:
        - sandbox:/mnt/data/...
        - /mnt/data/...

        Conservative extraction: stops at whitespace and common delimiters.
        """
        if not s:
            return set()
        pattern = r"(sandbox:/mnt/data/[^\s\"'<>\)\]\}]+|/mnt/data/[^\s\"'<>\)\]\}]+)"
        out: set[str] = set()
        for m in re.finditer(pattern, s):
            cand = m.group(0).rstrip(".,;:!?")
            out.add(cand)
        return out

    @staticmethod
    def _extract_mnt_data_refs_from_any(tool_args: Any) -> set[str]:
        """Deep-scan tool args and extract any sandbox/mnt-data references."""

        def _walk(v: Any) -> set[str]:
            if isinstance(v, str):
                return OpenaiResponsesV2._extract_mnt_data_refs_from_str(v)
            if isinstance(v, dict):
                acc: set[str] = set()
                for vv in v.values():
                    acc |= _walk(vv)
                return acc
            if isinstance(v, list):
                acc: set[str] = set()
                for x in v:
                    acc |= _walk(x)
                return acc
            return set()

        if tool_args is None:
            return set()

        if isinstance(tool_args, str):
            try:
                parsed = json.loads(tool_args)
            except Exception:
                return OpenaiResponsesV2._extract_mnt_data_refs_from_str(tool_args)
            return _walk(parsed)

        return _walk(tool_args)

    @staticmethod
    def _mnt_data_ref_to_relative_path(ref: str) -> str | None:
        if not isinstance(ref, str) or not ref:
            return None
        if ref.startswith(OpenaiResponsesV2.SANDBOX_PATH_PREFIX):
            return ref[len(OpenaiResponsesV2.SANDBOX_PATH_PREFIX):].lstrip("/")
        if ref.startswith(OpenaiResponsesV2.MNT_DATA_PATH_PREFIX):
            return ref[len(OpenaiResponsesV2.MNT_DATA_PATH_PREFIX):].lstrip("/")
        return None

    @staticmethod
    def _find_latest_completed_code_interpreter_container_id(
        responses_req: LlmResponsesOpenaiV2Config,
    ) -> str | None:
        """Find the latest completed code interpreter container_id from input history."""
        items = getattr(responses_req, "input", None)
        if not isinstance(items, list):
            return None
        for item in reversed(items):
            if isinstance(item, BaseModel):
                item = item.model_dump(exclude_none=True)
            if not isinstance(item, dict):
                continue
            if (
                item.get("type") == "code_interpreter_call"
                and item.get("status") == "completed"
                and item.get("container_id")
            ):
                return item.get("container_id")
        return None

    async def _resolve_mnt_data_refs_to_s3(
        self,
        ctx: Context,
        responses_req: LlmResponsesOpenaiV2Config,
        refs: set[str],
        client: AsyncAzureOpenAI | ChatCompletionsClient | AsyncOpenAI,
        user: UserORM | None,
        notify: Notify | None = None,
    ) -> tuple[dict[str, str], list[str]]:
        """Resolve sandbox/mnt-data refs to public S3 URLs.

        Returns (replacements, unresolved_refs).
        - replacements maps encountered ref strings to S3 URLs
        - unresolved_refs are refs we couldn't resolve (caller should abort tool call)
        """
        if not refs:
            return {}, []
        if not user:
            return {}, sorted(refs)

        # Reuse any already-known mappings first.
        replacements: dict[str, str] = {}
        refs_to_resolve: set[str] = set()
        for ref in refs:
            mapped = responses_req.sandbox_url_to_s3_url.get(ref)
            if not mapped:
                rel = OpenaiResponsesV2._mnt_data_ref_to_relative_path(ref)
                if rel:
                    mapped = (
                        responses_req.sandbox_url_to_s3_url.get(
                            f"{OpenaiResponsesV2.SANDBOX_PATH_PREFIX}{rel}"
                        )
                        or responses_req.sandbox_url_to_s3_url.get(
                            f"{OpenaiResponsesV2.MNT_DATA_PATH_PREFIX}{rel}"
                        )
                    )
            if mapped:
                replacements[ref] = mapped
            else:
                refs_to_resolve.add(ref)

        if not refs_to_resolve:
            return replacements, []

        container_id = OpenaiResponsesV2._find_latest_completed_code_interpreter_container_id(responses_req)
        if not container_id:
            return replacements, sorted(refs_to_resolve)

        # Some client implementations (e.g., Azure AI Inference ChatCompletionsClient) do not support the Containers API.
        if not hasattr(client, "containers"):
            logger.bind(ctx=ctx, client_type=type(client).__name__).warning(
                "Client does not support containers API; cannot resolve sandbox paths in tool arguments"
            )
            return replacements, sorted(refs_to_resolve)

        s3 = get_public_s3()
        user_files = UserFiles(user=user, s3=s3)

        # List container files once and build lookups.
        rel_to_file: dict[str, Any] = {}
        base_to_files: dict[str, list[Any]] = {}
        try:
            files_list = await client.containers.files.list(container_id=container_id)
            async for f in files_list:
                p = getattr(f, "path", None)
                if not isinstance(p, str) or not p:
                    continue
                rel = p.split("/mnt/data/")[-1].lstrip("/") if "/mnt/data/" in p else p.split("/")[-1]
                rel_to_file[rel] = f
                base = rel.split("/")[-1]
                base_to_files.setdefault(base, []).append(f)
        except Exception as e:
            logger.bind(ctx=ctx, container_id=container_id, error=str(e)).exception(
                "Failed listing container files while resolving sandbox tool args"
            )
            return replacements, sorted(refs_to_resolve)

        unresolved: list[str] = []
        for ref in sorted(refs_to_resolve):
            rel = OpenaiResponsesV2._mnt_data_ref_to_relative_path(ref)
            if not rel:
                unresolved.append(ref)
                continue

            f = rel_to_file.get(rel)
            if f is None:
                base = rel.split("/")[-1]
                candidates = base_to_files.get(base) or []
                if candidates:
                    # Basename-only matching is risky if multiple files share the same basename.
                    # Only use it when unambiguous.
                    assistant = [x for x in candidates if getattr(x, "source", None) == "assistant"]
                    if len(assistant) == 1:
                        f = assistant[0]
                    elif len(candidates) == 1:
                        f = candidates[0]
                    else:
                        logger.bind(
                            ctx=ctx,
                            container_id=container_id,
                            sandbox_ref=ref,
                            relative_path=rel,
                            basename=base,
                            candidate_paths=[getattr(x, "path", None) for x in candidates],
                        ).warning(
                            "Ambiguous basename match for sandbox tool-arg path; refusing to guess"
                        )
                        f = None

            if f is None:
                unresolved.append(ref)
                continue

            try:
                res = await client.containers.files.content.retrieve(file_id=f.id, container_id=container_id)
                assert res is not None, "container file content response is None"
                assert res.response.status_code == 200, f"container file download failed: {res.response.status_code}"
                file_obj = io.BytesIO(res.content)
                uuid = str(gen_uuid())
                basename = rel.split("/")[-1]
                s3_filename = f"{uuid}-{basename}"
                notify(ctx=ctx, message="Uploading sandbox file to S3 for tool call") if notify else None
                file_metas = await user_files.upload_binary(ctx=ctx, file_obj=file_obj, filename=s3_filename)
                assert file_metas and file_metas[0].url
                s3_url = str(file_metas[0].url)

                for k in {ref, f"{OpenaiResponsesV2.SANDBOX_PATH_PREFIX}{rel}", f"{OpenaiResponsesV2.MNT_DATA_PATH_PREFIX}{rel}"}:
                    replacements[k] = s3_url
                    responses_req.sandbox_url_to_s3_url[k] = s3_url

                logger.bind(ctx=ctx, container_id=container_id, sandbox_ref=ref, relative_path=rel, s3_url=s3_url).info(
                    "Resolved sandbox tool-arg path to S3"
                )
            except Exception as e:
                logger.bind(ctx=ctx, container_id=container_id, sandbox_ref=ref, relative_path=rel, error=str(e)).exception(
                    "Failed resolving sandbox tool-arg path to S3"
                )
                unresolved.append(ref)

        return replacements, unresolved

    @staticmethod
    def _rewrite_tool_args_with_replacements(tool_args: Any, replacements: dict[str, str]) -> Any:
        if not replacements or tool_args is None:
            return tool_args

        def _rw_str(s: str) -> str:
            out = s
            # Replace longest keys first to avoid prefix-collisions (e.g., /mnt/data/a vs /mnt/data/a/b).
            for old, new in sorted(replacements.items(), key=lambda kv: len(kv[0]), reverse=True):
                if old in out:
                    out = out.replace(old, new)
            return out

        def _walk(v: Any) -> Any:
            if isinstance(v, str):
                return _rw_str(v)
            if isinstance(v, dict):
                return {k: _walk(vv) for k, vv in v.items()}
            if isinstance(v, list):
                return [_walk(x) for x in v]
            return v

        if isinstance(tool_args, str):
            try:
                parsed = json.loads(tool_args)
            except Exception:
                return _rw_str(tool_args)
            if isinstance(parsed, (dict, list)):
                rewritten = _walk(parsed)
                try:
                    return json.dumps(rewritten)
                except Exception:
                    return tool_args
            rewritten_scalar = _walk(parsed)
            if isinstance(rewritten_scalar, str):
                return rewritten_scalar
            try:
                return json.dumps(rewritten_scalar)
            except Exception:
                return tool_args

        return _walk(tool_args)


    @staticmethod
    async def update_openai_urls_with_public_s3_urls(
        ctx: Context,
        responses_req: LlmResponsesOpenaiV2Config,
        output_item: ResponseOutputMessage,
        client: AsyncAzureOpenAI | ChatCompletionsClient | AsyncOpenAI,
        user_files: UserFiles | None = None,
    ) -> bool:
        if not user_files:
            logger.bind(ctx=ctx,  #output_item=output_item
                        ).error(
                "user_files is None, cant update openai urls with public s3 urls"
            )
            return False

        for content in output_item.content:
            for annotation in content.annotations:
                # Convert BaseModel to dict for logging
                annotation_dict = {}
                if isinstance(annotation, BaseModel):
                    annotation_dict = annotation.model_dump()
                else:
                    annotation_dict = annotation
                
                # Log annotation details
                # logger.bind(
                #     ctx=ctx,
                #     annotation=annotation_dict,
                #     annotation_type=annotation_dict.get('type'),
                #     filename=annotation_dict.get('filename'),
                #     file_id=annotation_dict.get('file_id'),
                #     container_id=annotation_dict.get('container_id'),
                #     start_index=annotation_dict.get('start_index'),
                #     end_index=annotation_dict.get('end_index'),
                # ).info("Processing annotation")
                
                try:
                    res: _legacy_response.HttpxBinaryResponseContent | None = None
                    annotation_type = annotation_dict.get("type")
                    match annotation_type:
                        case "file_citation":
                            annotation_filename = annotation_dict.get("filename")
                            file_id = annotation_dict.get("file_id")
                            if not annotation_filename or not file_id:
                                logger.bind(ctx=ctx, annotation=annotation_dict).warning(
                                    "file_citation annotation missing filename or file_id; skipping"
                                )
                                continue
                            res = await client.files.content(file_id=file_id)
                        case "container_file_citation":
                            annotation_filename = annotation_dict.get("filename")
                            file_id = annotation_dict.get("file_id")
                            container_id = annotation_dict.get("container_id")
                            if not annotation_filename or not file_id or not container_id:
                                logger.bind(ctx=ctx, annotation=annotation_dict).warning(
                                    "container_file_citation annotation missing filename, file_id, or container_id; skipping"
                                )
                                continue
                            res = await client.containers.files.content.retrieve(
                                file_id=file_id, container_id=container_id
                            )
                        case "url_citation" | "file_path":
                            # These annotations don't correspond to downloadable sandbox files.
                            continue
                        case _:
                            logger.bind(
                                ctx=ctx, annotation_type=annotation_type, annotation=annotation_dict
                            ).info("Unknown annotation type; skipping")
                            continue
                    assert res is not None, "res for content is None"
                    assert res.response.status_code == 200, (
                        f"res.response.status_code is not 200: {res.response.status_code}"
                    )
                    file_obj = io.BytesIO(res.content)
                    uuid = str(gen_uuid())
                    filename = f"{uuid}-{annotation_filename}"
                    file_metas = await user_files.upload_binary(
                        ctx=ctx, file_obj=file_obj, filename=filename
                    )
                    assert len(file_metas) > 0
                    responses_req.sandbox_url_to_s3_url[
                        f"{OpenaiResponsesV2.SANDBOX_PATH_PREFIX}{annotation_filename}"
                    ] = str(file_metas[0].url)
                except Exception as e:
                    if isinstance(e, BadRequestError) and e.message.lower().__contains__(
                        "not allowed to download"
                    ):
                        logger.bind(ctx=ctx,  #output_item=output_item
                                    ).info(
                            "Wasn't allowed to download possibly coz it is uploaded by us and not generated by llm"
                        )
                    else:
                        logger.bind(
                            ctx=ctx, annotation=annotation, error=str(e)
                        ).exception(f"Error in update_openai_urls_with_public_s3_urls: {str(e)}")

        # Fallback: Process code interpreter files without annotations
        # Find most recent successful code_interpreter_call
        container_id = None
        for item in reversed(responses_req.input):
            # If it's a BaseModel object, convert it to dict first
            if isinstance(item, BaseModel):
                logger.bind(ctx=ctx, item=item, item_type=type(item).__name__).info(
                    "Input item is a BaseModel, converting via model_dump()"
                )
                item = item.model_dump()

            # Handle TypedDict (dictionary) types - safe access with .get()
            item_type = item.get("type")
            item_status = item.get("status")
            item_container_id = item.get("container_id")

            # Log what we found in this item, commenting this for now as it is increasing the number of logs for now
            # logger.bind(
            #     ctx=ctx,
            #     item_type=item_type,
            #     item_status=item_status,
            #     extracted_container_id=item_container_id,
            # ).info("Extracted attributes from input item")

            if (
                item_type == "code_interpreter_call"
                and item_status == "completed"
                and item_container_id
            ):
                container_id = item_container_id
                logger.bind(ctx=ctx, container_id=container_id).info(
                    "Found completed code_interpreter_call with container_id"
                )
                break

        if not container_id:
            logger.bind(ctx=ctx).info("No successful code_interpreter_call found for fallback")

        if container_id:
            logger.bind(ctx=ctx, container_id=container_id).info(
                "Starting fallback container file processing"
            )
            try:
                # List all files in container
                files_list = await client.containers.files.list(container_id=container_id)
                logger.bind(ctx=ctx, container_id=container_id).info(
                    "Retrieved files list from container"
                )
                
                # Debug: Log complete container files data using sync client
                try:
                    container_files_data = get_container_files_sync(container_id=container_id)
                    logger.bind(
                        ctx=ctx,
                        container_id=container_id,
                        files_data=container_files_data["data"],
                        total_files=len(container_files_data["data"]),
                        has_more=container_files_data["has_more"],
                    ).info("Container files data from sync client")
                except Exception as e:
                    logger.bind(ctx=ctx, container_id=container_id, error=str(e)).warning(
                        "Failed to get container files data via sync client"
                    )

                # Track processed files for summary
                processed_files = []

                # Process assistant-generated files not yet in mapping
                async for file in files_list:
                    # logger.bind(
                    #     ctx=ctx, container_id=container_id, file_id=file.id, file_path=file.path, file_source=file.source
                    # ).info("Processing file from container")
                    if file.source == "assistant":
                        filename = file.path.split("/")[-1]
                        old_sandbox_url = f"{OpenaiResponsesV2.SANDBOX_PATH_PREFIX}{filename}"

                        # Correct sandbox URL key to use full subfolder path (annotation creates flat keys, text may have subfolders)
                        full_relative = file.path.split("/mnt/data/")[-1] if "/mnt/data/" in file.path else filename
                        correct_sandbox_url = f"{OpenaiResponsesV2.SANDBOX_PATH_PREFIX}{full_relative}"

                        # Skip if already processed with correct key
                        if correct_sandbox_url in responses_req.sandbox_url_to_s3_url:
                            continue

                        # If annotation created a flat key but actual path has subfolder, fix the key
                        if old_sandbox_url in responses_req.sandbox_url_to_s3_url:
                            s3_url = responses_req.sandbox_url_to_s3_url.pop(old_sandbox_url)
                            responses_req.sandbox_url_to_s3_url[correct_sandbox_url] = s3_url
                            logger.bind(
                                ctx=ctx,
                                old_sandbox_url=old_sandbox_url,
                                correct_sandbox_url=correct_sandbox_url,
                                s3_url=s3_url,
                            ).info("Corrected sandbox URL key from flat to full subfolder path")
                            processed_files.append(filename)
                            continue

                        # Download file content
                        res = await client.containers.files.content.retrieve(
                            file_id=file.id, container_id=container_id
                        )

                        assert res is not None, "res for content is None"
                        assert res.response.status_code == 200, (
                            f"res.response.status_code is not 200: {res.response.status_code}"
                        )

                        # Upload to S3
                        file_obj = io.BytesIO(res.content)
                        uuid = str(gen_uuid())
                        s3_filename = f"{uuid}-{filename}"
                        file_metas = await user_files.upload_binary(
                            ctx=ctx, file_obj=file_obj, filename=s3_filename
                        )

                        assert len(file_metas) > 0

                        # Update mapping with correct sandbox URL
                        responses_req.sandbox_url_to_s3_url[correct_sandbox_url] = str(file_metas[0].url)
                        processed_files.append(filename)

                        logger.bind(
                            ctx=ctx, file_id=file.id, filename=filename, s3_url=file_metas[0].url
                        ).info("Processed code interpreter file without annotation")

                # Log summary
                logger.bind(
                    ctx=ctx,
                    container_id=container_id,
                    total_files_processed=len(processed_files),
                    processed_files=processed_files,
                ).info("Fallback container file processing completed")

            except Exception as e:
                logger.bind(ctx=ctx, container_id=container_id, error=str(e)).exception(
                    "Error in fallback container file processing"
                )

        # First loop - exact sandbox URL replacements from sandbox_url_to_s3_url map
        for sandbox_url, public_s3_url in responses_req.sandbox_url_to_s3_url.items():
            # Use case-insensitive matching
            father_sandbox_url = "sandbox:" + sandbox_url
            # logger.bind(ctx=ctx).info(f"Father sandbox URL: {father_sandbox_url}")
            if sandbox_url.lower() in output_item.content[0].text.lower() and  father_sandbox_url.lower() not in output_item.content[0].text.lower():
                pattern = re.escape(sandbox_url)
                output_item.content[0].text = re.sub(
                    pattern,
                    public_s3_url,
                    output_item.content[0].text,
                    flags=re.IGNORECASE
                )
                logger.bind(ctx=ctx, sandbox_url=sandbox_url).info(
                    "Replaced exact sandbox URL with S3 URL (case-insensitive)"
                )

        # Second loop - construct sandbox URLs from file_url_to_file_object_map
        # OpenAI uses format: sandbox:/mnt/data/{file_id}-{filename} or sandbox:/mnt/data/subfolder/{file_id}-{filename}
        for s3_url, file_obj in responses_req.file_url_to_file_object_map.items():
            # Find matching sandbox URL by checking if file_obj.filename ends with sandbox filename
            # This works for any prefix format (UUID, file-id, etc.) without assumptions
            subfolder_path = None
            for sandbox_url in responses_req.sandbox_url_to_s3_url.keys():
                if sandbox_url.startswith(OpenaiResponsesV2.SANDBOX_PATH_PREFIX):
                    # Extract relative path and filename from sandbox URL
                    relative_path = sandbox_url.replace(OpenaiResponsesV2.SANDBOX_PATH_PREFIX, "")
                    sandbox_filename = relative_path.split("/")[-1]
                    
                    # Check if file_obj.filename ends with this sandbox filename
                    # This matches the base filename regardless of prefix
                    if file_obj.filename.endswith(sandbox_filename):
                        # Extract subfolder if exists
                        if "/" in relative_path:
                            subfolder_path = "/".join(relative_path.split("/")[:-1]) + "/"
                        break
            
            # Construct both URLs upfront
            constructed_url_with_subfolder = f"{OpenaiResponsesV2.SANDBOX_PATH_PREFIX}{subfolder_path}{file_obj.id}-{file_obj.filename}" if subfolder_path else None
            constructed_url_without_subfolder = f"{OpenaiResponsesV2.SANDBOX_PATH_PREFIX}{file_obj.id}-{file_obj.filename}"
            
            # First try with subfolder (if exists)
            if constructed_url_with_subfolder and constructed_url_with_subfolder.lower() in output_item.content[0].text.lower():
                pattern = re.escape(constructed_url_with_subfolder)
                output_item.content[0].text = re.sub(
                    pattern,
                    s3_url,
                    output_item.content[0].text,
                    flags=re.IGNORECASE
                )
                logger.bind(
                    ctx=ctx,
                    file_id=file_obj.id,
                    filename=file_obj.filename,
                    subfolder_path=subfolder_path,
                    constructed_sandbox_url=constructed_url_with_subfolder
                ).info(
                    "Replaced constructed sandbox URL (with subfolder) with S3 URL (case-insensitive)"
                )
            # Then try without subfolder (always check this)
            elif constructed_url_without_subfolder.lower() in output_item.content[0].text.lower():
                pattern = re.escape(constructed_url_without_subfolder)
                output_item.content[0].text = re.sub(
                    pattern,
                    s3_url,
                    output_item.content[0].text,
                    flags=re.IGNORECASE
                )
                logger.bind(
                    ctx=ctx,
                    file_id=file_obj.id,
                    filename=file_obj.filename,
                    constructed_sandbox_url=constructed_url_without_subfolder
                ).info(
                    "Replaced constructed sandbox URL (without subfolder) with S3 URL (case-insensitive)"
                )
            else:
                continue  # No match found, move to next file

        # Check if sandbox path prefix still remains after replacement (e.g. file not in mapping)
        has_sandbox_path = (
            output_item.content
            and output_item.content[0].text
            and OpenaiResponsesV2.SANDBOX_PATH_PREFIX in output_item.content[0].text
        )
        return bool(has_sandbox_path)

    async def update_new_public_s3_urls_with_openai_files(
        self,
        ctx: Context,
        client: AsyncAzureOpenAI | ChatCompletionsClient | AsyncOpenAI,
        responses_req: LlmResponsesOpenaiV2Config,
        user: UserORM,
        db: AsyncIOMotorDatabase,
        user_files: UserFiles | None = None,
        additional_file_ids: List[str] = [],
        notify: Notify | None = None,
    ):
        # upload files to openai
        urls_to_upload = []
        urls_from_config = [
            url
            for url in responses_req.file_urls
            if url not in responses_req.file_url_to_file_object_map
        ]
        # V2: Only include sandbox paths (code interpreter generated files), not other tool output URLs
        urls_from_llm = [
            url
            for sandbox_url, url in responses_req.sandbox_url_to_s3_url.items()
            if url not in responses_req.file_url_to_file_object_map
            and (
                sandbox_url.startswith(OpenaiResponsesV2.SANDBOX_PATH_PREFIX)
                or sandbox_url.startswith(OpenaiResponsesV2.MNT_DATA_PATH_PREFIX)
            )
        ]
        urls_to_upload = urls_from_config + urls_from_llm
        file_url_to_file_object_map: Dict[str, FileObject] = await self.upload_urls_to_openai(
            ctx=ctx,
            user=user,
            db=db,
            urls=urls_to_upload,
            openai_client=client,
            supported_file_formats=OPENAI_FILE_FORMATS + OPENAI_FILE_FORMATS_CODE_INTERPRETER,
        )
        for file_url, file_object in file_url_to_file_object_map.items():
            responses_req.file_url_to_file_object_map[file_url] = file_object
        # file ids to use
        file_ids = []
        for file_url, file_object in responses_req.file_url_to_file_object_map.items():
            file_ids.append(file_object.id)
        # Merge with additional file_ids from tool outputs (already uploaded to OpenAI)
        logger.bind(
            ctx=ctx,
            uploaded_file_ids=len(file_ids),
            additional_file_ids=len(additional_file_ids),
            additional_file_ids_list=additional_file_ids,
        ).info("Merging file_ids from tool outputs with uploaded file_ids")

        # Merge with additional file_ids while maintaining order
        for fid in additional_file_ids:
            if fid not in file_ids:
                file_ids.append(fid)

        logger.bind(ctx=ctx, total_file_ids=len(file_ids), merged_file_ids=file_ids).info(
            "Final merged file_ids for tool parameters"
        )

        # handle openai code_interpreter
        await self.build_responses_code_interpreter_tool_param(
            ctx=ctx, responses_req=responses_req, file_ids=file_ids, client=client,
            user=user, db=db,
        )

        # handle openai file_search
        await self.build_responses_file_search_tool_param(
            ctx=ctx, responses_req=responses_req, openai_client=client, file_ids=file_ids
        )

    async def upload_urls_to_openai(
        self,
        ctx: Context,
        user: UserORM,
        db: AsyncIOMotorDatabase,
        openai_client: AsyncAzureOpenAI | ChatCompletionsClient | AsyncOpenAI,
        urls: List[str],
        supported_file_formats: List[str],
        cache: Dict[str, FileObject] | None = None,  # update in-place
    ) -> Dict[str, FileObject]:
        if cache is None:
            cache = {}
        urls_to_tmpfiles: Dict[str, str] = {}
        try:
            new_urls = [url for url in urls if url not in cache]
            # supported_file_formats: List[str] = OPENAI_FILE_FORMATS
            urls_to_tmpfiles: Dict[str, str] = await FileDownloader._download_urls_to_tmpfiles(
                ctx=ctx,
                user=user,
                db=db,
                urls=new_urls,
                supported_file_formats=supported_file_formats,
            )
            logger.bind(ctx=ctx, downloaded_urls=urls_to_tmpfiles).info(
                "Downloaded urls to temp files"
            )
            urls_to_openai_file_obj: Dict[str, FileObject] = await self._upload_files_to_openai(
                ctx=ctx,
                user=user,
                db=db,
                file_urls_to_paths=urls_to_tmpfiles,
                openai_client=openai_client,
                cache=cache,
            )
            logger.bind(ctx=ctx, uploaded_urls=urls_to_openai_file_obj).info(
                "Uploaded urls to openai"
            )
            return urls_to_openai_file_obj
        finally:
            for url, tmpfile in urls_to_tmpfiles.items():
                if os.path.exists(tmpfile):
                    os.remove(tmpfile)
                    logger.bind(ctx=ctx, url=url).info("Removed temp file")

    # async def _download_urls_to_tmpfiles(
    #         self,
    #         ctx: Context,
    #         user: UserORM,
    #         db: AsyncIOMotorDatabase,
    #         urls: List[str],
    #         supported_file_formats: List[str] | None = None,
    #         cache: Dict[str, str] | None = None,  # url -> file_path # update cache in-place
    # ) -> Dict[str, str]:  # url -> file_path
    #     if cache is None:
    #         cache = {}
    #     result: Dict[str, str] = {}
    #     for url in urls:
    #         if url in cache:
    #             continue
    #         try:
    #             file_bytes: bytes | None = None
    #             file_name: str = await UserFiles.extract_file_name(ctx=ctx, file_url=url)
    #             file_ext: str = await UserFiles.extract_file_ext(ctx=ctx, file_url=url)
    #             if url.__contains__(".s3.") or url.startswith("s3://"):
    #                 s3_key = ""
    #                 s3: AwsS3 | None = None
    #                 if url.__contains__(SettingsProvider.sget().AWS_S3_BUCKET_NAME):
    #                     s3 = get_s3()
    #                 elif url.__contains__(SettingsProvider.sget().AWS_S3_PUBLIC_BUCKET_NAME):
    #                     s3 = get_public_s3()
    #                 else:
    #                     app_creds_doc = await UserAppCreds.get_app_creds_by_app_name(
    #                         ctx=Context(), db=db, user=user, app_name=AppTypes.aws
    #                     )
    #                     user_bucket_name = app_creds_doc.creds["bucket_name"]
    #                     assert user_bucket_name in url
    #                     s3 = get_aws_s3(
    #                         region_name=app_creds_doc.creds["region_name"],
    #                         aws_access_key_id=app_creds_doc.creds["aws_access_key_id"],
    #                         aws_secret_access_key=app_creds_doc.creds["aws_secret_access_key"],
    #                         bucket_name=app_creds_doc.creds["bucket_name"],
    #                         object_prefix=app_creds_doc.creds["object_prefix"],
    #                     )
    #                 assert s3 is not None
    #                 if url.__contains__(".s3."):
    #                     s3_key = await s3.get_key_from_https_url(url=url)
    #                 if url.startswith("s3://"):
    #                     s3_key = await s3.get_key_from_s3_uri(url=url)
    #                 assert s3_key
    #                 file_bytes = await s3.get_bytes(ctx=ctx, key=s3_key)
    #             elif url.startswith("http://") or url.startswith("https://"):
    #                 # special handling for Box links
    #                 if "box" in url and "?" in url:
    #                     base_url = url.split("?")[0]
    #                     redirect = url.split("?")[-1]
    #                     url = f"{base_url}?{redirect}"
    #                 async with httpx.AsyncClient(follow_redirects=True) as http_client:
    #                     response = await http_client.get(url)
    #                     if response.status_code != 200:
    #                         logger.bind(ctx=ctx, url=url).error(
    #                             f"openai_responses.upload_urls_to_openai: Error in getting file from url: {url}"
    #                         )
    #                         assert response.status_code == 200
    #                     file_bytes = response.content
    #             assert file_bytes
    #
    #             # 2. if it's a .gsheet, parse JSON and re-download as CSV
    #             if file_ext == ".gsheet":
    #                 logger.bind(ctx=ctx).debug("downloading for google sheet sheet")
    #                 meta = json.loads(file_bytes.decode("utf-8"))
    #                 sheet_url = meta.get("url")
    #                 if not sheet_url:
    #                     raise HTTPException(
    #                         status_code=400, detail='Invalid .gsheet: missing "url" field'
    #                     )
    #                 export_url = sheet_url.rstrip("/edit") + "/export?format=csv"
    #                 async with httpx.AsyncClient(follow_redirects=True) as session:
    #                     resp2 = await session.get(export_url)
    #                     if resp2.status_code != 200:
    #                         raise HTTPException(
    #                             status_code=resp2.status_code,
    #                             detail=f"Failed to export Google Sheet: HTTP {resp2.status_code}",
    #                         )
    #                     file_bytes = resp2.content
    #                 file_ext = ".csv"
    #
    #             tmp = tempfile.NamedTemporaryFile(delete=False, suffix=file_ext)
    #             tmp.write(file_bytes)
    #             logger.bind(ctx=ctx, tmpfile=tmp).debug("written data to temp file")
    #             tmp.close()
    #
    #             cache[url] = tmp.name
    #             result[url] = tmp.name
    #
    #             # supported_file_formats can be like [".pdf"], or it can be like ["pdf"] or a mix
    #             # file_ext is always of type ".pdf" (instead of "pdf")
    #             if (
    #                     supported_file_formats
    #                     and file_ext not in supported_file_formats
    #                     and file_ext[1:] not in supported_file_formats
    #                     and (".zip" in supported_file_formats or "zip" in supported_file_formats)
    #             ):
    #                 tmp_zip = tempfile.NamedTemporaryFile(delete=False, suffix=".zip")
    #                 with zipfile.ZipFile(tmp_zip.name, "w", zipfile.ZIP_DEFLATED) as zipf:
    #                     zipf.write(tmp.name, arcname=file_name)
    #                     logger.bind(ctx=ctx, url=url, tmpfile=tmp_zip).info(
    #                         "zipped data to temp file"
    #                     )
    #                 tmp_zip.close()
    #
    #                 cache[url] = tmp_zip.name
    #                 result[url] = tmp_zip.name
    #                 # remove the original file as no longer needed
    #                 if os.path.exists(tmp.name):
    #                     os.remove(tmp.name)
    #
    #             elif (
    #                     supported_file_formats
    #                     and file_ext not in supported_file_formats
    #                     and file_ext[1:] not in supported_file_formats
    #             ):
    #                 file_str = self._convert_file_to_text(file_path=tmp.name)
    #                 new_tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".txt")
    #                 new_tmp.write(file_str.encode("utf-8"))
    #                 logger.bind(ctx=ctx, url=url, tmpfile=new_tmp.name).debug(
    #                     "written data to temp file txt"
    #                 )
    #                 new_tmp.close()
    #
    #                 cache[url] = new_tmp.name
    #                 result[url] = new_tmp.name
    #                 # remove the original file as no longer needed
    #                 if os.path.exists(tmp.name):
    #                     os.remove(tmp.name)
    #
    #         except Exception as e:
    #             logger.bind(ctx=ctx, url=url).exception(
    #                 f"Error in downloading url to tmpfile: {str(e)}"
    #             )
    #
    #     logger.bind(ctx=ctx, downloaded_urls=result).info("Downloaded urls to temp files")
    #     return result

    async def _upload_files_to_openai(
        self,
        ctx: Context,
        user: UserORM,
        db: AsyncIOMotorDatabase,
        file_urls_to_paths: Dict[str, str],
        openai_client: AsyncAzureOpenAI | ChatCompletionsClient | AsyncOpenAI,
        cache: Dict[str, FileObject] | None = None,  # update cache in-place
    ) -> Dict[str, FileObject]:
        if cache is None:
            cache = {}
        result: Dict[str, FileObject] = {}
        openai_files = OpenaiFiles(openai_client=openai_client)
        for url, path in file_urls_to_paths.items():
            if url in cache:
                continue
            # Skip empty files (0 bytes) - OpenAI rejects them with "File is empty" error
            file_size = os.path.getsize(path)
            if file_size == 0:
                logger.bind(ctx=ctx, url=url, path=path).warning(
                    "Skipping file upload to OpenAI: file is empty (0 bytes)"
                )
                continue
            # Skip files larger than 510MB (OpenAI limit is 512MB)
            if file_size > self.MAX_OPENAI_FILE_SIZE_BYTES:
                logger.bind(
                    ctx=ctx,
                    url=url,
                    file_size_mb=round(file_size / (1024 * 1024), 2),
                ).warning("Skipping file upload to OpenAI: file exceeds 510MB limit")
                continue
            try:
                path_ext = await UserFiles.extract_file_ext(ctx=ctx, file_url=path)
                url_ext = await UserFiles.extract_file_ext(ctx=ctx, file_url=url)
                file_name: str = await UserFiles.extract_file_name(ctx=ctx, file_url=url)
                if path_ext != url_ext:
                    file_name = file_name.replace(".", "_")
                    file_name = file_name + path_ext

                file_tuple = (file_name, Path(path).read_bytes())
                file_create = FileCreate(
                    file=file_tuple,
                    purpose="assistants",
                )

                upload_t0 = time.perf_counter()
                logger.bind(
                    ctx=ctx,
                    url=url,
                    path=path,
                    file_size_bytes=file_size,
                    openai_filename=file_name,
                ).info("Uploading file to OpenAI")
                file_object: FileObject = await openai_files.create(
                    ctx=ctx, file_create=file_create
                )
                upload_ms = int((time.perf_counter() - upload_t0) * 1000)
                logger.bind(
                    ctx=ctx,
                    url=url,
                    file_size_bytes=file_size,
                    upload_ms=upload_ms,
                    openai_file_id=getattr(file_object, "id", None),
                    openai_status=getattr(file_object, "status", None),
                ).info("Uploaded file to OpenAI")
                cache[url] = file_object
                result[url] = file_object
            except Exception as e:
                logger.bind(ctx=ctx, url=url, path=path).exception(
                    f"Error in uploading tmpfiles to openai: {str(e)}"
                )
        return result

    async def build_responses_file_search_tool_param(
        self,
        ctx: Context,
        responses_req: LlmResponsesOpenaiV2Config,
        openai_client: AsyncAzureOpenAI | AsyncOpenAI,
        file_ids: List[str] = list,
    ) -> None:
        file_batch_timeout_seconds = 600

        try:
            tool_list = [] if responses_req.tools is None else responses_req.tools
            sanitized_tool_list: list | None = []
            file_search_index = -1
            file_search_tool_param: FileSearchToolParam | None= None
            for i, tool in enumerate(tool_list):
                if (isinstance(tool, str) and tool == "file_search"):
                    if file_search_index == -1:
                        # create new file_search_tool_param
                        file_search_tool_param = FileSearchToolParam(
                            type="file_search",
                            vector_store_ids=[],
                            max_num_results=10,
                            ranking_options=RankingOptions(ranker="auto", score_threshold=0.5),
                        )
                        file_search_index = i
                        sanitized_tool_list.append(file_search_tool_param)
                elif isinstance(tool, dict) and "type" in tool and tool["type"] == "file_search":
                    if file_search_index == -1:
                        file_search_index = i
                        file_search_tool_param = FileSearchToolParam(**tool)
                        sanitized_tool_list.append(file_search_tool_param)
                else:
                    sanitized_tool_list.append(tool)
            responses_req.tools = sanitized_tool_list if sanitized_tool_list else None

            if file_search_tool_param is None:
                return

            vector_store_ids = file_search_tool_param["vector_store_ids"]
            vector_store_client = OpenaiVectorStores(openai_client)
            vector_store_id = None

            if vector_store_ids and len(vector_store_ids) > 0:
                try:
                    vs = await vector_store_client.retrieve(
                        RetrieveVectorStore(vector_store_id=vector_store_ids[0])
                    )
                    if vs.status != "expired":
                        vector_store_id = vector_store_ids[0]
                        logger.bind(ctx=ctx, vs_id=vector_store_id, vs_status=vs.status).info(
                            "Using existing vector store"
                        )
                    else:
                        logger.bind(ctx=ctx, vs_id=vector_store_ids[0]).warning(
                            "Vector store expired, will create new"
                        )
                        raise Exception("Vector store expired")
                except Exception as e:
                    logger.bind(ctx=ctx, vs_id=vector_store_ids[0], error=str(e)).warning(
                        "Vector store invalid or not found, clearing state"
                    )
                    responses_req.file_id_to_vector_store_map.clear()
                    file_search_tool_param["vector_store_ids"] = []

            if vector_store_id is None:
                settings = SettingsProvider.sget()
                create_vector_store = CreateVectorStore(
                    expires_after={
                        "anchor": "last_active_at",
                        "days": settings.OPENAI_VECTOR_STORE_EXPIRES_AFTER_DAYS,
                    }
                )
                new_vector_store = await vector_store_client.create(
                    create_vector_store=create_vector_store,
                )
                vector_store_id = new_vector_store.id
                file_search_tool_param["vector_store_ids"] = [vector_store_id]
                logger.bind(ctx=ctx, vs_id=vector_store_id, expires_after_days=settings.OPENAI_VECTOR_STORE_EXPIRES_AFTER_DAYS).info(
                    "Created new vector store with expiration policy"
                )

            new_file_ids = []
            for file_id in file_ids:
                for file_url, file_object in responses_req.file_url_to_file_object_map.items():
                    if file_id == file_object.id:
                        ext = "." + file_object.filename.split(".")[-1]
                        ext = ext.lower()
                        if ext in OPENAI_FILE_FORMATS:
                            if file_id not in responses_req.file_id_to_vector_store_map:
                                new_file_ids.append(file_id)
                        else:
                            logger.bind(
                                ctx=ctx, filename=file_object.filename, file_object=file_object
                            ).info("openai_responses.file_search file_ext not in FILE_FORMATS")
                            continue

            logger.bind(ctx=ctx, file_ids=new_file_ids).info(
                "openai_responses.file_search new file_ids"
            )

            if len(new_file_ids) == 0:
                return

            async with asyncio.timeout(file_batch_timeout_seconds):
                await openai_client.vector_stores.file_batches.create_and_poll(
                    vector_store_id=vector_store_id, file_ids=new_file_ids
                )
            logger.bind(ctx=ctx, vector_store_id=vector_store_id).info(
                "Added files to vector store"
            )
            vector_store = await openai_client.vector_stores.retrieve(
                vector_store_id=vector_store_id
            )
            logger.bind(ctx=ctx, vector_store=vector_store).info(
                "openai_responses.create vector_store"
            )
            for file_id in new_file_ids:
                responses_req.file_id_to_vector_store_map[file_id] = VectorStoreLite(
                    **vector_store.model_dump()
                )
        except Exception as e:
            if isinstance(e, TimeoutError):
                logger.bind(
                    ctx=ctx,
                    error=str(e),
                    error_type=type(e).__name__,
                    file_ids=file_ids,
                    timeout_seconds=file_batch_timeout_seconds
                ).warning(
                    "Timeout waiting for vector store file batch processing. "
                    "Continuing without file_search to let the action proceed and continue."
                )
            else:
                logger.bind(
                    ctx=ctx,
                    error=str(e),
                    error_type=type(e).__name__,
                    file_ids=file_ids
                ).exception(
                    "Failed to build file_search tool param due to OpenAI API error. "
                    "Continuing without file_search to preserve tool input/output pair."
                )
            # Remove file_search tool with empty vector_store_ids to prevent OpenAI 400 error
            if responses_req.tools:
                responses_req.tools = [
                    t for t in responses_req.tools
                    if not (isinstance(t, dict) and t.get("type") == "file_search" and not t.get("vector_store_ids"))
                ]
            return


    async def _consolidate_file_ids_to_zip(
        self,
        ctx: Context,
        file_ids: List[str],
        client: AsyncAzureOpenAI | AsyncOpenAI,
        file_url_to_file_object_map: Dict[str, FileObject],
        user: UserORM,
        db: AsyncIOMotorDatabase,
        zip_index: int = 1,
    ) -> tuple[List[str], Set[str]]:
        """
        When file_ids exceed the threshold, download files from their original URLs,
        zip them into a single archive, upload the zip, and return its file_id.
        Returns: (new_file_ids, set of original file_ids that were zipped)
        """
        zip_path = None
        downloaded_tmpfiles: List[str] = []
        try:
            logger.bind(ctx=ctx, file_count=len(file_ids), zip_index=zip_index).info(
                "Starting file_ids consolidation into zip"
            )

            # Build file_id -> (url, filename) reverse lookup from the map
            file_id_to_url: Dict[str, str] = {}
            file_id_to_filename: Dict[str, str] = {}
            for url, file_obj in file_url_to_file_object_map.items():
                if file_obj.id in file_ids:
                    file_id_to_url[file_obj.id] = url
                    file_id_to_filename[file_obj.id] = file_obj.filename

            # For file_ids not in map (e.g. tool-generated), retrieve metadata directly
            for fid in file_ids:
                if fid not in file_id_to_filename:
                    try:
                        file_obj = await client.files.retrieve(file_id=fid)
                        file_id_to_filename[fid] = file_obj.filename if file_obj and file_obj.filename else fid
                    except Exception as e:
                        logger.bind(ctx=ctx, file_id=fid).warning(
                            f"Could not retrieve filename for file_id, using id as name: {str(e)}"
                        )
                        file_id_to_filename[fid] = fid

            # Collect URLs to download, normalizing whitespace from Pydantic HttpUrl serialization
            urls_to_download: List[str] = []
            normalized_url_to_fid: Dict[str, str] = {}
            for fid in file_ids:
                source_url = file_id_to_url.get(fid)
                if not source_url:
                    logger.bind(ctx=ctx, file_id=fid).warning(
                        "No source URL found for file_id, skipping"
                    )
                    continue
                normalized = "".join(str(source_url).split())
                urls_to_download.append(normalized)
                normalized_url_to_fid[normalized] = fid

            # Download all files using FileDownloader (handles httpx + S3 boto3 fallback)
            url_to_tmpfile: Dict[str, str] = await FileDownloader._download_urls_to_tmpfiles(
                ctx=ctx,
                user=user,
                db=db,
                urls=urls_to_download,
            )
            downloaded_tmpfiles = list(url_to_tmpfile.values())

            # Create temp zip and add downloaded files
            tmp_zip = tempfile.NamedTemporaryFile(delete=False, suffix=".zip")
            zip_path = tmp_zip.name
            tmp_zip.close()

            seen_filenames: Dict[str, int] = {}
            zipped_fids: Set[str] = set()
            zip_build_t0 = time.perf_counter()
            with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zipf:
                for url, tmpfile_path in url_to_tmpfile.items():
                    fid = normalized_url_to_fid.get(url)
                    if not fid:
                        continue
                    try:
                        # Handle filename collisions
                        filename = file_id_to_filename.get(fid, fid)
                        if filename in seen_filenames:
                            seen_filenames[filename] += 1
                            name, ext = os.path.splitext(filename)
                            filename = f"{name}_{seen_filenames[filename]}{ext}"
                        else:
                            seen_filenames[filename] = 0

                        zipf.write(tmpfile_path, arcname=filename)
                        zipped_fids.add(fid)
                    except Exception as e:
                        logger.bind(ctx=ctx, file_id=fid, source_url=url).exception(
                            f"Failed to read downloaded file for zipping: {str(e)}"
                        )

            zip_build_ms = int((time.perf_counter() - zip_build_t0) * 1000)
            logger.bind(
                ctx=ctx,
                zip_path=zip_path,
                file_count=len(seen_filenames),
                zip_build_ms=zip_build_ms,
            ).info("Built consolidated zip")

            # Guard: if no files were zipped, fall back to truncation
            if not seen_filenames:
                logger.bind(ctx=ctx).warning(
                    "No files could be downloaded for zipping, falling back to truncation"
                )
                return file_ids[:50], set()

            # Upload the zip to OpenAI (using OpenaiFiles for retry/polling)
            zip_size = os.path.getsize(zip_path)
            logger.bind(ctx=ctx, zip_path=zip_path, file_count=len(seen_filenames), zip_size_bytes=zip_size).info(
                "Uploading consolidated zip to OpenAI"
            )
            openai_files = OpenaiFiles(openai_client=client)
            zip_upload_t0 = time.perf_counter()
            zip_file_obj = await openai_files.create(
                ctx=ctx,
                file_create=FileCreate(
                    file=(f"consolidated_files_{zip_index}.zip", Path(zip_path).read_bytes()),
                    purpose="assistants",
                ),
            )
            zip_upload_ms = int((time.perf_counter() - zip_upload_t0) * 1000)

            logger.bind(
                ctx=ctx,
                zip_file_id=zip_file_obj.id,
                original_count=len(file_ids),
                zip_size_bytes=zip_size,
                zip_upload_ms=zip_upload_ms,
            ).info(
                "Consolidated file_ids into single zip"
            )
            return [zip_file_obj.id], zipped_fids

        except Exception as e:
            # Fallback: truncate to 50 file_ids instead of failing
            logger.bind(ctx=ctx).exception(
                f"Failed to consolidate file_ids to zip, falling back to truncation: {str(e)}"
            )
            return file_ids[:50], set()

        finally:
            if zip_path and os.path.exists(zip_path):
                os.remove(zip_path)
            for tmpfile in downloaded_tmpfiles:
                if os.path.exists(tmpfile):
                    os.remove(tmpfile)

    async def build_responses_code_interpreter_tool_param(
        self,
        ctx: Context,
        responses_req: LlmResponsesOpenaiV2Config,
        file_ids: List[str],
        client: AsyncAzureOpenAI | AsyncOpenAI | None = None,
        user: UserORM | None = None,
        db: AsyncIOMotorDatabase | None = None,
    ) -> None:
        try:
            new_file_ids = []
            tool_list = [] if responses_req.tools is None else responses_req.tools
            for i, tool in enumerate(tool_list):
                if (
                    isinstance(tool, dict)
                    and "type" in tool
                    and tool["type"] == "code_interpreter"
                    and "container" in tool
                    and "file_ids" in tool["container"]
                ):
                    new_file_ids += tool["container"]["file_ids"]
            new_file_ids += file_ids
            new_file_ids = list(set(new_file_ids))

            # Filter out file_ids already consolidated into a zip (prevents re-zipping every loop)
            new_file_ids = [fid for fid in new_file_ids if fid not in responses_req.consolidated_file_ids]

            # Separate: our consolidation zips vs regular files (never re-zip our zips)
            zip_file_ids = [fid for fid in new_file_ids if fid in responses_req.consolidation_zip_ids]
            regular_file_ids = [fid for fid in new_file_ids if fid not in responses_req.consolidation_zip_ids]

            logger.bind(
                ctx=ctx,
                total_file_ids=len(new_file_ids),
                zip_file_ids_count=len(zip_file_ids),
                regular_file_ids_count=len(regular_file_ids),
                consolidated_file_ids_count=len(responses_req.consolidated_file_ids),
                consolidation_zip_ids_count=len(responses_req.consolidation_zip_ids),
                consolidation_zip_ids_contents=responses_req.consolidation_zip_ids,
                consolidation_zip_ids_type=str(type(responses_req.consolidation_zip_ids)),
            ).info("openai_responses.code_interpreter file_ids before consolidation check")

            # Consolidate regular files if total exceeds threshold
            if client and (len(zip_file_ids) + len(regular_file_ids)) > self.FILE_IDS_CONSOLIDATION_THRESHOLD:
                zip_index = len(responses_req.consolidation_zip_ids) + 1
                consolidated, zipped_ids = await self._consolidate_file_ids_to_zip(
                    ctx=ctx,
                    file_ids=regular_file_ids,
                    client=client,
                    file_url_to_file_object_map=responses_req.file_url_to_file_object_map,
                    user=user,
                    db=db,
                    zip_index=zip_index,
                )
                if zipped_ids:
                    # Consolidation succeeded — track the originals and the new zip
                    responses_req.consolidated_file_ids.extend(fid for fid in zipped_ids if fid not in responses_req.consolidated_file_ids)
                    responses_req.consolidation_zip_ids.extend(fid for fid in consolidated if fid not in responses_req.consolidation_zip_ids)
                    new_file_ids = zip_file_ids + consolidated
                else:
                    # Consolidation failed — fallback returned truncated list, don't track
                    logger.bind(ctx=ctx).warning(
                        "Consolidation failed, using truncated file_ids without tracking"
                    )
                    new_file_ids = zip_file_ids + consolidated
            else:
                new_file_ids = zip_file_ids + regular_file_ids

            logger.bind(ctx=ctx, file_ids=new_file_ids).info(
                "openai_responses.code_interpreter file_ids"
            )

            memory_limit = responses_req.container_memory_limit if responses_req.container_memory_limit is not None else "1g"

            logger.bind(ctx=ctx, memory_limit=memory_limit).info(
                "Setting openai_responses.code_interpreter memory_limit"
            )

            new_code_interpreter = CodeInterpreter(
                type="code_interpreter",
                container=CodeInterpreterContainerCodeInterpreterToolAuto(
                    type="auto", file_ids=new_file_ids, memory_limit=memory_limit
                ),
            )

            sanitized_tool_list: list | None = []
            code_interpreter_index = -1
            for i, tool in enumerate(tool_list):
                if (isinstance(tool, str) and tool == "code_interpreter") or (
                    isinstance(tool, dict) and "type" in tool and tool["type"] == "code_interpreter"
                ):
                    if code_interpreter_index == -1:
                        code_interpreter_index = i
                        sanitized_tool_list.append(new_code_interpreter)
                else:
                    sanitized_tool_list.append(tool)

            responses_req.tools = sanitized_tool_list if sanitized_tool_list else None
        except Exception as e:
            logger.bind(
                ctx=ctx,
                error=str(e),
                error_type=type(e).__name__,
                file_ids=file_ids
            ).exception(
                "Failed to build code_interpreter tool param due to OpenAI API error. "
                "Continuing without code_interpreter to preserve tool input/output pair."
            )
            return

    def _convert_file_to_text(
        self,
        file_path: str,
        sheet_name: Optional[str] = None,
        text_format: str = "plain",  # "plain" for tab-delimited, or "jsonl"
    ) -> str:
        """
        Load the downloaded file and convert it to text:
          - tab-delimited or JSON-lines for CSV/Excel
          - plain text with paragraph and table extraction for Word
          - plain text for PowerPoint
        """
        suffix = Path(file_path).suffix.lower()

        # CSV or Excel
        if suffix in [".csv", ".xls", ".xlsx"]:
            df = None
            if suffix == ".csv":
                df = pd.read_csv(file_path)
            # Excel
            elif suffix in [".xls", ".xlsx"]:
                if sheet_name:
                    df = pd.read_excel(file_path, sheet_name=sheet_name)
                else:
                    df = pd.read_excel(file_path)
                # Handle multiple sheets dict
                if isinstance(df, dict):
                    df = pd.concat(
                        [sheet.assign(_sheet=name) for name, sheet in df.items()], ignore_index=True
                    )
            assert df is not None
            if text_format == "jsonl":
                return df.to_json(orient="records", lines=True)
            elif text_format == "plain":
                return df.to_csv(sep="\t", index=False)
            else:
                raise HTTPException(
                    status_code=420, detail=f"Unsupported text format: {text_format}"
                )
        # Word docs
        elif suffix in {".docx", ".doc"}:
            doc = Document(file_path)
            texts = []
            # extract paragraphs
            for p in doc.paragraphs:
                if p.text and p.text.strip():
                    texts.append(p.text.strip())
            # extract tables
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    texts.append("\t".join(cells))
            return "\n".join(texts)
        # PowerPoint
        elif suffix in {".pptx", ".ppt"}:
            prs = Presentation(file_path)
            slides_text = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text and shape.text.strip():
                        slides_text.append(shape.text.strip())
            return "\n---\n".join(slides_text)
        else:
            raise HTTPException(status_code=420, detail=f"Unsupported file type: {suffix}")
