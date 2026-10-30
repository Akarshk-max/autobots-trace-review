# src/autobots/ingestion/service.py

import asyncio
import base64
import hashlib
import io
import json
import uuid
from pathlib import Path
from typing import Any, Dict, Iterable, List, Literal, Optional, Set, Tuple
from urllib.parse import unquote

from src.autobots.data_model.app_state import AppState
from loguru import logger
from motor.motor_asyncio import AsyncIOMotorDatabase

from src.autobots.integrations.app_conn.app_conn import AppConn
from src.autobots.conn.chroma.chroma import Document
from src.autobots.conn.object_storage.factory import get_public_storage
from src.autobots.conn.object_storage.protocol import ObjectStorage
from src.autobots.conn.unstructured_io.unstructured_io import (
    PartitionParametersParams,
    UnstructuredIO,
)
from src.autobots.conn.weaviate.weaviate_store import WeaviateStore
from src.autobots.core.settings import SettingsProvider
from src.autobots.data_model.context import Context
from src.autobots.ingestion.exceptions import RetryableIngestionError
from src.autobots.ingestion.graphrag_processor import GraphRAGProcessor
from src.autobots.ingestion.job_status import JobStatusCreate, JobStatusCRUD
from src.autobots.ingestion.media_utils import (
    pre_parse_document_with_gemini,
    summarize_media_content,
)
from src.autobots.ingestion.neo4j_loader import Neo4jLoader
from src.autobots.ingestion.spreadsheet.locator import parse_source_locator
from src.autobots.ingestion.weaviate_processor import WeaviateLoader, WeaviateProcessor
from src.autobots.user.user_orm_model import UserORM

# Optional: LangChain language-aware splitters (fallback to internal splitter if missing)
try:
    from langchain_text_splitters import Language, RecursiveCharacterTextSplitter
except Exception:  # pragma: no cover
    Language = None
    RecursiveCharacterTextSplitter = None

# Optional: openpyxl for XLSM parsing
try:
    import openpyxl
except Exception:  # pragma: no cover
    openpyxl = None
from src.autobots.core.redis.async_pool import get_async_redis

# --- MODULE-LEVEL CONCURRENCY GUARD for Gemini Document Parsing ---
# This semaphore is shared across all IngestionService instances, correctly
# throttling concurrent calls to the Gemini API across all jobs in the worker process.
GEMINI_DOC_PARSE_SEMAPHORE = asyncio.Semaphore(4)

"""
Format Policy
-------------
We keep a single readable “checklist” policy organized by:
  (A) support-status group: SUPPORTED / PLANNED / BLOCKED / UNKNOWN_BEST_EFFORT
  (B) category: document/spreadsheet/presentation/web/structured/text_like/image/audio/video/archive/binary

Routing rules:
  - special handlers (SPECIAL_EXTENSION_HANDLERS): extension-specific custom ingestion (e.g., .json, .xlsm); NOT Unstructured
  - blocked: fail-fast (archives/binaries/executables)
  - media: handled via existing media pipeline
  - text_like: custom text-like ingestion (code/config extensions); NOT Unstructured
  - doc-like formats: Unstructured document pipeline (e.g., pdf/docx/html/pptx/xml)
  - planned: documentation-only (not fail-fast)
  - unknown: try Unstructured first; if empty/raises (and file looks like text), fall back to text-like ingestion

IMPORTANT: We do not change Unstructured code or its call flow.
"""

CATEGORIES: Tuple[str, ...] = (
    "document",
    "spreadsheet",
    "presentation",
    "web",
    "structured",
    "text_like",
    "image",
    "audio",
    "video",
    "archive",
    "binary",
)

# 1) SUPPORTED (works now)
SUPPORTED: Dict[str, List[str]] = {
    # "True documents" => Unstructured
    "document": [".pdf", ".doc", ".docx", ".odt", ".rtf", ".rst", ".txt", ".md", ".epub", ".eml", ".msg", ".p7s"],
    "spreadsheet": [".csv", ".tsv", ".xls", ".xlsx", ".xlsm", ".xlsb"],  # Special handler (SPECIAL_EXTENSION_HANDLERS); NOT Unstructured
    "presentation": [".ppt", ".pptx"],                 # Unstructured doc pipeline
    "web": [".html", ".htm", ".xhtml", ".xml"],        # Unstructured doc pipeline
    "structured": [".json"],                           # Special handler (SPECIAL_EXTENSION_HANDLERS); NOT Unstructured
    # Code/config => custom text-like handler (NOT Unstructured; language-aware splitting when available)
    "text_like": [
        # Python
        ".py", ".pyi",
        # JS / TS
        ".js", ".jsx", ".mjs", ".cjs", ".ts", ".tsx",
        # JVM
        ".java", ".kt", ".kts",
        # Systems / backend
        ".go", ".rs", ".rb", ".php", ".cs",
        # C / C++
        ".c", ".h", ".cc", ".cpp", ".cxx", ".hpp", ".hh", ".hxx",
        # Scripting / misc
        ".sh", ".ps1", ".lua",
        # Data / query
        ".sql",
        # Config
        ".yml", ".yaml", ".toml", ".ini", ".cfg", ".env", ".properties", ".hcl", ".tf",
    ],
    "image": [".png", ".jpg", ".jpeg", ".webp", ".heic", ".heif"],
    "audio": [".mp3", ".wav", ".m4a", ".ogg", ".oga", ".flac", ".aac", ".opus"],
    "video": [".mp4", ".m4v", ".mov", ".avi", ".webm", ".mkv", ".3gp", ".hevc", ".h265"],
    "archive": [],
    "binary": [],
}

# 2) PLANNED (kept for documentation only; we do not fail-fast on planned extensions)
PLANNED: Dict[str, List[str]] = {
    "document": [],
    "spreadsheet": [],
    "presentation": [],
    "web": [],
    "structured": [".jsonl", ".ndjson"],  # optional future
    "text_like": [],
    "image": [],
    "audio": [],
    "video": [],
    "archive": [],
    "binary": [],
}

# 3) BLOCKED (fail-fast: archives + binaries/executables)
BLOCKED: Dict[str, List[str]] = {
    "document": [],
    "spreadsheet": [],
    "presentation": [],
    "web": [],
    "structured": [],
    "text_like": [],
    "image": [],
    "audio": [],
    "video": [],
    "archive": [".zip", ".tar", ".gz", ".tgz", ".7z", ".rar", ".bz2", ".xz"],
    "binary": [
        ".exe", ".dll", ".so", ".dylib", ".bin", ".iso", ".dmg", ".pkg", ".deb", ".rpm",
        # extra common binary/container/data formats worth hard-blocking
        ".jar", ".class", ".pkl", ".pickle", ".sqlite", ".parquet",
    ],
}

# 4) UNKNOWN_BEST_EFFORT (optional “known unknowns”; everything else is also best-effort by default)
UNKNOWN_BEST_EFFORT: Dict[str, List[str]] = {k: [] for k in CATEGORIES}


def _flatten(policy: Dict[str, List[str]]) -> Set[str]:
    out: Set[str] = set()
    for cat in CATEGORIES:
        for ext in policy.get(cat, []):
            out.add(ext.lower())
    return out


SUPPORTED_EXTENSIONS = _flatten(SUPPORTED)
PLANNED_EXTENSIONS = _flatten(PLANNED)
BLOCKED_EXTENSIONS = _flatten(BLOCKED)
BEST_EFFORT_KNOWN_EXTENSIONS = _flatten(UNKNOWN_BEST_EFFORT)

# Special-case handlers by extension. Values are method names on IngestionService.
SPECIAL_EXTENSION_HANDLERS: Dict[str, str] = {
    ".json": "_process_json_from_url",
    # Spreadsheets: dedicated ingestion pipeline (NOT Unstructured)
    ".csv": "_process_spreadsheet_from_url",
    ".tsv": "_process_spreadsheet_from_url",
    ".xls": "_process_spreadsheet_from_url",
    ".xlsx": "_process_spreadsheet_from_url",
    ".xlsm": "_process_spreadsheet_from_url",
    ".xlsb": "_process_spreadsheet_from_url",
}

def _is_spreadsheet_suffix(suffix: str) -> bool:
    """
    Single source of truth for spreadsheet routing:
    if the suffix maps to the spreadsheet handler, it's a spreadsheet.
    This avoids keeping a separate SPREADSHEET_EXTENSIONS set in sync.
    """
    return SPECIAL_EXTENSION_HANDLERS.get(suffix) == "_process_spreadsheet_from_url"

# Media sets (existing pipeline)
IMAGE_EXTENSIONS = set(e.lower() for e in SUPPORTED["image"])
AUDIO_EXTENSIONS = set(e.lower() for e in SUPPORTED["audio"])
VIDEO_EXTENSIONS = set(e.lower() for e in SUPPORTED["video"])

# Code/config sets
TEXT_LIKE_EXTENSIONS = {e.lower() for e in SUPPORTED["text_like"]}

# Doc-like (Unstructured pipeline): ONLY extensions that should go through Unstructured partitioning.
# This is used for BOTH:
#   (1) logging/telemetry for “unknown extension” handling, and
#   (2) controlling fallback behavior (e.g., when Unstructured fails or returns no elements).
# NOTE: Spreadsheets are supported, but they are NOT part of the Unstructured doc-like pipeline anymore.
DOC_LIKE_EXTENSIONS = set(
    (SUPPORTED["document"] + SUPPORTED["presentation"] + SUPPORTED["web"])
)
DOC_LIKE_EXTENSIONS = {e.lower() for e in DOC_LIKE_EXTENSIONS}

# JSON / structured data should not go through the doc pipeline; it is handled via special handler.


_JSON_INJECTION_TOKENS = ("<|", "###", "SYSTEM:", "ASSISTANT:", "USER:", "DEVELOPER:")


def _sanitize_json_text(value: Any) -> str:
    s = value if isinstance(value, str) else str(value)
    out = s
    for tok in _JSON_INJECTION_TOKENS:
        out = out.replace(tok, "")
    return out.strip()


def _iter_json_leaves(obj: Any, path: str = "$") -> Iterable[Tuple[str, Any]]:
    """Yield (json_path, leaf_value) pairs for JSON leaf nodes."""
    if isinstance(obj, dict):
        for k, v in obj.items():
            # keep it path-safe and stable
            k2 = str(k).replace("~", "-").replace("/", "_")
            yield from _iter_json_leaves(v, f"{path}.{k2}")
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            yield from _iter_json_leaves(v, f"{path}[{i}]")
    else:
        yield path, obj

def _decode_text_best_effort(raw: bytes) -> str:
    """
    Best-effort decoder for code/config/unknown text-ish files.
    - Handles UTF-8 BOM
    - Falls back to 'utf-8' replacement to avoid hard failures
    """
    if raw is None:
        return ""
    try:
        return raw.decode("utf-8-sig")
    except Exception:
        return raw.decode("utf-8", errors="replace")


def _looks_like_text(raw: bytes, sample_size: int = 8192) -> bool:
    """
    UTF-8-friendly heuristic:
    - If NUL bytes exist => binary
    - Otherwise accept if "printable-ish" ratio is reasonable
      (treat bytes >= 0x20 as printable, including non-ASCII UTF-8 bytes)
    """
    if not raw:
        return True
    sample = raw[:sample_size]
    if b"\x00" in sample:
        return False
    printable = sum((b >= 32) or (b in (9, 10, 13)) for b in sample)
    ratio = printable / max(1, len(sample))
    return ratio >= 0.70

# Central language map used for:
#  - selecting LangChain language splitter (if installed)
#  - naming doc_type ("code_python", "config_yaml", etc.)
_LANGUAGE_MAP: Dict[str, Dict[str, Optional[str]]] = {
    # Python
    ".py":  {"kind": "code", "name": "python", "langchain": "PYTHON"},
    ".pyi": {"kind": "code", "name": "python", "langchain": "PYTHON"},

    # JavaScript / TypeScript
    ".js":  {"kind": "code", "name": "javascript", "langchain": "JS"},
    ".jsx": {"kind": "code", "name": "javascript", "langchain": "JS"},
    ".mjs": {"kind": "code", "name": "javascript", "langchain": "JS"},
    ".cjs": {"kind": "code", "name": "javascript", "langchain": "JS"},
    ".ts":  {"kind": "code", "name": "typescript", "langchain": "TS"},
    ".tsx": {"kind": "code", "name": "typescript", "langchain": "TS"},

    # Common backend languages (LangChain enum availability varies; getattr() keeps it safe)
    ".java": {"kind": "code", "name": "java", "langchain": "JAVA"},
    ".kt":   {"kind": "code", "name": "kotlin", "langchain": "KOTLIN"},
    ".kts":  {"kind": "code", "name": "kotlin", "langchain": "KOTLIN"},
    ".go":   {"kind": "code", "name": "go", "langchain": "GO"},
    ".rs":   {"kind": "code", "name": "rust", "langchain": "RUST"},
    ".rb":   {"kind": "code", "name": "ruby", "langchain": "RUBY"},
    ".php":  {"kind": "code", "name": "php", "langchain": "PHP"},
    ".cs":   {"kind": "code", "name": "csharp", "langchain": "CSHARP"},

    # C / C++
    ".c":   {"kind": "code", "name": "c", "langchain": "C"},
    ".h":   {"kind": "code", "name": "c", "langchain": "C"},
    ".cc":  {"kind": "code", "name": "cpp", "langchain": "CPP"},
    ".cpp": {"kind": "code", "name": "cpp", "langchain": "CPP"},
    ".cxx": {"kind": "code", "name": "cpp", "langchain": "CPP"},
    ".hpp": {"kind": "code", "name": "cpp", "langchain": "CPP"},
    ".hh":  {"kind": "code", "name": "cpp", "langchain": "CPP"},
    ".hxx": {"kind": "code", "name": "cpp", "langchain": "CPP"},

    # Shell / scripting
    ".sh":  {"kind": "code", "name": "shell", "langchain": "BASH"},
    ".ps1": {"kind": "code", "name": "powershell", "langchain": "POWERSHELL"},
    ".lua": {"kind": "code", "name": "lua", "langchain": "LUA"},

    # SQL
    ".sql": {"kind": "code", "name": "sql", "langchain": "SQL"},

    # Config
    ".yml":  {"kind": "config", "name": "yaml", "langchain": "YAML"},
    ".yaml": {"kind": "config", "name": "yaml", "langchain": "YAML"},
    ".toml": {"kind": "config", "name": "toml", "langchain": "TOML"},
    ".ini":  {"kind": "config", "name": "ini", "langchain": None},
    ".cfg":  {"kind": "config", "name": "cfg", "langchain": None},
    ".env":  {"kind": "config", "name": "env", "langchain": None},
    ".properties": {"kind": "config", "name": "properties", "langchain": None},
    ".hcl": {"kind": "config", "name": "hcl", "langchain": None},
    ".tf":  {"kind": "config", "name": "terraform", "langchain": None},
}


def _split_text_language_aware(
    text: str,
    ext: str,
    chunk_size: int = 2000,
    chunk_overlap: int = 200,
) -> List[str]:
    """
    Use langchain_text_splitters when available; otherwise fall back to deterministic
    newline-aware chunking.
    """
    text = text or ""
    if not text.strip():
        return []

    if RecursiveCharacterTextSplitter and Language:
        lang = None
        info = _LANGUAGE_MAP.get(ext)
        enum_name = info.get("langchain") if info else None
        if enum_name:
            # Safe: if enum doesn't exist in this langchain version, getattr returns None
            lang = getattr(Language, enum_name, None)

        splitter = (
            RecursiveCharacterTextSplitter.from_language(
                language=lang, chunk_size=chunk_size, chunk_overlap=chunk_overlap
            )
            if lang is not None
            else RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=chunk_overlap)
        )
        docs = splitter.create_documents([text])
        return [d.page_content for d in docs if getattr(d, "page_content", None)]

    # Fallback: simple chunk-by-lines
    lines = text.splitlines(keepends=True)
    out: List[str] = []
    buf: List[str] = []
    buf_len = 0

    def flush():
        nonlocal buf, buf_len
        if not buf:
            return
        chunk = "".join(buf).strip()
        if chunk:
            out.append(chunk)
        buf = []
        buf_len = 0

    for ln in lines:
        buf.append(ln)
        buf_len += len(ln)
        if buf_len >= chunk_size:
            flush()
    flush()

    # Add overlap by prefixing with tail of previous chunk
    if chunk_overlap > 0 and len(out) > 1:
        overlapped = [out[0]]
        for i in range(1, len(out)):
            prev = overlapped[-1]
            tail = prev[-chunk_overlap:] if len(prev) > chunk_overlap else prev
            overlapped.append((tail + "\n" + out[i]).strip())
        out = overlapped

    return out


class IngestionService:
    def __init__(
        self,
        user_id: str,
        tenant_id: Optional[str] = None,
        analyze_tokens: bool = False,
    ):
        if not user_id:
            raise ValueError("IngestionService requires a user_id.")
        self.user_id = user_id
        self.tenant_id = tenant_id or user_id  # Effective tenant for Weaviate/Neo4j
        self.analyze_tokens = analyze_tokens

        # Legacy attribute name retained for compatibility; the value is provider-neutral.
        self.s3_client: ObjectStorage = get_public_storage(object_prefix="")

        # The processors and loaders are now instantiated just-in-time within the
        # specific ingestion methods, after we determine the correct strategy.
        # So we no longer instantiate them here in the constructor.
        # self.weaviate_processor = WeaviateProcessor() # done later
        # self.weaviate_loader = WeaviateLoader() # done later
        # In the future, you would initialize the GraphRAG processor/loader here
        # self.graphrag_processor = GraphRAGProcessor()
        # self.neo4j_loader = Neo4jLoader(...)

    @staticmethod
    async def get_permanent_url(
        s3_client: ObjectStorage,
        job_data: dict,
        staging_key: str,
    ) -> str:
        """
        Gets the permanent S3 URL with explicit preference and fallback reconstruction.

        Priority:
        1. Use explicit `persisted_s3_url` from payload (most reliable).
        2. Reconstruct from `persisted_s3_key` if available.
        3. Derive from staging key + `user_id` (fallback for legacy jobs).
        4. Use staging URL as last resort (logs warning).
        """
        # Priority 1: Explicit URL (best)
        if job_data.get("persisted_s3_url"):
            logger.debug("Using explicit permanent URL from payload")
            return str(job_data["persisted_s3_url"])

        # Priority 2: Reconstruct from explicit key
        if job_data.get("persisted_s3_key"):
            logger.debug("Reconstructing URL from persisted_s3_key")
            return str(await s3_client.get_https_url(job_data["persisted_s3_key"]))

        # Priority 3: Derive from pattern (fallback for legacy jobs)
        if staging_key.startswith("staging/") and job_data.get("user_id"):
            logger.warning(
                f"Deriving permanent URL for job {job_data.get('job_id')} from staging pattern. "
                f"This is a fallback - consider updating job payload format."
            )
            try:
                settings = SettingsProvider.sget()
                filename = staging_key.split("/")[-1]
                permanent_key = f"{settings.AWS_S3_FILES_FOLDER}/{job_data['user_id']}/{filename}"
                return str(await s3_client.get_https_url(permanent_key))
            except Exception as e:
                logger.error(f"Failed to derive permanent URL: {e}")

        # Priority 4: Last resort (log error and use processing URL)
        processing_url = str(await s3_client.get_https_url(staging_key))
        logger.error(
            f"Cannot determine permanent URL for job {job_data.get('job_id')}. "
            f"Using processing URL: {processing_url} - THIS WILL LIKELY BREAK AFTER CLEANUP!"
        )
        return processing_url

    async def _process_spreadsheet_from_url(
        self,
        file_url: str,
        original_filename: str,
        url_for_metadata: str,
        unstructure_io_client: Optional[UnstructuredIO],
        ctx: Context,
    ) -> Dict[str, Any]:
        """
        Dedicated spreadsheet ingestion pipeline (CSV/TSV/XLS/XLSX/XLSM/XLSB).
        - streams download to tempfile (no bytes-in-memory)
        - rejects on caps (no sampling)
        - outputs header-aware row-batch chunks + workbook/sheet summaries

        Note: unstructure_io_client is unused; kept for compatibility with the special-handler signature.
        """
        from src.autobots.ingestion.spreadsheet.ingest import ingest_spreadsheet_from_url

        res = await ingest_spreadsheet_from_url(
            ctx=ctx,
            file_url=file_url,
            original_filename=original_filename,
            url_for_metadata=url_for_metadata,
            tenant_id=self.tenant_id,
        )
        return res.processed_data

    async def _process_json_from_url(
        self,
        file_url: str,
        original_filename: str,
        url_for_metadata: str,
        unstructure_io_client: UnstructuredIO,
        ctx: Context,
    ) -> Dict[str, Any]:
        """
        JSON special-case ingestion:
          - Download JSON
          - Parse
          - Chunk as (json_path -> value) leaves
          - Produce processed_data compatible with WeaviateLoader/Neo4jLoader
        """
        raw_bytes = await unstructure_io_client._stream_download_file(file_url)
        if raw_bytes is None:
            raise RetryableIngestionError(f"Failed to download JSON file '{original_filename}'.")

        try:
            parsed = json.loads(raw_bytes.decode("utf-8-sig"))  # handle BOM
        except Exception as e:
            # Permanent failure (don’t retry)
            raise ValueError(f"Invalid JSON in '{original_filename}': {e}") from e

        document_id = hashlib.sha256(str(url_for_metadata).encode()).hexdigest()
        document_namespace = uuid.uuid5(uuid.NAMESPACE_DNS, document_id)

        documents: List[Document] = []
        chunk_uuids: List[str] = []

        order = 0
        MAX_VALUE_CHARS = 4000

        for json_path, value in _iter_json_leaves(parsed, "$"):
            if value is None or isinstance(value, (dict, list)):
                continue

            value_str = _sanitize_json_text(value)
            if not value_str:
                continue

            parts = [
                value_str[i : i + MAX_VALUE_CHARS]
                for i in range(0, len(value_str), MAX_VALUE_CHARS)
            ]

            for part_idx, part in enumerate(parts):
                stable_key = json_path if part_idx == 0 else f"{json_path}#part{part_idx}"

                # Keep metadata limited to schema-known props (source_locator is the JSON path-like locator)
                meta = {
                    "source": original_filename,
                    "document_id": document_id,
                    "s3_url": str(url_for_metadata),
                    "order": order,
                    "doc_type": "json",
                    "source_locator": json_path,
                    "page_number": None,
                }
                documents.append(Document(page_content=f"{json_path}: {part}", metadata=meta))
                chunk_uuids.append(str(uuid.uuid5(document_namespace, stable_key)))
                order += 1

        return {
            "documents": documents,
            "images_for_vectorization": [],
            "chunk_uuids": chunk_uuids,
        }

    async def _process_text_like_from_url(
        self,
        file_url: str,
        original_filename: str,
        url_for_metadata: str,
        unstructure_io_client: UnstructuredIO,
        ctx: Context,
        force_doc_type: Optional[str] = None,
    ) -> Dict[str, Any]:
        """
        Code/config/text-like ingestion (no Unstructured):
          - download bytes
          - decode to text
          - split with language-aware splitter (if available) else fallback
        """
        raw_bytes = await unstructure_io_client._stream_download_file(file_url)
        if raw_bytes is None:
            raise RetryableIngestionError(
                f"Failed to download text-like file '{original_filename}'."
            )

        if not _looks_like_text(raw_bytes):
            raise ValueError(
                f"File '{original_filename}' does not appear to be text; cannot ingest as text-like."
            )

        suffix = Path(original_filename).suffix.lower()
        text = _decode_text_best_effort(raw_bytes)

        # lightweight injection token stripping
        for tok in _JSON_INJECTION_TOKENS:
            text = text.replace(tok, "")

        chunks = _split_text_language_aware(text, suffix)
        if not chunks:
            return {"documents": [], "images_for_vectorization": [], "chunk_uuids": []}

        document_id = hashlib.sha256(str(url_for_metadata).encode()).hexdigest()
        document_namespace = uuid.uuid5(uuid.NAMESPACE_DNS, document_id)

        if force_doc_type:
            doc_type = force_doc_type
        else:
            info = _LANGUAGE_MAP.get(suffix)
            if info and info.get("kind") and info.get("name"):
                doc_type = f"{info['kind']}_{info['name']}"
            else:
                doc_type = "text_like"

        documents: List[Document] = []
        chunk_uuids: List[str] = []

        for order, chunk in enumerate(chunks):
            chunk = chunk.strip()
            if not chunk:
                continue
            meta = {
                "source": original_filename,
                "document_id": document_id,
                "s3_url": str(url_for_metadata),
                "order": order,
                "doc_type": doc_type,
                "page_number": None,
            }
            documents.append(Document(page_content=chunk, metadata=meta))

            stable_key = f"{order}:{hashlib.sha1(chunk.encode('utf-8', errors='ignore')).hexdigest()}"
            chunk_uuids.append(str(uuid.uuid5(document_namespace, stable_key)))

        return {
            "documents": documents,
            "images_for_vectorization": [],
            "chunk_uuids": chunk_uuids,
        }

    async def _process_xlsm_from_url(
        self,
        file_url: str,
        original_filename: str,
        url_for_metadata: str,
        unstructure_io_client: UnstructuredIO,
        ctx: Context,
    ) -> Dict[str, Any]:
        """
        DEPRECATED: .xlsm is now routed to the dedicated spreadsheet ingestion pipeline
        via SPECIAL_EXTENSION_HANDLERS (i.e., _process_spreadsheet_from_url). This method
        remains for backward compatibility / reference and may be removed later.

        XLSM ingestion (no Unstructured, no zip repack):
          - download bytes
          - parse with openpyxl in safe mode
          - extract row text into chunks

        Fixes included vs earlier version:
          - no ridx leakage across sheets
          - correct start_row/end_row tracking when skipping empty rows
          - best-effort workbook close
        """
        if openpyxl is None:
            raise ValueError("openpyxl is required to ingest .xlsm files but is not installed.")

        raw_bytes = await unstructure_io_client._stream_download_file(file_url)
        if raw_bytes is None:
            raise RetryableIngestionError(f"Failed to download XLSM file '{original_filename}'.")

        document_id = hashlib.sha256(str(url_for_metadata).encode()).hexdigest()
        document_namespace = uuid.uuid5(uuid.NAMESPACE_DNS, document_id)

        MAX_CHARS_PER_CHUNK = 12000
        MAX_ROWS_PER_CHUNK = 200

        def extract_wb(wb) -> Tuple[List[Document], List[str]]:
            documents: List[Document] = []
            chunk_uuids: List[str] = []
            order = 0

            for ws in wb.worksheets:
                sheet_name = ws.title or "Sheet"
                rows_buf: List[str] = []
                rows_count = 0
                chars_count = 0
                start_row: Optional[int] = None
                end_row_included: Optional[int] = None

                def flush():
                    nonlocal order, rows_buf, rows_count, chars_count, start_row, end_row_included
                    if start_row is None or end_row_included is None or not rows_buf:
                        # nothing buffered
                        rows_buf = []
                        rows_count = 0
                        chars_count = 0
                        start_row = None
                        end_row_included = None
                        return

                    content = (
                        f"Sheet: {sheet_name}\nRows {start_row}-{end_row_included}\n"
                        + "\n".join(rows_buf)
                    )
                    meta = {
                        "source": original_filename,
                        "document_id": document_id,
                        "s3_url": str(url_for_metadata),
                        "order": order,
                        "doc_type": "xlsm",
                        "page_number": None,
                    }
                    documents.append(Document(page_content=content, metadata=meta))
                    stable_key = f"{sheet_name}:{start_row}:{end_row_included}"
                    chunk_uuids.append(str(uuid.uuid5(document_namespace, stable_key)))
                    order += 1

                    rows_buf = []
                    rows_count = 0
                    chars_count = 0
                    start_row = None
                    end_row_included = None

                for ridx, row in enumerate(ws.iter_rows(values_only=True), start=1):

                    # Convert row to a TSV-ish line
                    vals: List[str] = []
                    for v in row:
                        if v is None:
                            vals.append("")
                        else:
                            s = str(v)
                            for tok in _JSON_INJECTION_TOKENS:
                                s = s.replace(tok, "")
                            vals.append(s)
                    line = "\t".join(vals).strip()
                    if not line:
                        continue

                    if start_row is None:
                        start_row = ridx
                    end_row_included = ridx

                    rows_buf.append(line)
                    rows_count += 1
                    chars_count += len(line) + 1

                    if rows_count >= MAX_ROWS_PER_CHUNK or chars_count >= MAX_CHARS_PER_CHUNK:
                        flush()

                # flush remaining buffered lines for this sheet
                flush()

            return documents, chunk_uuids

        wb = None
        try:
            # First try: data_only=True (prefer values), safe mode
            wb = openpyxl.load_workbook(
                filename=io.BytesIO(raw_bytes),
                read_only=True,
                data_only=True,
                keep_vba=False,
            )
            docs, uuids = extract_wb(wb)

            # If the output is suspiciously empty, retry with formulas (data_only=False)
            total_chars = sum(len(d.page_content) for d in docs) if docs else 0
            if total_chars < 200:
                logger.bind(ctx=ctx).warning(
                    f"XLSM extraction for '{original_filename}' yielded very little text "
                    f"({total_chars} chars). Retrying with data_only=False to capture formulas."
                )
                try:
                    wb.close()
                except Exception:
                    pass
                wb = openpyxl.load_workbook(
                    filename=io.BytesIO(raw_bytes),
                    read_only=True,
                    data_only=False,
                    keep_vba=False,
                )
                docs, uuids = extract_wb(wb)

            return {"documents": docs, "images_for_vectorization": [], "chunk_uuids": uuids}
        except Exception as e:
            raise ValueError(f"Could not open/parse XLSM '{original_filename}': {e}") from e
        finally:
            try:
                if wb is not None:
                    wb.close()
            except Exception:
                pass

    async def _get_media_strategy(
        self, collection_name: str, store: Optional[WeaviateStore] = None
    ) -> Literal["vectorize", "summarize", "skip"]:
        """
        Connects to Weaviate to determine the best media handling strategy for a given collection.
        This is the single source of truth for media processing logic.
        """
        media_strategy: Literal["vectorize", "summarize", "skip"] = "skip"

        # Pass skip_init_checks=True to the store factory
        temp_store_for_check = store or await WeaviateStore.create_store(
            collection_name, user_id=self.user_id, skip_init_checks=True
        )

        try:
            collection_config = await asyncio.to_thread(
                temp_store_for_check._client.client.collections.get(collection_name).config.get
            )

            # 1. Check if vectorizer_config exists before trying to access its properties.
            if (
                hasattr(collection_config, "vectorizer_config")
                and collection_config.vectorizer_config is not None
            ):
                # 2. Access the vectorizer name safely.
                vectorizer_name = collection_config.vectorizer_config.vectorizer
            else:
                # 3. If vectorizer_config is None, it means no vectorizer is explicitly set at the top level.
                #    This happens with some multi-tenant configurations or if set to "none".
                vectorizer_name = "none"  # Default to "none" if the attribute doesn't exist.

            if "multi2vec" in vectorizer_name:
                logger.info(f"Collection '{collection_name}' supports multimodal vectorization.")
                media_strategy = "vectorize"
            else:
                # This now correctly handles both "text2vec-openai" and "none".
                logger.info(
                    f"Collection '{collection_name}' is text-only. Falling back to media summarization."
                )
                media_strategy = "summarize"
        except Exception as e:
            logger.error(
                f"Could not determine vectorizer for '{collection_name}', defaulting to skipping media. Error: {e}"
            )
            media_strategy = "skip"
        finally:
            if not store:
                await temp_store_for_check.close()
        return media_strategy

    async def queue_ingestion_from_s3(
        self,
        s3_key: str,
        original_filename: str,
        collection_name: str,
        db: AsyncIOMotorDatabase,
        targets: Optional[List[str]] = None,
    ) -> str:
        """
        Creates a job record, queues the job in Redis, and returns the job ID.
        """
        # If the calling worker code doesn't provide targets,
        # default to both. This ensures backward compatibility while supporting new logic.
        if targets is None:
            targets = ["weaviate", "graphrag"]

        logger.info(
            f"Starting ingestion for S3 key: {s3_key} for user: {self.user_id} -> Targets: {targets}"
        )

        job_id = str(uuid.uuid4())

        # Compute the permanent URL for history display (Nightly Sync Path)
        # Since this method is used for internal/sync jobs where s3_key is already
        # the permanent location, we can derive the URL directly.
        file_url = await self.s3_client.get_https_url(s3_key)

        job_crud = JobStatusCRUD(db)
        await job_crud.create(
            JobStatusCreate(
                _id=job_id,
                user_id=self.user_id,
                kb_id=None,
                tenant_id=self.tenant_id,
                collection_name=collection_name,
                original_filename=original_filename,
                targets=targets,
                s3_url=str(file_url),
            )
        )

        # The worker needs the job_id to update its status

        job_payload = {
            "job_id": job_id,
            "user_id": self.user_id,
            "tenant_id": self.tenant_id,
            "s3_key": str(s3_key),
            "original_filename": original_filename,
            "collection_name": collection_name,
            "targets": targets,  # Pass to worker for logic branching
        }
        r = await get_async_redis()
        await r.lpush(
            SettingsProvider.sget().REDIS_INGESTION_HIGH_PRIORITY_QUEUE,
            json.dumps(job_payload),
        )
        return job_id

    async def _execute_ingestion_job(
        self,
        s3_key: str,
        original_filename: str,
        collection_name: str,
        db: AsyncIOMotorDatabase,
        targets: List[str],
        user: UserORM,
        ctx: Context,
        app_state: AppState,
        store: Optional[WeaviateStore] = None,
        job_data: Optional[Dict] = None,  # Accept full job payload
    ) -> dict:
        """
        REFACTORED for sequential execution: Weaviate first, then GraphRAG.
        Returns a dictionary with results for the worker to interpret.
        """
        """Internal method for the worker to execute the full ingestion pipeline."""
        # NOTE: Unstructured client is acquired lazily (spreadsheets do not require it).
        unstructure_io_client: Optional[UnstructuredIO] = None

        _job_data = job_data or {}
        processed_data = None  # This will hold the common chunk data
        successful_targets = []
        errors = []
        _store = None  # Ensure _store is defined in the outer scope

        # URL for processing remains the staging URL for isolation
        processing_url = await self.s3_client.get_https_url(s3_key)
        if not processing_url:
            raise RuntimeError(f"Could not construct public URL for S3 key: {s3_key}")

        # The URL for metadata is the authoritative one, resolved by the helper
        url_for_metadata = await self.get_permanent_url(self.s3_client, _job_data, s3_key)

        # Defensively ensure we are working with a string for the rest of the function
        url_for_metadata = str(url_for_metadata)

        logger.bind(ctx=ctx).debug(
            f"[URL TRACE 2/4] Service received - Processing: {processing_url}, Metadata: {url_for_metadata}"
        )

        suffix = Path(original_filename).suffix.lower()

        # FIX: Create the store ONCE and pass it to ALL child functions
        # This prevents each child from creating its own store instance
        _store = store or await WeaviateStore.create_store(
            collection_name=collection_name,
            user_id=self.tenant_id,
            # Spreadsheet-friendly batching:
            # spreadsheets can produce token-heavy chunks; Weaviate's native vectorizer batches
            # many objects into a single embeddings request with an aggregate token cap.
            # Use a smaller batch size only for spreadsheet ingestion.
            text_batch_size=(38 if _is_spreadsheet_suffix(suffix) else None),
        )

        try:
            # Fail fast: blocked extensions BEFORE any external lookups (like Unstructured key).
            if suffix in BLOCKED_EXTENSIONS:
                raise ValueError(
                    f"File type '{suffix}' is not supported (blocked). "
                    f"Please convert '{original_filename}' to a supported format such as PDF, DOCX, TXT, HTML, CSV, XLSX, or JSON."
                )

            # Determine whether we need Unstructured for this file type.
            handler_name = SPECIAL_EXTENSION_HANDLERS.get(suffix)
            needs_unstructured = handler_name != "_process_spreadsheet_from_url"

            if needs_unstructured:
                client_result = await AppConn.get_unstructured_io(ctx=ctx, user=user, db=db)
                if client_result is None or isinstance(client_result, Exception):
                    logger.bind(ctx=ctx).error(
                        f"AppConn failed to return valid client: {client_result}"
                    )
                    raise ValueError(
                        "Unstructured.io API Key not found or invalid. Please connect your API key."
                    )
                unstructure_io_client = client_result

            # --- STAGE 1: Process file into a unified format (documents + chunk_uuids) ---
            # Special handler dispatch (JSON now; extensible later)
            if handler_name:
                handler = getattr(self, handler_name, None)
                if not handler:
                    raise RuntimeError(f"Missing special handler '{handler_name}' for '{suffix}'.")
                processed_data = await handler(
                    file_url=str(processing_url),
                    original_filename=original_filename,
                    url_for_metadata=str(url_for_metadata),
                    unstructure_io_client=unstructure_io_client,  # may be None for spreadsheets
                    ctx=ctx,
                )

            # Media pipeline
            elif suffix in (IMAGE_EXTENSIONS | AUDIO_EXTENSIONS | VIDEO_EXTENSIONS):
                if unstructure_io_client is None:
                    raise ValueError(
                        "Unstructured client is required for media ingestion in this implementation."
                    )
                processed_data = await self._ingest_direct_media_from_url(
                    str(processing_url),
                    original_filename,
                    collection_name,
                    user,
                    db,
                    ctx,
                    unstructure_io_client,
                    _store,
                    force_strategy="summarize",
                    url_for_metadata=str(url_for_metadata),
                )

            # Code/config pipeline (no Unstructured partitioning; uses Unstructured downloader in current implementation)
            elif suffix in TEXT_LIKE_EXTENSIONS:
                if unstructure_io_client is None:
                    raise ValueError(
                        "Unstructured client is required for text-like ingestion in this implementation."
                    )
                processed_data = await self._process_text_like_from_url(
                    file_url=str(processing_url),
                    original_filename=original_filename,
                    url_for_metadata=str(url_for_metadata),
                    unstructure_io_client=unstructure_io_client,
                    ctx=ctx,
                )
            else:
                # Best-effort: try Unstructured first for unknown extensions (not blocked).
                if suffix and (suffix not in DOC_LIKE_EXTENSIONS) and (suffix not in BEST_EFFORT_KNOWN_EXTENSIONS):
                    logger.bind(ctx=ctx).warning(
                        f"Unknown extension '{suffix}' for '{original_filename}'. "
                        "Attempting document ingestion anyway."
                    )
                try:
                    if unstructure_io_client is None:
                        raise ValueError(
                            "Unstructured client is required for document ingestion in this implementation."
                        )
                    processed_data = await self._process_document_from_url(
                        str(processing_url),
                        original_filename,
                        collection_name,
                        user,
                        db,
                        ctx,
                        unstructure_io_client,
                        _store,
                        url_for_metadata=str(url_for_metadata),
                    )
                except RetryableIngestionError:
                    # Let the worker retry transient issues
                    raise
                except Exception as e:
                    # IMPORTANT: fallback on exceptions ONLY for unknown extensions (not doc-like)
                    if (not suffix) or (suffix not in DOC_LIKE_EXTENSIONS):
                        logger.bind(ctx=ctx).warning(
                            f"Unstructured raised for unknown extension '{suffix}' "
                            f"({original_filename}): {e}. Falling back to text-like ingestion."
                        )
                        processed_data = await self._process_text_like_from_url(
                            file_url=str(processing_url),
                            original_filename=original_filename,
                            url_for_metadata=str(url_for_metadata),
                            unstructure_io_client=unstructure_io_client,
                            ctx=ctx,
                            force_doc_type="text_like_fallback",
                        )
                    else:
                        raise

                # If Unstructured yields nothing for unknown extensions, fall back to text-like
                if (not processed_data or not processed_data.get("documents")) and (
                    (not suffix) or (suffix not in DOC_LIKE_EXTENSIONS)
                ):
                    logger.bind(ctx=ctx).warning(
                        f"Unstructured produced no elements for '{original_filename}' ({suffix}). "
                        "Falling back to text-like ingestion."
                    )
                    processed_data = await self._process_text_like_from_url(
                        file_url=str(processing_url),
                        original_filename=original_filename,
                        url_for_metadata=str(url_for_metadata),
                        unstructure_io_client=unstructure_io_client,
                        ctx=ctx,
                        force_doc_type="text_like_fallback",
                    )

            if not processed_data or not processed_data.get("documents"):
                raise RetryableIngestionError(
                    f"File '{original_filename}' did not yield any processable text chunks."
                )

            # --- STAGE 2: MANDATORY Weaviate Ingestion ---
            if "weaviate" in targets:
                try:
                    weaviate_loader = WeaviateLoader()
                    load_successful = await weaviate_loader.load(
                        processed_data, self.user_id, collection_name, store=_store
                    )
                    if load_successful:
                        successful_targets.append("weaviate")
                    else:
                        raise Exception(
                            "WeaviateLoader returned False, indicating a failure to load documents."
                        )
                except Exception as e:
                    logger.exception(
                        "CRITICAL: Weaviate loading failed. Aborting further processing for this job."
                    )
                    errors.append(f"Weaviate Error: {str(e)}")
                    # Return immediately as Weaviate is a hard requirement
                    return {
                        "processed_data": processed_data,
                        "successful_targets": [],
                        "store": _store,
                        "errors": errors,
                    }

            # --- STAGE 3: OPTIONAL GraphRAG Ingestion ---
            # This stage only runs if Weaviate succeeded and 'graphrag' is a target.
            if "graphrag" in targets:
                try:
                    logger.info(
                        f"Processing content from '{original_filename}' for GraphRAG target."
                    )
                    # Use the permanent URL for the document ID
                    unique_content_identifier = url_for_metadata
                    document_id = hashlib.sha256(
                        str(unique_content_identifier).encode()
                    ).hexdigest()
                    neo4j_loader = Neo4jLoader()
                    # Build the list of URLs corresponding to each document/chunk
                    chunk_s3_urls = [d.metadata.get("s3_url") for d in processed_data["documents"]]

                    # Spreadsheet-aware strategy:
                    # - allow LLM graph extraction only for small spreadsheets (or if mode forces it)
                    # - otherwise use deterministic graph builder (no LLM) for Neo4j chunk adjacency
                    if _is_spreadsheet_suffix(suffix):
                        from src.autobots.ingestion.spreadsheet.ingest import (
                            SPREADSHEET_LLM_MAX_BYTES,
                            SPREADSHEET_LLM_MAX_CHUNKS,
                            SPREADSHEET_LLM_MAX_TOTAL_CHARS,
                            build_deterministic_spreadsheet_graph,
                            choose_spreadsheet_neo4j_strategy,
                        )

                        file_size_bytes = int(
                            processed_data.get("spreadsheet_file_size_bytes") or 0
                        )
                        chunk_count = len(processed_data.get("documents") or [])
                        total_chars = int(processed_data.get("spreadsheet_total_chars") or 0)
                        strat = choose_spreadsheet_neo4j_strategy(file_size_bytes, chunk_count, total_chars)

                        if strat == "llm":
                            # Reject rather than sample if too large for LLM
                            if (
                                file_size_bytes > SPREADSHEET_LLM_MAX_BYTES
                                or chunk_count > SPREADSHEET_LLM_MAX_CHUNKS
                                or total_chars > SPREADSHEET_LLM_MAX_TOTAL_CHARS
                            ):
                                raise ValueError(
                                    "Spreadsheet is too large for LLM graph extraction. "
                                    f"(file_size_bytes={file_size_bytes}, max_bytes={SPREADSHEET_LLM_MAX_BYTES}, "
                                    f"chunks={chunk_count}, max_chunks={SPREADSHEET_LLM_MAX_CHUNKS}, "
                                    f"total_chars={total_chars}, max_total_chars={SPREADSHEET_LLM_MAX_TOTAL_CHARS}). "
                                    "Please split by sheet/table or export to CSV."
                                )

                            graph_rag_processor = GraphRAGProcessor(
                                analyze_tokens=self.analyze_tokens
                            )
                            graph_metadata = {
                                "document_id": document_id,
                                "s3_url": str(url_for_metadata),
                            }
                            extracted_graph = await graph_rag_processor.process(
                                processed_data["documents"], graph_metadata, user, db, ctx, app_state
                            )
                            load_successful = await neo4j_loader.load(
                                processed_data=extracted_graph,
                                user_id=self.tenant_id,
                                collection_name=collection_name,
                                document_id=document_id,
                                chunk_uuids=processed_data["chunk_uuids"],
                                chunk_s3_urls=chunk_s3_urls,
                                s3_url=str(url_for_metadata),
                            )
                            if load_successful:
                                successful_targets.append("graphrag")
                            else:
                                errors.append("Neo4jLoader returned False for LLM spreadsheet graph.")
                            # Continue to next stage (skip the single-load block below)
                            return {
                                "processed_data": processed_data,
                                "successful_targets": successful_targets,
                                "store": _store,
                                "errors": errors,
                            }

                        # deterministic: load per-sheet group to avoid cross-sheet NEXT edges created by Neo4jClient.add_graph
                        docs_all = processed_data["documents"]
                        uuids_all = processed_data["chunk_uuids"]
                        if len(docs_all) != len(uuids_all):
                            raise ValueError("Spreadsheet docs and chunk_uuids length mismatch.")

                        def _sheet_from_locator(loc: str) -> str:
                            import re
                            m = re.search(r"Sheet=([^;]+)", loc or "")
                            raw = (m.group(1).strip() if m else "__UNKNOWN__")
                            # Best-effort decode in case spreadsheet ingestion encodes sheet names.
                            # Safe even if not encoded.
                            try:
                                raw = unquote(raw)
                            except Exception:
                                pass
                            return raw[:200]

                        # Group indices by (sheet, table) to avoid cross-table NEXT edges.
                        by_group: Dict[Tuple[str, int], List[int]] = {}
                        for i, d in enumerate(docs_all):
                            md = d.metadata or {}
                            loc = str(md.get("source_locator") or "")
                            info = parse_source_locator(loc)
                            sheet = (info.sheet or "").strip()
                            # parse_source_locator() returns "__UNKNOWN__" when Sheet= is missing,
                            # which is truthy, so we must explicitly treat it as missing.
                            if (not sheet) or (sheet == "__UNKNOWN__"):
                                sheet = _sheet_from_locator(loc)
                            sheet = sheet[:200]
                            table = int(info.table) if info.table is not None else 1
                            by_group.setdefault((sheet, table), []).append(i)

                        any_fail = False
                        for (sheet, table), idxs in by_group.items():
                            # Ensure stable semantic ordering within each sheet subgraph.
                            idxs.sort(
                                key=lambda i: int((docs_all[i].metadata or {}).get("order", i))
                            )
                            # NOTE (readiness probe contract):
                            # The worker's readiness_probe() verifies Neo4j availability by checking for the presence
                            # of the job's chunk_uuids (UUID-specific), not just "row chunks".
                            # See polling.readiness_probe/_probe_neo4j: MATCH (n:Chunk) ... n.id IN $chunkIds.
                            # If deterministic Neo4j ingestion indexed only spreadsheet_rows, notes/summaries-heavy
                            # spreadsheets would fail the probe and be marked PARTIAL_AVAILABLE even though row
                            # adjacency is correct.
                            #
                            # Therefore we index BOTH rows and meta chunks in Neo4j, but load them as separate Neo4j
                            # documents so any client-side auto-NEXT chaining does not interleave notes with row chains.
                            # Variant A2: load row chunks and meta chunks separately to keep NEXT chains clean,
                            # while still indexing *all* chunk UUIDs so readiness_probe matches expectations.
                            rows_idxs = [
                                i for i in idxs
                                if str((docs_all[i].metadata or {}).get("doc_type") or "") == "spreadsheet_rows"
                            ]
                            meta_idxs = [
                                i for i in idxs
                                if str((docs_all[i].metadata or {}).get("doc_type") or "") != "spreadsheet_rows"
                            ]

                            async def _load_subgraph(sub_idxs: List[int], kind: str) -> bool:
                                if not sub_idxs:
                                    return True
                                sub_docs = [docs_all[i] for i in sub_idxs]
                                sub_uuids = [uuids_all[i] for i in sub_idxs]
                                sub_processed = {"documents": sub_docs, "chunk_uuids": sub_uuids}
                                sub_doc_id = hashlib.sha256(
                                    f"{document_id}:{sheet}:table{table}:{kind}".encode()
                                ).hexdigest()

                                extracted_graph = build_deterministic_spreadsheet_graph(
                                    processed_data=sub_processed,
                                    document_id=sub_doc_id,
                                    sheet_namespace_id=document_id,
                                )
                                sub_chunk_s3_urls = [d.metadata.get("s3_url") for d in sub_docs]
                                return await neo4j_loader.load(
                                    processed_data=extracted_graph,
                                    user_id=self.tenant_id,
                                    collection_name=collection_name,
                                    document_id=sub_doc_id,
                                    chunk_uuids=sub_uuids,
                                    chunk_s3_urls=sub_chunk_s3_urls,
                                    s3_url=str(url_for_metadata),
                                )

                            ok_rows = await _load_subgraph(rows_idxs, "rows")
                            ok_meta = await _load_subgraph(meta_idxs, "meta")
                            if not (ok_rows and ok_meta):
                                any_fail = True

                        if not any_fail:
                            successful_targets.append("graphrag")
                        else:
                            errors.append("One or more deterministic spreadsheet subgraphs failed to load.")

                        # Spreadsheet deterministic path complete
                        # NOTE: GraphRAG is the final stage today; this early return is safe.
                        # If stages are added after GraphRAG in the future, refactor to a single exit path.
                        return {
                            "processed_data": processed_data,
                            "successful_targets": successful_targets,
                            "store": _store,
                            "errors": errors,
                        }
                    else:
                        graph_rag_processor = GraphRAGProcessor(analyze_tokens=self.analyze_tokens)
                        graph_metadata = {
                            "document_id": document_id,
                            "s3_url": str(url_for_metadata),
                        }
                        extracted_graph = await graph_rag_processor.process(
                            processed_data["documents"], graph_metadata, user, db, ctx, app_state
                        )

                    load_successful = await neo4j_loader.load(
                        processed_data=extracted_graph,
                        user_id=self.tenant_id,
                        collection_name=collection_name,
                        document_id=document_id,
                        chunk_uuids=processed_data["chunk_uuids"],
                        chunk_s3_urls=chunk_s3_urls,  # Pass the list of URLs for each chunk
                        # Pass s3_url to the Neo4j loader
                        s3_url=str(url_for_metadata),
                    )
                    if load_successful:
                        successful_targets.append("graphrag")
                    else:
                        errors.append(
                            "Neo4jLoader returned False, indicating a failure to load the graph."
                        )
                except Exception as e:
                    logger.exception(
                        "GraphRAG/Neo4j processing failed. This is a non-critical failure."
                    )
                    errors.append(f"Neo4j Error: {str(e)}")

            return {
                "processed_data": processed_data,
                "successful_targets": successful_targets,
                "store": _store,
                "errors": errors,
            }
        finally:
            # --- FIX: Do NOT close the store here. The worker is now responsible. ---
            if not store and _store:
                pass

    async def _process_document_from_url(
        self,
        file_url: str,
        original_filename: str,
        collection_name: str,
        user: UserORM,
        db: AsyncIOMotorDatabase,
        ctx: Context,
        unstructure_io_client: UnstructuredIO,
        store: Optional[WeaviateStore] = None,
        url_for_metadata: Optional[str] = None,
    ):
        """Internal helper to process a document. Implements the dual parsing strategy."""
        settings = SettingsProvider.sget()
        strategy = settings.INGESTION_WORKER_PARSING_STRATEGY
        logger.info(
            f"Processing document '{original_filename}' with initial strategy: '{strategy}'"
        )

        _store = store or await WeaviateStore.create_store(
            collection_name=collection_name, user_id=self.tenant_id
        )

        try:
            elements = None
            final_strategy_used = strategy

            if strategy == "fast_with_gemini":
                try:
                    # 1. Pre-parse with Gemini, guarded by the module-level semaphore
                    async with GEMINI_DOC_PARSE_SEMAPHORE:
                        gemini_text_content = await pre_parse_document_with_gemini(
                            user, db, ctx, file_url, original_filename
                        )

                    # 2. Set Unstructured parameters for 'fast' mode with consistent chunking
                    partition_params = PartitionParametersParams(
                        strategy="fast",
                        chunking_strategy="by_title",
                        languages=["eng"],
                        combine_under_n_chars=2500,
                        new_after_n_chars=7500,
                        max_characters=7500,
                    )

                    # 3. Pass in-memory text to Unstructured with a neutral filename
                    elements = await unstructure_io_client.partition_from_content(
                        file_content=gemini_text_content.encode("utf-8"),
                        filename=f"{Path(original_filename).stem}.txt",
                        partition_parameters=partition_params,
                    )

                    # 4. If Unstructured fast mode yields no elements, also trigger fallback
                    if not elements:
                        logger.warning(
                            "Unstructured 'fast' mode yielded no elements from Gemini text. Falling back to 'auto' strategy."
                        )
                        strategy = "auto"
                        final_strategy_used = "auto_fallback"

                except Exception as e:
                    logger.bind(ctx=ctx).warning(
                        f"Gemini pre-parsing failed for '{original_filename}': {e}. "
                        "Falling back to 'auto' strategy for resilience."
                    )
                    # Set strategy to auto to trigger the fallback logic
                    strategy = "auto"
                    final_strategy_used = "auto_fallback"

            # This block now serves as both the primary 'auto' path and the fallback path
            if strategy == "auto":
                final_strategy_used = "auto" if final_strategy_used == "auto" else "auto_fallback"
                partition_params = PartitionParametersParams(
                    strategy="auto",
                    chunking_strategy="by_title",
                    extract_image_block_types=["Image"],
                    pdf_infer_table_structure=True,
                    languages=["eng"],
                    combine_under_n_chars=2500,
                    new_after_n_chars=7500,
                    max_characters=7500,
                )
                elements = await unstructure_io_client.partition_from_remote_url(
                    remote_url=file_url, partition_parameters=partition_params
                )

                # Restore the one-time retry specifically for the auto path for transient issues
                if not elements and final_strategy_used != "auto_fallback":
                    logger.warning(
                        "Unstructured 'auto' mode returned no elements; retrying once after 5s."
                    )
                    await asyncio.sleep(5)
                    elements = await unstructure_io_client.partition_from_remote_url(
                        remote_url=file_url, partition_parameters=partition_params
                    )

            logger.bind(ctx=ctx).info(
                f"Document '{original_filename}' was parsed using final strategy: '{final_strategy_used}'"
            )

            # --- COMMON LOGIC AFTER PARSING ---
            if not elements:  # Using 'not elements' handles both None and empty list
                logger.error(
                    f"Document '{original_filename}' could not be processed into elements after all attempts."
                )
                return None

            media_strategy = await self._get_media_strategy(collection_name, _store)
            weaviate_processor = WeaviateProcessor(media_strategy=media_strategy)

            # Use the authoritative URL for the document ID to ensure determinism
            final_metadata_url = str(url_for_metadata or file_url)
            unique_content_identifier = final_metadata_url

            document_id = hashlib.sha256(str(unique_content_identifier).encode()).hexdigest()

            # Pass the s3_url in the metadata dictionary
            processor_metadata = {
                "filename": original_filename,
                "document_id": document_id,
                "s3_url": final_metadata_url,
            }
            return await weaviate_processor.process(elements, processor_metadata, user, db, ctx)

        finally:
            # If we created the store in this function, we must close it.
            if not store and _store:
                await _store.close()

    async def _ingest_direct_media_from_url(
        self,
        file_url: str,
        original_filename: str,
        collection_name: str,
        user: UserORM,
        db: AsyncIOMotorDatabase,
        ctx: Context,
        unstructure_io_client: UnstructuredIO,
        store: Optional[WeaviateStore] = None,
        force_strategy: Optional[Literal["summarize"]] = None,
        url_for_metadata: Optional[str] = None,
    ) -> Optional[dict]:
        """Handles media files. Now ONLY responsible for extracting text chunks (summaries)."""
        logger.info(f"Processing standalone media file: '{original_filename}'")
        media_strategy = await self._get_media_strategy(collection_name, store)

        # Allow the caller (like the GraphRAG path) to force summarization
        if force_strategy == "summarize":
            logger.info(f"Strategy override: Forcing 'summarize' for {original_filename}.")
            media_strategy = "summarize"

        if media_strategy == "skip":
            logger.warning(f"Skipping media file '{original_filename}' due to 'skip' strategy.")
            return

        _store = store or await WeaviateStore.create_store(
            collection_name=collection_name, user_id=self.tenant_id
        )
        summary_doc, chunk_uuid = None, None  # Initialize for the return statement
        try:
            metadata = [{"source": original_filename}]
            suffix = Path(original_filename).suffix.lower()

            if media_strategy == "vectorize":
                logger.info(f"Attempting to directly vectorize media: {original_filename}")
                if suffix in IMAGE_EXTENSIONS:
                    await _store.add_images(inputs=[file_url], metadatas=metadata)
                elif suffix in AUDIO_EXTENSIONS:
                    await _store.add_audios(inputs=[file_url], metadatas=metadata)
                elif suffix in VIDEO_EXTENSIONS:
                    await _store.add_videos(inputs=[file_url], metadatas=metadata)

            elif media_strategy == "summarize":
                logger.info(f"Summarizing/transcribing media file as text: {original_filename}")
                summary = ""

                media_type = ""
                if suffix in IMAGE_EXTENSIONS:
                    media_type = "image"
                elif suffix in AUDIO_EXTENSIONS:
                    media_type = "audio"
                elif suffix in VIDEO_EXTENSIONS:
                    media_type = "video"

                if media_type:
                    # For images, we need to download the content first to get the base64 string
                    if media_type == "image":
                        file_content = await unstructure_io_client._stream_download_file(file_url)
                        if file_content is None:
                            # Treat as transient so the worker’s outer retry–loop will re-run
                            raise RetryableIngestionError("image_download_failed")

                        b64_string = base64.b64encode(file_content).decode("utf-8")
                        summary = await summarize_media_content(
                            "image",
                            user,
                            db,
                            ctx,
                            content_b64=b64_string,
                            original_filename=original_filename,
                        )
                    else:  # For audio/video, the utility function handles the URL directly
                        summary = await summarize_media_content(
                            media_type,
                            user,
                            db,
                            ctx,
                            content_url=file_url,
                            original_filename=original_filename,
                        )

                # Ingest the summary/transcript as a standard text document.
                from src.autobots.conn.chroma.chroma import Document

                # Generate a UUID for the new summary document, just like in weaviate_processor
                final_metadata_url = str(url_for_metadata or file_url)
                document_id = hashlib.sha256(str(final_metadata_url).encode()).hexdigest()

                doc_metadata = {
                    "source": original_filename,
                    "element_type": f"{Path(original_filename).suffix[1:].upper()}Summary",
                    "document_id": document_id,
                    "s3_url": final_metadata_url,
                }
                summary_doc = Document(page_content=summary, metadata=doc_metadata)
                doc_namespace = uuid.uuid5(uuid.NAMESPACE_DNS, document_id)
                chunk_uuid = str(uuid.uuid5(doc_namespace, summary_doc.page_content))
                # DO NOT load here. Loading is now centralized in _execute_ingestion_job.
                # This function just prepares the data.
        finally:
            if not store:
                await _store.close()
        # Return a dict in the same shape as document processing for consistency
        return (
            {"documents": [summary_doc], "chunk_uuids": [chunk_uuid]}
            if summary_doc and chunk_uuid
            else None
        )
